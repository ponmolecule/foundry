"""r418: exhibits resolve reporting membership by product ID; the allowance uses the engine's own opening allowance;
missing legacy members are reported, never looped on."""
import copy
import io
import json
import os
import subprocess
import unittest
from pathlib import Path

from openpyxl import load_workbook

from .bpt_cover import build_bpt_cover
from .run_q import run_v2

_FIX = os.path.join(os.path.dirname(__file__), "..", "fixtures", "universal_template_bank.json")


def _rows(ws):
    return {r[0]: list(r[1:]) for r in ws.iter_rows(values_only=True) if r and r[0]}


class ReportingIdentityTests(unittest.TestCase):
    def test_id_groups_drive_every_exhibit_through_renames(self):
        with open(_FIX, encoding="utf-8") as fh:
            cfg = json.load(fh)
        dep = cfg["assumptions"]["deposit_products"]
        obs = cfg["assumptions"]["obs_exposures"]
        dep[0]["product_id"], dep[1]["product_id"], obs[2]["product_id"] = "p-lower", "p-upper", "p-fees"
        cfg["reporting_manifest"] = {"reporting_products": [{"id": "rp-oa", "name": "Operating Accounts", "line_overrides": {},
            "members": [{"family": "deposit", "product_id": "p-lower", "name": "stale label"},
                        {"family": "deposit", "product_id": "p-upper", "name": "MMDA / Savings"},
                        {"family": "obs", "product_id": "p-fees", "name": obs[2]["name"]}]}]}
        dep[0]["name"] = "Operating Accounts — Lower Tier"      # renamed after grouping
        res = run_v2(copy.deepcopy(cfg))
        self.assertTrue(res["reporting"]["ok"], res["reporting"]["problems"])
        wb = load_workbook(io.BytesIO(build_bpt_cover(cfg, res)[0]), data_only=True)
        nii, nie, pm = (_rows(wb[s]) for s in ("Non-Interest Income", "Non-Interest Expense", "Product Map & Notes"))
        self.assertIn("Operating Accounts", nii)
        self.assertNotIn(obs[2]["name"], nii, "member not listed separately")
        self.assertNotIn("Operating Accounts — Lower Tier Fees", nii)
        self.assertIn("Cost of Revenue: Operating Accounts", nie)
        self.assertNotIn(f"Cost of Revenue: {obs[2]['name']}", nie)
        for name in ("Operating Accounts — Lower Tier", "MMDA / Savings", obs[2]["name"]):
            self.assertEqual(pm[name][6], "Operating Accounts", name)
        self.assertEqual(pm["Brokered"][6], "Itself")
        rp = _rows(wb["Reporting Products"])
        self.assertIn("Operating Accounts — reporting product ($000)", rp)
        members_row = next(k for k in rp if str(k).startswith("Members: "))
        self.assertIn("Operating Accounts — Lower Tier", members_row)

    def test_entered_opening_allowance_reconciles(self):
        from .tests_r369_lending_contracts import config
        cfg = config(n=24)
        loan = cfg["assumptions"]["lending_products"][0]
        loan["loan_rollforward"]["opening_allowance"] = 123456
        loan["product_id"] = "p-loan"
        cfg["reporting_manifest"] = {"reporting_products": [{"id": "rp-l", "name": "Business Lending", "line_overrides": {},
            "members": [{"family": "lending", "product_id": "p-loan", "name": loan["name"]}]}]}
        res = run_v2(copy.deepcopy(cfg))
        prod = next(p for p in res["products"] if p["name"] == loan["name"])
        self.assertEqual(prod["alllOpen"], [123.46], "engine's own opening allowance, $000s")
        v = res["reporting"]
        self.assertTrue(v["ok"], [p["message"] for p in v["problems"]] +
                        [m["class"] for m in v["reconciliation"]["measures"] if not m["ok"]])
        allow = next(m for m in v["reconciliation"]["measures"] if m["class"] == "allowance")
        self.assertAlmostEqual(allow["engine_products"][0], 123.46, places=6)
        self.assertAlmostEqual(allow["engine_bank"][0], 123.46, places=6)
        self.assertTrue(build_bpt_cover(cfg, res)[0][:2] == b"PK", "valid engagement exports")

    def test_browser(self):
        root = Path(__file__).resolve().parents[2]
        r = subprocess.run(["node", str(root / "web/tests_r418_reporting_identity.cjs")], capture_output=True,
                           text=True, cwd=str(root), timeout=480)
        self.assertEqual(r.returncode, 0, (r.stdout[-1500:] + r.stderr[-1500:]))


if __name__ == "__main__":
    unittest.main(verbosity=2)
