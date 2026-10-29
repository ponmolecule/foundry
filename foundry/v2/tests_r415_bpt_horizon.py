"""Business Plan Tables follow forecast scope, independently of filing scope."""
import copy
import io
import unittest
from openpyxl import load_workbook
from .bpt_cover import _annual_periods, _quarterly_periods, build_bpt_cover, build_bpt_cover_legacy
from .bpt_enhanced import _Model
from .tests_bpt_enhanced import payments_engagement, rows
from .run_q import run_v2


class ForecastExportTests(unittest.TestCase):
    def test_full_workbooks(self):
        for ppy, years in [(12, 5), (4, 7), (4, 2)]:
            with self.subTest(ppy=ppy, years=years):
                cfg = payments_engagement()
                cfg['assumptions'].update(periods_per_year=ppy, n_periods=ppy * years)
                cfg.setdefault('charter_profile', {})['submission_quarters'] = 4
                res = run_v2(copy.deepcopy(cfg))
                original = copy.deepcopy((cfg, res))
                m = _Model(cfg, res)
                self.assertEqual(len(m.years), years)
                self.assertEqual(len(m.quarterly) - 1, years * 4)
                self.assertAlmostEqual(m.window_total(res['financials']['is']['ni']), sum(res['financials']['is']['ni']))
                for build in (build_bpt_cover, build_bpt_cover_legacy):
                    data, _ = build(cfg, res)
                    wb = load_workbook(io.BytesIO(data), data_only=True)
                    texts = [v for ws in wb for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str)]
                    self.assertIn(f'Year {years}', texts)
                    self.assertTrue(f'Y{years} Q4' in texts or 'Q4' in texts)
                wb = load_workbook(io.BytesIO(build_bpt_cover(cfg, res)[0]), data_only=True)
                ni = rows(wb['Income Statement (Annual)'])['Net Income']
                self.assertAlmostEqual(ni[-1], sum(res['financials']['is']['ni'][-ppy:]), places=5)
                self.assertEqual((cfg, res), original)

    def test_partial_and_aggregation(self):
        annual = _annual_periods(12, 14)
        quarterly = _quarterly_periods(12, 14)
        self.assertEqual(annual[-1]['label'], 'Year 2 (partial)')
        self.assertEqual(quarterly[-1]['label'], 'Y2 Q1 (partial)')
        self.assertEqual(quarterly[-1]['stock_idx'], 13)
        self.assertEqual(quarterly[-1]['flow_idxs'], [12, 13])
        self.assertEqual([i for p in quarterly[1:] for i in p['flow_idxs']], list(range(14)))
        self.assertEqual([p['stock_idx'] for p in _annual_periods(1, 5)[1:]], list(range(5)))


if __name__ == '__main__':
    unittest.main()
