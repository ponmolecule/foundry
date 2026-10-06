"""r304 Other Assets stock, funding, exact audit and backward-compatibility tests."""
import copy,json,io,unittest
from openpyxl import load_workbook
from .other_assets import prepare_other_assets,other_asset_period
from .run_q import run_v2
from .parity import run_parity
from .audit_workbook import calculation_audit_workbook
from .present import derived_lines

def flat(x):return {'source':'entered','trajectory':'flat','value':x}
def ar_model(days=25,opening=0):
 return {'opening_balance':opening,'base_spec':flat(0),'components':[{'component_id':'ar','name':'Accounts receivable','terms':[{'term_id':'fees-ar','driver_spec':{'source':'link','link':{'kind':'bank_income_flow','series_id':'fees','aggregation':'sum'}},'multiplier_spec':flat(1),'days_spec':flat(days),'year_days':360}]}]}
class Tests(unittest.TestCase):
 def test_days_monthly_quarterly_and_release(self):
  for ppy,fees,expected in [(12,49538.4,41282),(4,148615.2,41282)]:
   a={'other_assets_mode':'formula_level','other_assets_model':ar_model()};p=prepare_other_assets(a,2,ppy)
   self.assertAlmostEqual(other_asset_period(p,0,income_flows={'fees':fees})['total'],expected)
   self.assertEqual(other_asset_period(p,1,income_flows={'fees':0})['total'],0)
 def test_entered_levels_and_signed_terms(self):
  m={'opening_balance':300,'base_spec':flat(100),'components':[{'component_id':'x','name':'Other receivable','terms':[{'term_id':'x1','driver_spec':{'source':'entered','trajectory':'explicit','cadence':'month','values':[500,200,0]},'multiplier_spec':flat(1)},{'term_id':'x2','driver_spec':flat(50),'multiplier_spec':flat(-1)}]}]}
  p=prepare_other_assets({'other_assets_model':m},3,12)
  self.assertEqual([other_asset_period(p,i)['total'] for i in range(3)],[550,250,50])
 def test_validation(self):
  for change in ('missing','cycle','negative-days','zero-year','duplicate'):
   m=ar_model();t=m['components'][0]['terms'][0]
   if change=='missing':t['driver_spec']['link']['series_id']='missing'
   if change=='cycle':t['driver_spec']['link']['series_id']='ni'
   if change=='negative-days':t['days_spec']=flat(-1)
   if change=='zero-year':t['year_days']=0
   if change=='duplicate':m['components'].append(copy.deepcopy(m['components'][0]))
   with self.assertRaises(ValueError):prepare_other_assets({'other_assets_model':m},3,12)
 def test_bank_stock_cash_and_audit(self):
  c=json.loads(__import__('pathlib').Path('foundry/fixtures/core_bank_test_base.json').read_text());a=c['assumptions'];a['cash_yield']=0;a['securities_yield']=0;a['surplus_allocation_policy']='cash'
  # No independent cash-linked component interest in this fixture.
  b=run_parity(c);d=copy.deepcopy(c);d['assumptions']['other_assets_mode']='formula_level';d['assumptions']['other_assets_model']=ar_model(opening=10000);r=run_parity(d)
  self.assertEqual(r['is']['fees'],b['is']['fees']);self.assertEqual(r['is']['ni'],b['is']['ni'])
  self.assertEqual(r['bs']['otherAssets'][0],10)
  for i in range(1,len(r['bs']['otherAssets'])):
   stock=r['bs']['otherAssets'][i]
   self.assertAlmostEqual(stock,r['is']['fees'][i-1]*25/90,delta=.02)
   self.assertAlmostEqual(r['bs']['cash'][i]-b['bs']['cash'][i],-(stock-a['other_assets']/1000),delta=.02)
   self.assertAlmostEqual(r['bs']['totalAssets'][i],b['bs']['totalAssets'][i],delta=.02)
  pub=run_v2(d);self.assertIn('other_assets_detail',pub);w=calculation_audit_workbook(d,pub);f=io.BytesIO();w.save(f);f.seek(0);loaded=load_workbook(f);self.assertIn('Other Assets',loaded.sheetnames)
  rows=list(loaded['Other Assets'].values);self.assertTrue(any('Accounts receivable' in row for row in rows))
  self.assertTrue(any('otherAssets' in row for row in loaded['Balance Sheet'].values))
  self.assertIn('other_assets_detail.components[Accounts receivable].amount', '\n'.join(str(x) for x in loaded['All Series'].values))
 def test_monthly_fiw_and_capital(self):
  from .fiw import build_fiw,diff_import
  from .callreport import build_rc
  c=json.loads(__import__('pathlib').Path('foundry/fixtures/core_bank_test_base.json').read_text())
  a=c['assumptions'];a['periods_per_year']=12;a['n_periods']=12;a['capital_raises']=[];a['scheduled_borrowings']=[];a['other_assets_mode']='formula_level';a['other_assets_model']=ar_model()
  r=run_v2(c)
  for i,stock in enumerate(r['financials']['bs']['otherAssets'][1:]):
   self.assertAlmostEqual(stock,r['financials']['is']['fees'][i]*25/30,delta=.02)
  # Public trace preserves dimensionless days and factor, and money scales once.
  t=r['other_assets_detail']['components'][0]['terms'][0]
  self.assertEqual(t['days'],[25]*12);self.assertEqual(t['year_days'],360)
  self.assertEqual(t['authored_multiplier'],[1]*12)
  wbbytes=build_fiw(c)[0];wb=load_workbook(io.BytesIO(wbbytes));self.assertIn('ASSM_OTHER_ASSETS',wb.sheetnames)
  imported,report=diff_import(wbbytes,c)
  self.assertEqual(imported["assumptions"]["other_assets_model"],a["other_assets_model"])
  self.assertFalse(any(x["key"].startswith("other_assets_model") for x in report["edits"]))
  # Exact collection formula operands must be present in the editable workbook.
  keys=[row[0] for row in wb['ASSM_OTHER_ASSETS'].values]
  self.assertIn('other_assets_model.components.0.terms.0.days_spec.value',keys)
  self.assertIn('other_assets_model.components.0.terms.0.year_days',keys)
  b=copy.deepcopy(c);b['assumptions']['other_assets_mode']='flat';br=run_v2(b)
  # Asset balances enter the standard 100% other-assets risk bucket.
  parts=r['capital']['standardized']['rwa_components'];bp=br['capital']['standardized']['rwa_components']
  self.assertNotEqual(parts['other_assets'],bp['other_assets'])
 def test_profile_b(self):
  c=json.loads(__import__('pathlib').Path('foundry/fixtures/parity/configs/pf_b_base.json').read_text())
  c['assumptions']['other_assets_mode']='formula_level';c['assumptions']['other_assets_model']=ar_model()
  r=run_parity(c)
  self.assertEqual(len(r['bs']['otherAssets']),len(r['is']['fees']))
  for stock,fee in zip(r['bs']['otherAssets'],r['is']['fees']):self.assertAlmostEqual(stock,fee*25/90,delta=.02)
 def test_flat_legacy_and_inactive(self):
  c=json.loads(__import__('pathlib').Path('foundry/fixtures/core_bank_test_base.json').read_text());b=run_v2(c);d=copy.deepcopy(c);d['assumptions']['other_assets_mode']='flat';d['assumptions']['other_assets_model']=ar_model();r=run_v2(d)
  self.assertEqual(b['financials'],r['financials']);self.assertNotIn('other_assets_detail',r)
if __name__=='__main__':unittest.main()
