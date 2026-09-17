"""Calculation-audit workbook for Foundry v2.

This is intentionally an audit/reconciliation artifact, not a client exhibit.  It exposes
native-cadence calculated series in organized worksheets so a reviewer can line Foundry up
against a source model and identify the first period / subsystem where economics diverge.

The workbook is presentation-only: it consumes the same deterministic run/config contracts
as the UI and never feeds values back into the engine.
"""
from __future__ import annotations

from collections.abc import Mapping
from typing import Any
import json
import calendar
import copy
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .growth import growth_context_from_cfg
from .timebase import model_period_end_date
from .series import apply_amount_basis

_DARK = "182640"
_DARK2 = "24334D"
_GOLD = "D8A85E"
_LIGHT = "EDF0F5"
_MUTED = "8A92A3"
_BLUE = "0000FF"
_GREEN = "008000"
_RED = "FF0000"

_MONEY_FMT = '#,##0.000;[Red](#,##0.000);-'
_NUM_FMT = '#,##0.000000;[Red](#,##0.000000);-'
_COUNT_FMT = '#,##0.000;[Red](#,##0.000);-'
_RATE_FMT = '0.000000;[Red](0.000000);-'
_PCT_FMT = '0.000000%;[Red](0.000000%);-'
_RAW_NUM_FMT = '0.###############;[Red](0.###############);-'
_RAW_MONEY_FMT = '#,##0.###############;[Red](#,##0.###############);-'


def _slug_label(path: str) -> str:
    return path.rsplit(".", 1)[-1].replace("_", " ")


def _exact_base_run(cfg: Mapping[str, Any]) -> Mapping[str, Any]:
    """Re-run the deterministic base engine without the public parity rounding layer.

    The audit workbook is specifically a reconciliation artifact. Using the rounded API seam for
    intermediate calculations can itself manufacture tiny differences, so detailed audit sheets
    consume the exact engine result and convert monetary values to $000s without rounding.
    """
    profile = str(cfg.get("parity_profile") or "")
    if profile == "pf_a":
        from .engine_q_a import run_pf_a
        return run_pf_a(copy.deepcopy(cfg))
    if profile == "pf_b":
        from .engine_q_b import run_pf_b
        return run_pf_b(copy.deepcopy(cfg))
    raise ValueError(f"unknown parity_profile {profile!r} for calculation audit")


def _money_k_series(arr):
    return [None if x is None else float(x) / 1000.0 for x in (arr or [])]


def _exact_financial_rows(fin: Mapping[str, Any], section: str, units: str, n: int,
                          *, include_open: bool = False, ratio: bool = False):
    rows = []
    for key, arr in (fin or {}).items():
        if not isinstance(arr, list) or not all(x is None or isinstance(x, (int, float)) for x in arr):
            continue
        vals = list(arr) if ratio else _money_k_series(arr)
        target = n + (1 if include_open else 0)
        if include_open and len(vals) == n:
            vals = [None] + vals
        elif not include_open and len(vals) == n + 1:
            vals = vals[1:]
        rows.append((section, _slug_label(key), key, units, vals[:target], _RATE_FMT if ratio else _MONEY_FMT))
    return rows


def _period_labels(cfg: Mapping[str, Any], n: int, ppy: int) -> list[str]:
    prefix = "M" if ppy == 12 else ("Y" if ppy == 1 else "Q")
    out = []
    for p in range(1, n + 1):
        try:
            d = model_period_end_date(cfg, p, ppy)
            out.append(f"{prefix}{p}\n{d.isoformat()}")
        except Exception:
            out.append(f"{prefix}{p}")
    return out


def _sheet_title(ws, title: str, subtitle: str | None = None, width: int = 8) -> int:
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(4, width))
    c = ws.cell(1, 1, title)
    c.fill = PatternFill("solid", fgColor=_DARK)
    c.font = Font(color=_LIGHT, bold=True, size=13)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(4, width))
        c2 = ws.cell(2, 1, subtitle)
        c2.font = Font(color=_MUTED, italic=True, size=9)
        c2.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 28
        return 4
    return 3


def _style_header(ws, row: int, ncols: int) -> None:
    thin = Side(style="thin", color="3A4C6E")
    for c in ws[row][:ncols]:
        c.fill = PatternFill("solid", fgColor=_DARK2)
        c.font = Font(color=_LIGHT, bold=True, size=9)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=thin)
    ws.row_dimensions[row].height = 34


def _finish_wide(ws, header_row: int, n_period_cols: int, meta_cols: int = 4) -> None:
    ws.freeze_panes = ws.cell(header_row + 1, meta_cols + 1)
    widths = [18, 38, 34, 18]
    for i, w in enumerate(widths[:meta_cols], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for j in range(meta_cols + 1, meta_cols + 1 + n_period_cols):
        ws.column_dimensions[get_column_letter(j)].width = 15
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(meta_cols+n_period_cols)}{ws.max_row}"


def _write_wide_rows(ws, cfg, rows, *, title: str, subtitle: str, n: int, ppy: int,
                     include_open: bool = False) -> None:
    labels = _period_labels(cfg, n, ppy)
    pcols = (["Open"] if include_open else []) + labels
    hr = _sheet_title(ws, title, subtitle, 4 + len(pcols))
    ws.append([]) if ws.max_row < hr - 1 else None
    ws.cell(hr, 1, "Section")
    ws.cell(hr, 2, "Calculation / Series")
    ws.cell(hr, 3, "Key / Series ID")
    ws.cell(hr, 4, "Units")
    for j, x in enumerate(pcols, 5):
        ws.cell(hr, j, x)
    _style_header(ws, hr, 4 + len(pcols))
    for section, label, key, units, vals, fmt in rows:
        vals = list(vals or [])
        expected = n + (1 if include_open else 0)
        if len(vals) < expected:
            vals += [None] * (expected - len(vals))
        elif len(vals) > expected:
            vals = vals[:expected]
        ws.append([section, label, key, units] + vals)
        r = ws.max_row
        for c in range(5, 5 + expected):
            ws.cell(r, c).number_format = fmt or _NUM_FMT
            ws.cell(r, c).alignment = Alignment(horizontal="right")
        ws.cell(r, 2).alignment = Alignment(wrap_text=True)
        ws.cell(r, 3).font = Font(name="Consolas", size=8, color=_MUTED)
    _finish_wide(ws, hr, len(pcols), 4)


def _write_long_rows(ws, *, title: str, subtitle: str, headers: list[str], rows: list[list[Any]],
                     widths: list[int] | None = None, formats: dict[int, str] | None = None,
                     freeze_col: int = 1) -> None:
    """Write an audit-oriented long-form table. Column numbers in ``formats`` are 1-based."""
    hr = _sheet_title(ws, title, subtitle, len(headers))
    for j, h in enumerate(headers, 1):
        ws.cell(hr, j, h)
    _style_header(ws, hr, len(headers))
    for row in rows:
        ws.append(list(row))
    if formats:
        for col, fmt in formats.items():
            for r in range(hr + 1, ws.max_row + 1):
                ws.cell(r, col).number_format = fmt
                if isinstance(ws.cell(r, col).value, (int, float)):
                    ws.cell(r, col).alignment = Alignment(horizontal="right")
    for r in range(hr + 1, ws.max_row + 1):
        for c in range(1, min(len(headers), 12) + 1):
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
    widths = list(widths or [])
    for j in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(j)].width = widths[j - 1] if j - 1 < len(widths) else 16
    ws.freeze_panes = ws.cell(hr + 1, max(1, int(freeze_col)) + 1)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(headers))}{ws.max_row}"


def _is_sensitive_key(key: str) -> bool:
    k = str(key or "").lower().replace("-", "_")
    return any(tok in k for tok in ("password", "secret", "api_key", "apikey", "access_token",
                                     "refresh_token", "authorization", "credential"))


def _flatten_config_values(obj: Any, path: str = ""):
    """Yield exact authored scalar leaves for audit comparison without dumping secret material."""
    if isinstance(obj, Mapping):
        if not obj and path:
            yield path, "{}", "empty object"
        for k, v in obj.items():
            np = f"{path}.{k}" if path else str(k)
            if _is_sensitive_key(str(k)):
                yield np, "[REDACTED]", "redacted"
            else:
                yield from _flatten_config_values(v, np)
        return
    if isinstance(obj, (list, tuple)):
        if not obj and path:
            yield path, "[]", "empty list"
        for i, v in enumerate(obj):
            yield from _flatten_config_values(v, f"{path}[{i}]")
        return
    if obj is None:
        yield path, "null", "null"
    elif isinstance(obj, bool):
        yield path, obj, "boolean"
    elif isinstance(obj, (int, float)):
        yield path, obj, type(obj).__name__
    else:
        yield path, str(obj), type(obj).__name__


def _config_snapshot_rows(cfg: Mapping[str, Any]) -> list[list[Any]]:
    return [[path, value, typ] for path, value, typ in _flatten_config_values(cfg)]


def _series_provenance_rows(cfg: Mapping[str, Any]) -> list[list[Any]]:
    """Inventory stable Series identifiers and authored links visible in configuration."""
    rows, seen = [], set()

    def add(path, field, sid, owner="", semantic="", source="", trajectory="", cadence="", resolution="",
            link_kind="", link_target="", aggregation=""):
        key = (path, field, str(sid or ""))
        if not sid or key in seen:
            return
        seen.add(key)
        rows.append([path, field, str(sid), owner, semantic, source, trajectory, cadence, resolution,
                     link_kind, link_target, aggregation])

    def walk(obj, path=""):
        if isinstance(obj, Mapping):
            # Canonical Series spec or module-owned object with a stable series_id.
            sid = str(obj.get("series_id") or "").strip()
            link = obj.get("link") if isinstance(obj.get("link"), Mapping) else {}
            if sid:
                add(path, "series_id", sid, str(obj.get("owner_module") or ""),
                    str(obj.get("semantic_type") or ""), str(obj.get("source") or ""),
                    str(obj.get("trajectory") or obj.get("mode") or ""), str(obj.get("cadence") or obj.get("period") or ""),
                    str(obj.get("resolution") or ""), str((link or {}).get("kind") or ""),
                    str((link or {}).get("series_id") or (link or {}).get("name") or ""),
                    str((link or {}).get("aggregation") or ""))
            # Other stable-ID fields owned by a larger object.
            for k, v in obj.items():
                if k != "series_id" and str(k).endswith("_series_id") and str(v or "").strip():
                    add(path, str(k), str(v).strip(), str(obj.get("owner_module") or ""), str(k).removesuffix("_series_id"),
                        str(obj.get("source") or ""), str(obj.get("trajectory") or obj.get("mode") or ""),
                        str(obj.get("cadence") or obj.get("period") or ""), str(obj.get("resolution") or ""))
            if isinstance(obj.get("derived_series_ids"), Mapping):
                for sem, dsid in obj.get("derived_series_ids").items():
                    add(f"{path}.derived_series_ids" if path else "derived_series_ids", str(sem), dsid,
                        str(obj.get("owner_module") or "customer_acquisition"), str(sem), "derived")
            if str(obj.get("source") or "").lower() == "link" and link:
                target = str(link.get("series_id") or link.get("name") or "")
                if target:
                    add(path, "link target", target, "", "linked source", "link", "", "", "",
                        str(link.get("kind") or ""), target, str(link.get("aggregation") or ""))
            for k, v in obj.items():
                if isinstance(v, (Mapping, list, tuple)):
                    walk(v, f"{path}.{k}" if path else str(k))
        elif isinstance(obj, (list, tuple)):
            for i, v in enumerate(obj):
                if isinstance(v, (Mapping, list, tuple)):
                    walk(v, f"{path}[{i}]")
    walk(cfg)
    return rows


def _series_spec_summary(spec: Mapping[str, Any] | None) -> str:
    if not isinstance(spec, Mapping):
        return "legacy scalar / params"
    src = str(spec.get("source") or "entered")
    if src == "link":
        lk = spec.get("link") or {}
        return f"link:{lk.get('kind') or '?'}:{lk.get('series_id') or lk.get('name') or '?'}:{lk.get('aggregation') or 'default'}"
    traj = str(spec.get("trajectory") or spec.get("mode") or "flat")
    if traj == "explicit":
        return f"entered:explicit:{spec.get('cadence') or 'year'}:{spec.get('resolution') or 'step'}:extend={spec.get('extend') or 'hold'}"
    if traj == "growth":
        gs = spec.get("growth_spec") or {}
        return f"entered:growth:{gs.get('period') or 'year'}:{gs.get('method') or 'step'}"
    return "entered:flat"


def _cac_channel_rows(cfg, n, ppy):
    """Annual channel summary aggregated from the canonical monthly CAC calculation."""
    a = cfg.get("assumptions") or {}
    feeds = a.get("cac_feeds") or {}
    if not feeds:
        return []
    from .cac_feeder import cac_auc_rollforward
    gctx = growth_context_from_cfg(cfg, ppy)
    rows = []
    for fname, feed in feeds.items():
        raw = cac_auc_rollforward(feed or {}, n, ppy, assumptions=a, growth_context=gctx)
        channels = list((feed or {}).get("channels") or [])
        for yi, yr in enumerate(raw.get("annual") or [], 1):
            detail = list(yr.get("channels") or [])
            for ci, ch in enumerate(channels):
                d = detail[ci] if ci < len(detail) else {}
                specs = (ch or {}).get("driver_specs") or {}
                metadata = {k: _series_spec_summary(v) for k, v in specs.items()}
                if not metadata:
                    metadata = {"legacy_params": (ch or {}).get("params") or {},
                                "avg_auc_growth": (ch or {}).get("avg_auc_growth")}
                rows.append([
                    fname, str((feed or {}).get("series_id") or ""), yi,
                    str((ch or {}).get("name") or f"Channel {ci+1}"), str((ch or {}).get("method") or ""),
                    d.get("new_customers"), d.get("new_auc"), d.get("spend"), d.get("cac"),
                    json.dumps(metadata, sort_keys=True, separators=(",", ":"), default=str),
                ])
    return rows


def _cac_monthly_acquisition_rows(cfg, n, ppy):
    """Causal monthly channel operands/results at full engine precision."""
    a = cfg.get("assumptions") or {}
    feeds = a.get("cac_feeds") or {}
    if not feeds:
        return []
    from .cac_feeder import cac_auc_rollforward
    gctx = growth_context_from_cfg(cfg, ppy)
    rows = []
    for fname, feed in feeds.items():
        raw = cac_auc_rollforward(feed or {}, n, ppy, assumptions=a, growth_context=gctx)
        channels = list((feed or {}).get("channels") or [])
        for mr in raw.get("monthly") or []:
            detail = list(mr.get("channels") or [])
            for ci, ch in enumerate(channels):
                d = detail[ci] if ci < len(detail) else {}
                op = d.get("operands") or {}
                rows.append([
                    fname, str((feed or {}).get("series_id") or ""), mr.get("month"), mr.get("year"),
                    mr.get("month_in_year"), str((ch or {}).get("name") or f"Channel {ci+1}"),
                    str((ch or {}).get("method") or ""),
                    op.get("pool"), op.get("conversion_rate"), op.get("spend"), op.get("cac"),
                    op.get("ftes"), op.get("per_fte"), op.get("new_customers"),
                    d.get("avg_auc_per_customer"), d.get("new_customers"), d.get("new_auc"),
                    d.get("spend"), d.get("cac"),
                ])
    return rows

def _cac_rollforward_rows(cfg, n, ppy):
    a = cfg.get("assumptions") or {}
    if not (a.get("cac_feeds") or {}):
        return []
    from .cac_feeder import cac_auc_rollforward
    gctx = growth_context_from_cfg(cfg, ppy)
    rows = []
    for fname, feed in (a.get("cac_feeds") or {}).items():
        raw = cac_auc_rollforward(feed or {}, n, ppy, assumptions=a, growth_context=gctx)
        for yr in raw.get("annual") or []:
            rows.append([
                fname, str((feed or {}).get("series_id") or ""), yr.get("year"), "month",
                yr.get("attrition_period"),
                yr.get("beg_auc"), yr.get("new_auc"), yr.get("attrition_rate"), yr.get("auc_lost"), yr.get("end_auc"),
                yr.get("beg_cust"), yr.get("new_cust"), yr.get("cust_lost"), yr.get("end_cust"),
                yr.get("total_spend"), yr.get("blended_cac"),
            ])
    return rows


def _cac_monthly_rows(cfg, n, ppy):
    a = cfg.get("assumptions") or {}
    if not (a.get("cac_feeds") or {}):
        return []
    from .cac_feeder import cac_auc_rollforward
    from .balance_measures import monthly_balance_measure_series
    from .timebase import model_period_year_month
    gctx = growth_context_from_cfg(cfg, ppy)
    y0, m0 = model_period_year_month(cfg, 1, ppy)
    months_per_period = 12 // int(ppy)
    modeled_months = int(n) * months_per_period
    rows = []
    for fname, feed in (a.get("cac_feeds") or {}).items():
        raw = cac_auc_rollforward(feed or {}, n, ppy, assumptions=a, growth_context=gctx)
        auc_eop = list(raw.get("auc_end_by_month") or [])
        cust_eop = list(raw.get("customer_end_by_month") or [])
        beg_auc0 = float((feed or {}).get("beginning_auc") or 0.0)
        beg_cust0 = float((feed or {}).get("beginning_customers") or 0.0)
        auc_avg = monthly_balance_measure_series(beg_auc0, auc_eop, "period_average")
        cust_avg = monthly_balance_measure_series(beg_cust0, cust_eop, "period_average")
        total_months = max(len(auc_eop), len(cust_eop))
        for mi in range(total_months):
            idx = (y0 * 12 + m0 - 1) + mi
            yy, mm = idx // 12, idx % 12 + 1
            dt = date(yy, mm, calendar.monthrange(yy, mm)[1])
            native_period = (mi // months_per_period) + 1 if mi < modeled_months else None
            rows.append([
                fname, str((feed or {}).get("series_id") or ""), mi + 1, dt, (mi // 12) + 1, (mi % 12) + 1,
                native_period, mi < modeled_months,
                beg_auc0 if mi == 0 else auc_eop[mi - 1], auc_eop[mi] if mi < len(auc_eop) else None,
                auc_avg[mi] if mi < len(auc_avg) else None,
                beg_cust0 if mi == 0 else cust_eop[mi - 1], cust_eop[mi] if mi < len(cust_eop) else None,
                cust_avg[mi] if mi < len(cust_avg) else None,
                ((raw.get("monthly") or [])[mi].get("new_cust") if mi < len(raw.get("monthly") or []) else None),
                ((raw.get("monthly") or [])[mi].get("cust_lost") if mi < len(raw.get("monthly") or []) else None),
                ((raw.get("monthly") or [])[mi].get("new_auc") if mi < len(raw.get("monthly") or []) else None),
                ((raw.get("monthly") or [])[mi].get("auc_lost") if mi < len(raw.get("monthly") or []) else None),
                ((raw.get("monthly") or [])[mi].get("attrition_rate") if mi < len(raw.get("monthly") or []) else None),
                ((raw.get("monthly") or [])[mi].get("attrition_event") if mi < len(raw.get("monthly") or []) else None),
                ((raw.get("monthly") or [])[mi].get("total_spend") if mi < len(raw.get("monthly") or []) else None),
            ])
    return rows


def _series_rows(fin: Mapping[str, Any], section: str, units: str, n: int,
                 *, include_open: bool = False, fmt: str = _MONEY_FMT):
    rows = []
    for key, arr in (fin or {}).items():
        if not isinstance(arr, list) or not all(x is None or isinstance(x, (int, float)) for x in arr):
            continue
        vals = list(arr)
        target = n + (1 if include_open else 0)
        if include_open and len(vals) == n:
            vals = [None] + vals
        elif not include_open and len(vals) == n + 1:
            vals = vals[1:]
        rows.append((section, _slug_label(key), key, units, vals[:target], fmt))
    return rows


def _raw_opex_context(cfg: Mapping[str, Any], results: Mapping[str, Any], n: int, ppy: int, exact: Mapping[str, Any] | None = None):
    """Rebuild only observational inputs needed to audit additive Opex components.

    These helpers are the same deterministic primitives used by the engine.  The audit workbook
    never changes model outputs; it resolves component amounts independently for traceability.
    """
    a = (cfg.get("assumptions") or {})
    gctx = growth_context_from_cfg(cfg, ppy)

    cac_monthly, cac_beginning = {}, {}
    if a.get("cac_feeds"):
        from .cac_feeder import cac_auc_rollforward
        for nm, feed in (a.get("cac_feeds") or {}).items():
            sid = str((feed or {}).get("series_id") or nm or "").strip()
            if not sid:
                continue
            rr = cac_auc_rollforward(feed or {}, n, ppy, assumptions=a, growth_context=gctx)
            cac_monthly[sid] = list(rr.get("auc_end_by_month") or [])
            cac_beginning[sid] = float((feed or {}).get("beginning_auc") or 0.0)

    if exact is not None:
        qraw = {sid: [float(x or 0.0) for x in vals] for sid, vals in ((exact.get("fee_stream_quantities") or {}).items())}
    else:
        qpub = (((results.get("fee_stream_quantities") or {}).get("series")) or {})
        qraw = {sid: [float(x or 0.0) * 1000.0 for x in vals] for sid, vals in qpub.items()}
    known_ids = set(qraw)
    try:
        from .opex_extensions import fee_stream_quantity_catalog, fee_stream_balance_quantity_catalog
        known_ids.update(x["series_id"] for x in fee_stream_quantity_catalog(a))
        known_ids.update(x["series_id"] for x in fee_stream_balance_quantity_catalog(a))
    except Exception:
        pass

    pools = {}
    if a.get("cost_pools"):
        from .cost_pools import cost_pool_series_map
        pools = cost_pool_series_map(a, n, ppy, growth_context=gctx)

    if exact is not None:
        ta = [float(x or 0.0) for x in (((exact.get("bs") or {}).get("totalAssets")) or [])]
        is_ = exact.get("is") or {}
        fee_income = [float(x or 0.0) for x in (is_.get("fees") or [0.0] * n)]
        gain_on_sale = [float(x or 0.0) for x in (is_.get("gos") or [0.0] * n)]
        servicing_net = [float(x or 0.0) for x in (is_.get("servNet") or [0.0] * n)]
        wfout = exact.get("workforce") or {}
    else:
        ta = [float(x or 0.0) * 1000.0 for x in (((results.get("financials") or {}).get("bs") or {}).get("totalAssets") or [])]
        is_ = ((results.get("financials") or {}).get("is") or {})
        fee_income = [float(x or 0.0) * 1000.0 for x in (is_.get("fees") or [0.0] * n)]
        gain_on_sale = [float(x or 0.0) * 1000.0 for x in (is_.get("gos") or [0.0] * n)]
        servicing_net = [float(x or 0.0) * 1000.0 for x in (is_.get("servNet") or [0.0] * n)]
        wfout = results.get("workforce") or {}

    # Workforce Count is a level Series, so unlike money it is never rescaled between
    # exact and public audit contexts.  Re-key the row arrays by their persisted stable
    # Series IDs and expose the aggregate count Series when the engine returned it.
    workforce_count_history = {}
    wfcfg = (((a.get("nie_detail") or {}).get("workforce")) or {})
    role_ids = list(wfout.get("series_ids") or [])
    role_counts = list(wfout.get("counts") or [])
    if not role_ids:
        role_ids = [str((r or {}).get("series_id") or "") for r in (wfcfg.get("roles") or [])]
    for wi, sid in enumerate(role_ids):
        sid = str(sid or "").strip()
        if sid and wi < len(role_counts):
            workforce_count_history[sid] = [float(x or 0.0) for x in list(role_counts[wi] or [])[:n]]
    total_sid = str(wfout.get("total_count_series_id") or wfcfg.get("total_count_series_id") or "").strip()
    total_counts = list(wfout.get("total_counts") or [])
    if total_sid:
        if total_counts:
            workforce_count_history[total_sid] = [float(x or 0.0) for x in total_counts[:n]]
        elif workforce_count_history:
            workforce_count_history[total_sid] = [
                sum(float(arr[i] if i < len(arr) else 0.0) for arr in workforce_count_history.values())
                for i in range(n)
            ]

    return {
        "growth_context": gctx,
        "cac_monthly": cac_monthly,
        "cac_beginning": cac_beginning,
        "quantity_history": qraw,
        "known_quantity_ids": sorted(known_ids),
        "cost_pools": pools,
        "total_assets": ta,
        "fee_income": fee_income,
        "gain_on_sale": gain_on_sale,
        "servicing_net": servicing_net,
        "workforce_count_history": workforce_count_history,
    }


def _operating_expense_rows(cfg, results, n, ppy, exact=None):
    is_ = (exact.get("is") or {}) if exact is not None else ((results.get("financials") or {}).get("is") or {})
    rows = []
    for key, label in (
        ("prodOpex", "Product operating expense"),
        ("feeOpex", "Fee Product costs"),
        ("workforceComp", "Workforce compensation"),
        ("otherOpex", "Other operating expense"),
        ("depreciationExpense", "Depreciation expense"),
        ("overhead", "Corporate overhead"),
    ):
        if key in is_:
            vals = _money_k_series(is_[key]) if exact is not None else is_[key]
            rows.append(("Income Statement", label, key, "$000s / engine period", vals, _MONEY_FMT))
    if is_.get("overhead") is not None:
        total_nie_raw = [float((is_.get("prodOpex") or [0]*n)[i] or 0)
                         + float((is_.get("feeOpex") or [0]*n)[i] or 0)
                         + float((is_.get("overhead") or [0]*n)[i] or 0) for i in range(n)]
        total_nie = [v / 1000.0 for v in total_nie_raw] if exact is not None else total_nie_raw
        rows.append(("Income Statement", "Total NIE", "derived:total_nie", "$000s / engine period", total_nie, _MONEY_FMT))

    a = cfg.get("assumptions") or {}
    nd = a.get("nie_detail") or {}
    cats = nd.get("categories") or []
    if not cats:
        return rows

    from .income_modules import nie_category_series
    from .opex_extensions import (recognition_spec_for_category, resolve_recognition,
                                  resolve_settlement, normalize_settlement,
                                  resolve_linked_components, resolve_cost_pool_calculation,
                                  linked_component_amount)
    ctx = _raw_opex_context(cfg, results, n, ppy, exact=exact)
    component_sum = [0.0] * n

    for ci, cat in enumerate(cats):
        cname = str((cat or {}).get("name") or f"Category {ci+1}")
        sid = str((cat or {}).get("series_id") or "")
        econ = nie_category_series(cat, n, ppy, growth_context=ctx["growth_context"])
        recspec = recognition_spec_for_category(cat, ppy)
        recog = resolve_recognition(econ, recspec, ppy, context=ctx["growth_context"])
        econ_k = [float(x or 0.0) / 1000.0 for x in econ]
        recog_k = [float(x or 0.0) / 1000.0 for x in recog]
        rows.append((f"Recurring · {cname}", "Economic trajectory", sid, "$000s / engine period", econ_k, _MONEY_FMT))
        rows.append((f"Recurring · {cname}", "Recognized expense", sid, "$000s / engine period", recog_k, _MONEY_FMT))
        for i, v in enumerate(recog_k): component_sum[i] += v

        sett = resolve_settlement(recog, normalize_settlement((cat or {}).get("settlement"), ppy), ppy,
                                  context=ctx["growth_context"])
        rows.append((f"Recurring · {cname}", "Cash settlement", sid, "$000s / engine period",
                     [float(x or 0.0)/1000.0 for x in sett["cash"]], _MONEY_FMT))
        rows.append((f"Recurring · {cname}", "Prepaid balance", sid, "$000s EOP",
                     [float(x or 0.0)/1000.0 for x in sett["prepaid"]], _MONEY_FMT))
        rows.append((f"Recurring · {cname}", "Accrued balance", sid, "$000s EOP",
                     [float(x or 0.0)/1000.0 for x in sett["accrued"]], _MONEY_FMT))

        cp = resolve_cost_pool_calculation(cat, n, ppy, context=ctx["growth_context"], assumptions=a)
        comps = [cp] if cp is not None else resolve_linked_components(cat, n, ppy,
                                                                      context=ctx["growth_context"], assumptions=a)
        raw_comps = ([((cat or {}).get("calculation") or {})] if cp is not None else list((cat or {}).get("linked_components") or []))
        for j, comp in enumerate(comps):
            raw = raw_comps[j] if j < len(raw_comps) else {}
            drv = str((comp or {}).get("driver") or "component")
            nm = str((raw or {}).get("name") or (comp or {}).get("name") or drv.replace("_", " "))
            cid = str((raw or {}).get("component_id") or (raw or {}).get("series_id") or (comp or {}).get("series_id") or (comp or {}).get("ref") or "")
            vals = []
            for i in range(n):
                qmap = {sid0: (arr[i] if i < len(arr) else 0.0) for sid0, arr in ctx["quantity_history"].items()}
                pmap = {sid0: (arr[i] if i < len(arr) else 0.0) for sid0, arr in ctx["cost_pools"].items()}
                metrics = {
                    "fee_income": ctx["fee_income"][i] if i < len(ctx["fee_income"]) else 0.0,
                    "gain_on_sale": ctx["gain_on_sale"][i] if i < len(ctx["gain_on_sale"]) else 0.0,
                    "servicing_net": ctx["servicing_net"][i] if i < len(ctx["servicing_net"]) else 0.0,
                    "fee_stream_quantities": qmap,
                    "customer_acquisition_auc_monthly": ctx["cac_monthly"],
                    "customer_acquisition_auc_beginning": ctx["cac_beginning"],
                    "fee_stream_quantity_history": ctx["quantity_history"],
                    "fee_stream_quantity_known_ids": ctx["known_quantity_ids"],
                    "bank_total_assets_end_by_period": ctx["total_assets"],
                    "cost_pool": pmap,
                    "workforce_count": {sid0: (arr[i] if i < len(arr) else 0.0)
                                         for sid0, arr in ctx["workforce_count_history"].items()},
                    "periods_per_year": ppy,
                }
                try:
                    amt = linked_component_amount(comp, i, metrics)
                except Exception:
                    amt = None
                vals.append(None if amt is None else float(amt) / 1000.0)
            rows.append((f"Additive · {cname}", nm, cid or drv, "$000s / engine period", vals, _MONEY_FMT))
            for i, v in enumerate(vals):
                if v is not None: component_sum[i] += v

    rows.append(("Reconciliation", "Recurring + additive components", "audit:component_sum",
                 "$000s / engine period", component_sum, _MONEY_FMT))
    other = list(is_.get("otherOpex") or [0.0] * n)
    if exact is not None:
        other = [float(x or 0.0) / 1000.0 for x in other]
    residual = [float(other[i] or 0.0) - float(component_sum[i] or 0.0) for i in range(n)]
    rows.append(("Reconciliation", "Residual inside Other Opex (e.g. simplified assessments / gross-up)",
                 "audit:other_opex_residual", "$000s / engine period", residual, _MONEY_FMT))
    return rows


def _opex_component_detail_rows(cfg, results, n, ppy, exact=None):
    """Long-form component diagnostics with intermediate values used by the Opex engine."""
    a = cfg.get("assumptions") or {}
    cats = (((a.get("nie_detail") or {}).get("categories")) or [])
    if not cats:
        return []
    from .opex_extensions import (resolve_linked_components, resolve_cost_pool_calculation, linked_component_amount,
                                  PIECEWISE_LINKED_DRIVER, COST_POOL_CHARGE_DRIVER, CAC_AUC_DRIVER,
                                  FEE_STREAM_QUANTITY_DRIVER, WORKFORCE_COUNT_DRIVER, _piecewise_event_due, _piecewise_term_value,
                                  _normalize_piecewise_bands, _normalize_observation_lag, _lag_to_engine_periods)
    ctx = _raw_opex_context(cfg, results, n, ppy, exact=exact)
    rows = []

    def metrics_for(i):
        qmap = {sid0: (arr[i] if i < len(arr) else 0.0) for sid0, arr in ctx["quantity_history"].items()}
        pmap = {sid0: (arr[i] if i < len(arr) else 0.0) for sid0, arr in ctx["cost_pools"].items()}
        return {
            "fee_income": ctx["fee_income"][i] if i < len(ctx["fee_income"]) else 0.0,
            "gain_on_sale": ctx["gain_on_sale"][i] if i < len(ctx["gain_on_sale"]) else 0.0,
            "servicing_net": ctx["servicing_net"][i] if i < len(ctx["servicing_net"]) else 0.0,
            "fee_stream_quantities": qmap,
            "customer_acquisition_auc_monthly": ctx["cac_monthly"],
            "customer_acquisition_auc_beginning": ctx["cac_beginning"],
            "fee_stream_quantity_history": ctx["quantity_history"],
            "fee_stream_quantity_known_ids": ctx["known_quantity_ids"],
            "bank_total_assets_end_by_period": ctx["total_assets"],
            "cost_pool": pmap,
            "workforce_count": {sid0: (arr[i] if i < len(arr) else 0.0)
                                 for sid0, arr in ctx["workforce_count_history"].items()},
            "periods_per_year": ppy,
        }

    for ci, cat in enumerate(cats):
        cname = str((cat or {}).get("name") or f"Category {ci+1}")
        cp = resolve_cost_pool_calculation(cat, n, ppy, context=ctx["growth_context"], assumptions=a)
        comps = [cp] if cp is not None else resolve_linked_components(cat, n, ppy,
                                                                       context=ctx["growth_context"], assumptions=a)
        raws = ([((cat or {}).get("calculation") or {})] if cp is not None else list((cat or {}).get("linked_components") or []))
        for j, comp in enumerate(comps):
            raw = raws[j] if j < len(raws) else {}
            drv = str((comp or {}).get("driver") or "component")
            nm = str((raw or {}).get("name") or (comp or {}).get("name") or drv.replace("_", " "))
            cid = str((raw or {}).get("component_id") or (raw or {}).get("series_id") or (comp or {}).get("series_id") or (comp or {}).get("ref") or "")
            for i in range(n):
                metrics = metrics_for(i)
                pend = model_period_end_date(cfg, i + 1, ppy)
                try:
                    amount = linked_component_amount(comp, i, metrics)
                except Exception as exc:
                    rows.append([cname, nm, cid, drv, i + 1, pend, None, None, None, None, None, None, None, None,
                                 None, None, None, None, None, None, None, None, None, None, None, f"ERROR: {exc}"])
                    continue

                if drv == PIECEWISE_LINKED_DRIVER:
                    due = _piecewise_event_due(comp, i, ppy)
                    if not due:
                        rows.append([cname, nm, cid, drv, i + 1, pend, False, None, None, None, None, None, None, None,
                                     None, None, None, None, None, None, None, None, None, None, float(amount or 0.0),
                                     "Self-timed component; no event this period"])
                        continue
                    lagp = _lag_to_engine_periods(_normalize_observation_lag(comp.get("observation_lag")), int(ppy))
                    obs = (i + 1) - lagp
                    obs_date = model_period_end_date(cfg, obs, ppy) if obs > 0 else None
                    tvals = []
                    total = 0.0
                    for ti, term in enumerate(comp.get("terms") or [], 1):
                        tv = _piecewise_term_value(term, obs, metrics)
                        wt = float(term.get("weight") if term.get("weight") is not None else 1.0)
                        weighted = wt * tv
                        total += weighted
                        tvals.append((ti, term, tv, wt, weighted))
                    active = None
                    for b in _normalize_piecewise_bands(comp.get("bands")):
                        lo, hi = float(b["lower_bound"]), b["upper_bound"]
                        if total + 1e-12 < lo:
                            continue
                        if hi is None or total <= float(hi) + 1e-12:
                            active = b; break
                    for ti, term, tv, wt, weighted in tvals:
                        rows.append([
                            cname, nm, cid, drv, i + 1, pend, True, obs, obs_date, ti,
                            str(term.get("source") or ""), str(term.get("series_id") or ""), str(term.get("measure") or ""), wt,
                            tv, weighted, total,
                            (active or {}).get("lower_bound"), (active or {}).get("upper_bound"),
                            (active or {}).get("base_amount"), (active or {}).get("marginal_rate"),
                            None, None, None, float(amount or 0.0), "Literal band base + marginal rate × excess",
                        ])
                    continue

                rate = None
                source_base = None
                recovery = None
                notes = ""
                rates = (comp or {}).get("rates") or []
                if i < len(rates):
                    rate = float(rates[i] or 0.0)
                if drv == WORKFORCE_COUNT_DRIVER:
                    sid = str((comp or {}).get("series_id") or "")
                    source_base = float((metrics.get("workforce_count") or {}).get(sid) or 0.0)
                    amounts = list((comp or {}).get("amount_per_fte") or [])
                    rate = float(amounts[i] if i < len(amounts) else 0.0)
                    notes = "Active Workforce Count × amount per FTE / native engine period"
                elif drv == COST_POOL_CHARGE_DRIVER:
                    ref = str((comp or {}).get("ref") or "")
                    source_base = float((metrics.get("cost_pool") or {}).get(ref) or 0.0)
                    recovery = float((comp or {}).get("recovery_pct") or 0.0)
                    notes = "Cost pool × recovery % × (1 + markup)"
                elif drv == FEE_STREAM_QUANTITY_DRIVER:
                    sid = str((comp or {}).get("series_id") or "")
                    source_base = float((metrics.get("fee_stream_quantities") or {}).get(sid) or 0.0)
                    notes = "Fee-stream quantity × rate"
                elif drv in {"fee_income", "gain_on_sale", "servicing_net", "noninterest_income"}:
                    if drv == "noninterest_income":
                        source_base = float(metrics.get("fee_income") or 0.0) + float(metrics.get("gain_on_sale") or 0.0) + float(metrics.get("servicing_net") or 0.0)
                    else:
                        source_base = float(metrics.get(drv) or 0.0)
                    notes = "Same-period driver × rate"
                elif drv == CAC_AUC_DRIVER:
                    notes = f"Canonical monthly {comp.get('measure') or 'period_end'} AUC accrued using {comp.get('rate_period') or 'year'} rate period; see CAC Monthly Canonical"
                _measure = str((comp or {}).get("measure") or "")
                if drv == WORKFORCE_COUNT_DRIVER:
                    _measure = "FTE / headcount"
                rows.append([cname, nm, cid, drv, i + 1, pend, True, None, None, None, None,
                             str((comp or {}).get("series_id") or (comp or {}).get("ref") or ""),
                             _measure, None, None, None, None, None, None, None, None,
                             rate, recovery, source_base, float(amount or 0.0), notes])
    return rows


def _workforce_rows(cfg, results, n, ppy, exact=None):
    wfout = (exact.get("workforce") or {}) if exact is not None else (results.get("workforce") or {})
    wfcfg = (((cfg.get("assumptions") or {}).get("nie_detail") or {}).get("workforce") or {})
    roles = list(wfcfg.get("roles") or [])
    rows = []
    if not roles and not wfout:
        return rows
    hires = list(wfout.get("resolved_hire_periods") or [])
    counts = list(wfout.get("counts") or [])
    gctx = growth_context_from_cfg(cfg, ppy)
    from .workforce import (workforce_role_compensation_series, workforce_compensation_period,
                            workforce_compensation_basis)
    default_load = float(wfcfg.get("default_payroll_load_rate") or 0.0)
    default_spec = wfcfg.get("default_salary_growth_spec")
    total = [0.0] * n
    for i, role in enumerate(roles):
        name = str((role or {}).get("role") or f"Role {i+1}")
        sid = str((role or {}).get("series_id") or "")
        cnt = list(counts[i]) if i < len(counts) else [0.0] * n
        hire = hires[i] if i < len(hires) else (role or {}).get("hire_period")
        start = int(hire or 1)
        annual = workforce_role_compensation_series(role, n, ppy, growth_context=gctx,
                                                     start_period=start, default_growth_spec=default_spec)
        comp_period = workforce_compensation_period(role)
        comp_basis = workforce_compensation_basis(role)
        factor = {"year": 1.0, "quarter": 4.0, "month": 12.0}[comp_period]
        authored = [float(v or 0.0) / factor for v in annual]
        load = float((role or {}).get("payroll_load_rate") if (role or {}).get("payroll_load_rate") is not None else default_load)
        exp = []
        end_raw = (role or {}).get("end_period")
        end = int(end_raw) if end_raw not in (None, "") else None
        for p in range(n):
            period = p + 1
            active = hire is not None and period >= int(hire) and (end is None or period <= end)
            active_count = float(cnt[p] or 0.0) if p < len(cnt) else 0.0
            annual_now = float(annual[p] or 0.0) if p < len(annual) else 0.0
            e = (apply_amount_basis(annual_now / float(ppy), comp_basis, active_count)
                 * (1.0 + load) / 1000.0) if active else 0.0
            exp.append(e)
            total[p] += e
        rows.append((name, "Active count", sid, "FTE / headcount", cnt, _COUNT_FMT))
        rows.append((name, "Compensation amount basis", sid, "basis",
                     [("Total" if comp_basis == "total" else "Per FTE")] * n, None))
        if comp_basis == "total":
            rows.append((name, "Authored total compensation", sid,
                         f"$ / total population / {comp_period}", authored, _MONEY_FMT))
            rows.append((name, "Annualized total compensation", sid,
                         "$ / total population / year", annual, _MONEY_FMT))
        else:
            rows.append((name, "Authored compensation / FTE", sid,
                         f"$ / FTE / {comp_period}", authored, _MONEY_FMT))
            rows.append((name, "Annualized compensation / FTE", sid, "$ / FTE / year", annual, _MONEY_FMT))
        rows.append((name, f"Payroll load ({load:.4%})", sid, "decimal", [load] * n, _RATE_FMT))
        rows.append((name, f"Payroll expense · resolved hire {hire or '—'}", sid, "$000s / engine period", exp, _MONEY_FMT))
    if wfout.get("comp"):
        eng_comp = _money_k_series(wfout.get("comp")) if exact is not None else wfout.get("comp")
        rows.append(("Reconciliation", "Engine workforce compensation", "workforce.comp", "$000s / engine period",
                     eng_comp, _MONEY_FMT))
        role_comp = wfout.get("role_comp")
        if role_comp is not None:
            vals = _money_k_series(role_comp) if exact is not None else list(role_comp)
            rows.append(("Reconciliation", "Role payroll compensation", "workforce.role_comp", "$000s / engine period",
                         vals, _MONEY_FMT))
        else:
            rows.append(("Reconciliation", "Sum of role payroll expense", "audit:role_comp_sum", "$000s / engine period",
                         total, _MONEY_FMT))
        add_comp = wfout.get("additive_comp")
        if add_comp is not None:
            vals = _money_k_series(add_comp) if exact is not None else list(add_comp)
            rows.append(("Reconciliation", "Additive compensation components", "workforce.additive_comp", "$000s / engine period",
                         vals, _MONEY_FMT))
        for comp in (wfout.get("additive_components") or []):
            vals = list((comp or {}).get("amounts") or [])
            if exact is not None:
                vals = _money_k_series(vals)
            rows.append(("Additive compensation", str((comp or {}).get("name") or "Tiered / banded compensation component"),
                         str((comp or {}).get("component_id") or ""), "$000s / engine period", vals, _MONEY_FMT))
    return rows


def _workforce_component_detail_rows(cfg, results, n, ppy, exact=None):
    """Long-form diagnostics for additive Workforce compensation components."""
    a = cfg.get("assumptions") or {}
    wfcfg = (((a.get("nie_detail") or {}).get("workforce")) or {})
    raw_components = list(wfcfg.get("additive_components") or [])
    if not raw_components:
        return []
    from .workforce import (normalize_workforce_additive_component, workforce_additive_component_amount,
                            WORKFORCE_INCOME_FLOW_SOURCE, _income_flow_term_value)
    from .opex_extensions import (_piecewise_event_due, _piecewise_term_value, _normalize_piecewise_bands,
                                  _normalize_observation_lag, _lag_to_engine_periods)
    ctx = _raw_opex_context(cfg, results, n, ppy, exact=exact)
    source = exact if exact is not None else results
    is_ = source.get("is") or ((source.get("financials") or {}).get("is") or {})

    def _series(key):
        arr = list(is_.get(key) or [])
        if exact is not None and len(arr) == n + 1:
            arr = arr[1:]
        return [float(x or 0.0) * (1.0 if exact is not None else 1000.0) for x in arr[:n]]

    nii, fee, gos, srv = _series("nii"), _series("fees"), _series("gos"), _series("servNet")
    nonint = [(fee[i] if i < len(fee) else 0.0) + (gos[i] if i < len(gos) else 0.0) +
              (srv[i] if i < len(srv) else 0.0) for i in range(n)]
    totalrev = [(nii[i] if i < len(nii) else 0.0) + nonint[i] for i in range(n)]
    flow_hist = {"fee_income": fee, "gain_on_sale": gos, "servicing_net": srv,
                 "noninterest_income": nonint, "net_interest_income": nii,
                 "total_operating_revenue": totalrev}
    rows = []
    for ci, raw in enumerate(raw_components):
        comp = normalize_workforce_additive_component(raw, ppy)
        name = str(comp.get("name") or f"Component {ci+1}")
        cid = str(comp.get("component_id") or "")
        for i in range(n):
            metrics = {
                "periods_per_year": ppy, "income_statement_flow_history": flow_hist,
                "fee_stream_quantities": {sid0: (arr[i] if i < len(arr) else 0.0)
                                          for sid0, arr in ctx["quantity_history"].items()},
                "customer_acquisition_auc_monthly": ctx["cac_monthly"],
                "customer_acquisition_auc_beginning": ctx["cac_beginning"],
                "fee_stream_quantity_history": ctx["quantity_history"],
                "fee_stream_quantity_known_ids": ctx["known_quantity_ids"],
                "bank_total_assets_end_by_period": ctx["total_assets"],
            }
            pend = model_period_end_date(cfg, i + 1, ppy)
            due = _piecewise_event_due(comp, i, ppy)
            try:
                amount = workforce_additive_component_amount(comp, i, metrics)
            except Exception as exc:
                rows.append([name, cid, i + 1, pend, due, None, None, None, None, None, None, None, None, None, None,
                             None, None, None, None, f"ERROR: {exc}"])
                continue
            if not due:
                rows.append([name, cid, i + 1, pend, False, None, None, None, None, None, None, None, None, None, None,
                             None, None, None, float(amount or 0.0), "No event this period"])
                continue
            lagp = _lag_to_engine_periods(_normalize_observation_lag(comp.get("observation_lag")), int(ppy))
            obs = (i + 1) - lagp
            obs_date = model_period_end_date(cfg, obs, ppy) if obs > 0 else None
            tvals, composite = [], 0.0
            for ti, term in enumerate(comp.get("terms") or [], 1):
                if str(term.get("source") or "") == WORKFORCE_INCOME_FLOW_SOURCE:
                    tv = _income_flow_term_value(term, obs, metrics)
                else:
                    tv = _piecewise_term_value(term, obs, metrics)
                wt = float(term.get("weight") if term.get("weight") is not None else 1.0)
                weighted = wt * float(tv or 0.0); composite += weighted
                tvals.append((ti, term, tv, wt, weighted))
            active = None
            for b in _normalize_piecewise_bands(comp.get("bands")):
                lo, hi = float(b["lower_bound"]), b["upper_bound"]
                if composite + 1e-12 < lo:
                    continue
                if hi is None or composite <= float(hi) + 1e-12:
                    active = b; break
            for ti, term, tv, wt, weighted in tvals:
                rows.append([
                    name, cid, i + 1, pend, True, obs, obs_date, ti, str(term.get("source") or ""),
                    str(term.get("metric") or term.get("series_id") or ""), str(term.get("aggregation") or term.get("measure") or ""),
                    wt, tv, weighted, composite, (active or {}).get("lower_bound"), (active or {}).get("upper_bound"),
                    (active or {}).get("base_amount"), (active or {}).get("marginal_rate"), float(amount or 0.0),
                    "Posts to Workforce compensation; income flows are model-year-to-date through the observed event period",
                ])
    return rows


def _cac_rows(cfg, results, n, ppy):
    out = []
    a = cfg.get("assumptions") or {}
    pubs = results.get("customer_acquisition") or {}
    if not pubs:
        return out
    from .cac_feeder import cac_auc_rollforward
    from .balance_measures import monthly_balance_measure_series
    gctx = growth_context_from_cfg(cfg, ppy)
    months_per_period = 12 // ppy
    for nm, pub in pubs.items():
        feed = (a.get("cac_feeds") or {}).get(nm) or {}
        sid = str(feed.get("series_id") or nm)
        raw = cac_auc_rollforward(feed, n, ppy, assumptions=a, growth_context=gctx)
        mend = list(raw.get("auc_end_by_month") or [])
        beg = float(feed.get("beginning_auc") or 0.0)
        mavg = monthly_balance_measure_series(beg, mend, "period_average")
        avg_native = []
        for i in range(n):
            lo, hi = i * months_per_period, min(len(mavg), (i + 1) * months_per_period)
            avg_native.append((sum(float(x or 0.0) for x in mavg[lo:hi]) / max(1, hi - lo)) / 1000.0 if hi > lo else None)
        out.append((nm, "AUC period end", sid, "$000s EOP", pub.get("aucEndByPeriod") or [], _MONEY_FMT))
        out.append((nm, "AUC period average", sid, "$000s average", avg_native, _MONEY_FMT))
        out.append((nm, "Active customers period end", sid, "count", pub.get("customerEndByPeriod") or [], _COUNT_FMT))
        out.append((nm, "Active customers period average", sid, "count", pub.get("customerAverageByPeriod") or [], _COUNT_FMT))
        out.append((nm, "Annual-count measure by period", sid, "count", pub.get("customerAnnualCountByPeriod") or [], _COUNT_FMT))
    return out


def _product_rows(results, n, exact=None):
    rows = []
    source_products = (exact.get("products") or []) if exact is not None else (results.get("products") or [])
    for p in source_products:
        name = str(p.get("name") or "Product")
        fam = str(p.get("family") or "")
        for key, arr in p.items():
            if not isinstance(arr, list) or not all(x is None or isinstance(x, (int, float)) for x in arr):
                continue
            vals = list(arr)
            if len(vals) == n + 1:
                vals = vals[1:]
            if len(vals) != n:
                continue
            if key in ("rateQ", "ftp_rate"):
                units, fmt = ("%" if key == "rateQ" else "decimal annual rate"), _RATE_FMT
            else:
                units, fmt = "$000s / engine value", _MONEY_FMT
                if exact is not None:
                    vals = _money_k_series(vals)
            rows.append((f"{fam} · {name}", _slug_label(key), key, units, vals, fmt))
    return rows


def _fee_cost_rows(cfg, results, n, ppy, exact=None):
    rows = []

    # Reconciliation contract: the headline rows MUST be the same series Product Details renders.
    # r100 gives Product Details a parallel unrounded base-engine diagnostic because the historical
    # public parity seam rounds monetary product arrays to 0.01 $000s before the browser sees them.
    # Prefer that exact diagnostic here; fall back to the legacy public array for older results.
    public_products = results.get("products") or []
    for p in public_products:
        name = str(p.get("name") or "Product")
        detail = p.get("detailExact") or {}
        for key, label in (("fees", "Fee revenue"), ("passCost", "Fee Product cost"), ("opex", "Product operating expense")):
            arr = detail.get(key) if isinstance(detail.get(key), list) else p.get(key)
            if isinstance(arr, list) and len(arr) == n:
                exact_detail = isinstance(detail.get(key), list)
                rows.append((name, label, key,
                             "$000s / engine period · Product Details" + (" · unrounded" if exact_detail else " · legacy public"),
                             list(arr), _RAW_MONEY_FMT if exact_detail else _MONEY_FMT))

    # Preserve a separately labeled deterministic exact-engine row as a second audit trail. For
    # current r100 snapshots it should reconcile to the Product Details diagnostic above; it remains
    # useful for older snapshots that do not carry detailExact.
    if exact is not None:
        for p in (exact.get("products") or []):
            name = str(p.get("name") or "Product")
            for key, label in (("fees", "Fee revenue · exact engine"),
                               ("passCost", "Fee Product cost · exact engine"),
                               ("opex", "Product operating expense · exact engine")):
                arr = p.get(key)
                if isinstance(arr, list) and len(arr) == n:
                    rows.append((name, label, f"{key}:exact", "$000s / engine period · unrounded exact engine",
                                 _money_k_series(arr), _RAW_MONEY_FMT))

    # Authoring-level cost factors are economically important audit evidence.  The direct
    # cost factor/rate path and the dimensionless multiplier are separate layers; expose both
    # so a reviewer can reconcile a product-cost difference without reverse-engineering JSON.
    # A saved r82 factor_path remains economically identical because its implicit multiplier
    # is 1.0.
    from .income_modules import _fee_cost_factor_value
    gctx = growth_context_from_cfg(cfg, ppy)
    for prod in ((cfg.get("assumptions") or {}).get("obs_exposures") or []):
        pname = str((prod or {}).get("name") or "Fee Product")
        for si, st in enumerate((prod or {}).get("fee_streams") or []):
            cost = (st or {}).get("cost") or {}
            kind = str(cost.get("kind") or "none")
            if kind == "none":
                continue
            params = cost.get("params") or {}
            sname = str((st or {}).get("name") or f"Stream {si+1}")
            section = f"{pname} › {sname}"
            sid = str((st or {}).get("quantity_series_id") or "")
            if kind == "per_unit":
                base = float(params.get("cost_per_unit") or 0.0)
                base_units = "$ / unit"
                base_fmt = _NUM_FMT
            else:
                base = float(params.get("pct") or 0.0)
                base_units = ("decimal share of transaction throughput"
                              if kind == "pct_of_throughput_opex"
                              else "decimal share of gross fee revenue")
                base_fmt = _RATE_FMT
            rows.append((section, "Base cost factor", sid or f"cost:{kind}", base_units, [base] * n, base_fmt))

            factor_path = params.get("factor_path")
            mult_path = params.get("multiplier_path")
            ctx = {"growth_context": gctx}
            direct = ([_fee_cost_factor_value(factor_path, q, ppy, ctx, base) for q in range(1, n + 1)]
                      if factor_path is not None else [base] * n)
            if factor_path is not None:
                rows.append((section, "Direct cost rate / factor path", sid or "cost:factor_path",
                             base_units, direct, base_fmt))
            mult = ([_fee_cost_factor_value(mult_path, q, ppy, ctx, 1.0) for q in range(1, n + 1)]
                    if mult_path is not None else [1.0] * n)
            effective = [float(d or 0.0) * float(m or 0.0) for d, m in zip(direct, mult)]
            rows.append((section, "Cost multiplier", sid or "cost:multiplier", "dimensionless", mult, _NUM_FMT))
            rows.append((section, "Effective cost factor", sid or "cost:effective", base_units, effective, base_fmt))
    return rows


def _fee_quantity_unit_kinds(cfg):
    """Infer audit-display units for Fee Stream driver quantities without guessing.

    The engine quantity is not universally monetary: helper transaction streams may carry
    counts or other native units, while balance-derived transaction streams and
    amount-per-source-unit coefficients carry dollars.  Preserve
    source units through ×/% derivations and mark only provably monetary paths as money.
    """
    a = (cfg or {}).get("assumptions") or {}
    out = {}
    for key in ("lending_products", "deposit_products", "obs_exposures"):
        for prod in a.get(key) or []:
            streams = list((prod or {}).get("fee_streams") or [])
            by_name = {str((st or {}).get("name") or ""): i for i, st in enumerate(streams) if (st or {}).get("name")}
            cache = {}
            def resolve(i, stack=None):
                if i in cache:
                    return cache[i]
                stack = set(stack or ())
                if i in stack:
                    return "native"
                stack.add(i)
                st = streams[i] or {}
                basis = str(st.get("basis") or "").lower()
                if basis == "balance":
                    kind = "money"
                elif basis == "account":
                    kind = "count"
                elif basis != "transaction":
                    kind = "native"
                else:
                    drv = st.get("driver") or {}
                    params = drv.get("params") or {}
                    coef = params.get("coefficient") or {}
                    if str(coef.get("kind") or "").lower() == "amount_per_source_unit":
                        kind = "money"
                    else:
                        src = str(drv.get("source") or "constant").lower()
                        if src in {"own_balance", "managed_notional", "bank_aggregate", "cost_pool"}:
                            kind = "money"
                        elif src == "customer_acquisition_count":
                            kind = "count"
                        elif src == "stream_ref":
                            ref = str(drv.get("ref") or "")
                            kind = resolve(by_name[ref], stack) if ref in by_name else "native"
                        else:
                            kind = "native"
                cache[i] = kind
                return kind
            for i, st in enumerate(streams):
                sid = str((st or {}).get("quantity_series_id") or "").strip()
                if sid:
                    out[sid] = resolve(i)
    return out


def _fee_stream_meta(cfg):
    """Return stable-ID metadata for every authored Fee Product stream.

    The audit workbook must never rely on a revenue-oriented display name to imply what a
    numeric row represents.  This registry lets quantity/economics sheets append an explicit
    economic-layer label (throughput, rate, revenue, cost) while keeping the user's product and
    stream names as context only.
    """
    a = (cfg or {}).get("assumptions") or {}
    out = {}
    for fam, key in (("Lending", "lending_products"), ("Deposit", "deposit_products"),
                     ("Fee Product", "obs_exposures")):
        for pi, prod in enumerate(a.get(key) or []):
            pname = str((prod or {}).get("name") or f"{fam} {pi + 1}")
            for si, st in enumerate((prod or {}).get("fee_streams") or []):
                st = st or {}
                sid = str(st.get("quantity_series_id") or "").strip()
                if not sid:
                    continue
                out[sid] = {
                    "series_id": sid,
                    "product": pname,
                    "stream": str(st.get("name") or f"Stream {si + 1}"),
                    "family": fam,
                    "basis": str(st.get("basis") or "").strip().lower(),
                    "rate_behavior": str((st.get("rate") or {}).get("behavior") or "flat").strip().lower(),
                    "cost_kind": str((st.get("cost") or {}).get("kind") or "none").strip().lower(),
                }
    return out


def _quantity_metric_label(basis: str) -> str:
    b = str(basis or "").lower()
    if b == "transaction":
        return "Transaction throughput / driver quantity"
    if b == "balance":
        return "Balance driver quantity"
    if b == "account":
        return "Account-count driver quantity"
    if b == "event":
        return "Event driver quantity"
    return "Driver quantity"


def _quantity_rows(cfg, results, n, exact=None):
    meta = _fee_stream_meta(cfg)
    unit_kinds = _fee_quantity_unit_kinds(cfg)
    rows = []
    if exact is not None:
        qmap = exact.get("fee_stream_quantities") or {}
    else:
        qmap = (((results.get("fee_stream_quantities") or {}).get("series")) or {})
    for sid, vals in qmap.items():
        m = meta.get(sid) or {}
        basis = str(m.get("basis") or "").lower()
        label_base = " › ".join(x for x in (m.get("product"), m.get("stream")) if x) or sid
        label = f"{label_base} › {_quantity_metric_label(basis)}"
        sem = ("native_period_flow" if basis == "transaction" else
               "balance_level" if basis == "balance" else
               "count_level" if basis == "account" else "native observation")
        kind = unit_kinds.get(sid, "native")
        if kind == "money":
            qvals = _money_k_series(vals) if exact is not None else list(vals)
            units, fmt = f"$000s · {sem}", _MONEY_FMT
        else:
            # Public run_v2 historically scales all fee quantity Series by 1,000.  The
            # calculation audit uses the exact engine, but restore native units here too
            # for defensive callers that supply only the public result.
            qvals = list(vals) if exact is not None else [float(v or 0.0) * 1000.0 for v in vals]
            units = ("count" if kind == "count" else "native units") + f" · {sem}"
            fmt = _NUM_FMT
        rows.append((m.get("family") or "Fee stream", label, sid, units, qvals, fmt))
    return rows


def _fee_stream_economics_rows(cfg, exact, n):
    """Expose the causal fee-stream chain at the same native cadence as the engine.

    This is a troubleshooting surface, not a re-calculation.  Values are captured by the
    engine while it evaluates each stream, then rendered here with explicit economic-layer
    labels.  Monetary values are converted from engine dollars to $000s only for display.
    """
    econ = (exact or {}).get("fee_stream_economics") or {}
    meta = _fee_stream_meta(cfg)
    unit_kinds = _fee_quantity_unit_kinds(cfg)
    rows = []

    def _series(rec, key):
        vals = list((rec or {}).get(key) or [])
        if len(vals) < n:
            vals += [None] * (n - len(vals))
        return vals[:n]

    for sid, rec in econ.items():
        m = meta.get(sid) or {}
        product = str(m.get("product") or "Fee Product")
        stream = str(m.get("stream") or sid)
        basis = str(m.get("basis") or (rec or {}).get("basis") or "").lower()
        cost_kind = str(m.get("cost_kind") or (rec or {}).get("cost_kind") or "none").lower()
        section = f"{product} › {stream}"

        qty = _series(rec, "quantity")
        if unit_kinds.get(sid, "native") == "money":
            qvals, qunits, qfmt = _money_k_series(qty), "$000s · driver quantity", _RAW_MONEY_FMT
        elif unit_kinds.get(sid, "native") == "count":
            qvals, qunits, qfmt = qty, "count · driver quantity", _RAW_NUM_FMT
        else:
            qvals, qunits, qfmt = qty, "native units · driver quantity", _RAW_NUM_FMT
        rows.append((section, _quantity_metric_label(basis), sid, qunits, qvals, qfmt))

        pricing = _series(rec, "pricing_factor")
        if any(v is not None for v in pricing):
            if basis in {"transaction", "balance"}:
                plabel = "Fee rate / pricing factor"
                punits, pfmt, pvals = "%" + (" of throughput" if basis == "transaction" else " annual on balance"), _PCT_FMT, pricing
            elif basis == "account":
                plabel = "Effective account fee"
                punits, pfmt, pvals = "$ / account / engine period", _RAW_NUM_FMT, pricing
            else:
                plabel = "Effective fee amount"
                punits, pfmt, pvals = "$000s / engine period", _RAW_MONEY_FMT, _money_k_series(pricing)
            rows.append((section, plabel, f"{sid}:pricing", punits, pvals, pfmt))

        rows.append((section, "Gross fee revenue · before contra-revenue", f"{sid}:gross_revenue",
                     "$000s / engine period", _money_k_series(_series(rec, "gross_fee_revenue")), _RAW_MONEY_FMT))

        contra = _series(rec, "contra_revenue")
        if any(abs(float(v or 0.0)) > 0.0 for v in contra):
            rows.append((section, "Contra-revenue / revenue share", f"{sid}:contra_revenue",
                         "$000s / engine period", _money_k_series(contra), _RAW_MONEY_FMT))

        rows.append((section, "Reported fee income · stream output", f"{sid}:reported_fee_income",
                     "$000s / engine period", _money_k_series(_series(rec, "reported_fee_income")), _RAW_MONEY_FMT))

        direct = _series(rec, "direct_cost_factor")
        if any(v is not None for v in direct):
            if cost_kind == "per_unit":
                cunits, cfmt = "$ / throughput unit", _RAW_NUM_FMT
            elif cost_kind == "pct_of_throughput_opex":
                cunits, cfmt = "% of transaction throughput", _PCT_FMT
            else:
                cunits, cfmt = "% of gross fee revenue", _PCT_FMT
            rows.append((section, "Direct cost rate / factor", f"{sid}:direct_cost_factor",
                         cunits, direct, cfmt))
            rows.append((section, "Cost multiplier", f"{sid}:cost_multiplier", "dimensionless",
                         _series(rec, "cost_multiplier"), _RAW_NUM_FMT))
            rows.append((section, "Effective cost rate / factor", f"{sid}:effective_cost_factor",
                         cunits, _series(rec, "effective_cost_factor"), cfmt))

        opcost = _series(rec, "fee_product_cost")
        if any(abs(float(v or 0.0)) > 0.0 for v in opcost) or cost_kind in {"per_unit", "pct_of_revenue_opex", "pct_of_throughput_opex"}:
            rows.append((section, "Fee Product cost · stream output", f"{sid}:fee_product_cost",
                         "$000s / engine period", _money_k_series(opcost), _RAW_MONEY_FMT))
    return rows


def _managed_securities_rows(cfg, results, n):
    """Expose the managed-securities stock/flow/yield causal chain.

    ``results`` is the exact base-engine output here, so monetary values are raw
    dollars and are converted once to $000s for the workbook. Rates/ratios remain
    decimals and use percentage formatting.
    """
    rows = []
    for pi, p in enumerate((results or {}).get("managed_securities") or []):
        pname = str((p or {}).get("name") or f"Managed portfolio {pi+1}")
        psid = str((p or {}).get("series_id") or f"managed_portfolio_{pi+1}")
        section = pname
        rows.append((section, "Target ratio", f"{psid}:target_ratio", "% of linked target source",
                     list((p or {}).get("target_ratio") or []), _PCT_FMT))
        rows.append((section, "Target securities balance", f"{psid}:target", "$000s · period end",
                     _money_k_series((p or {}).get("target") or []), _RAW_MONEY_FMT))
        for si, sl in enumerate((p or {}).get("sleeves") or []):
            sname = str((sl or {}).get("name") or f"Security {si+1}")
            ssid = str((sl or {}).get("series_id") or f"{psid}:sleeve:{si+1}")
            sec = f"{pname} › {sname} [{(sl or {}).get('classification') or 'AFS'}]"
            rows.extend([
                (sec, "Allocation", f"{ssid}:allocation", "% of managed portfolio",
                 list((sl or {}).get("allocation") or []), _PCT_FMT),
                (sec, "Starting portfolio", f"{ssid}:starting", "$000s · beginning balance",
                 _money_k_series((sl or {}).get("starting") or []), _RAW_MONEY_FMT),
                (sec, "Authored maturity / runoff rate", f"{ssid}:maturity_rate_authored",
                 f"% per {(sl or {}).get('maturity_period') or 'model period'}",
                 list((sl or {}).get("maturity_rate_authored") or (sl or {}).get("maturity_rate") or []), _PCT_FMT),
                (sec, "Effective maturity / runoff rate", f"{ssid}:maturity_rate", "% per engine period",
                 list((sl or {}).get("maturity_rate") or []), _PCT_FMT),
                (sec, "Maturing / runoff", f"{ssid}:maturing", "$000s · engine period",
                 _money_k_series((sl or {}).get("maturing") or []), _RAW_MONEY_FMT),
                (sec, "Net purchases / (sales)", f"{ssid}:net_purchases",
                 "$000s · signed flow (+ purchase / − sale)",
                 _money_k_series((sl or {}).get("net_purchases") or []), _RAW_MONEY_FMT),
                (sec, "Ending portfolio", f"{ssid}:ending", "$000s · period end",
                 _money_k_series((sl or {}).get("ending") or []), _RAW_MONEY_FMT),
                (sec, "Annual portfolio yield", f"{ssid}:yield", "% p.a.",
                 list((sl or {}).get("yield") or []), _PCT_FMT),
                (sec, "Interest income", f"{ssid}:interest_income", "$000s · engine period",
                 _money_k_series((sl or {}).get("interest_income") or []), _RAW_MONEY_FMT),
            ])
    return rows


def _pool_rows(cfg, results, n, ppy):
    a = cfg.get("assumptions") or {}
    if a.get("cost_pools"):
        from .cost_pools import cost_pool_catalog, cost_pool_series_map
        gctx = growth_context_from_cfg(cfg, ppy)
        cmap = cost_pool_series_map(a, n, ppy, growth_context=gctx)
        names = {m["series_id"]: m["name"] for m in cost_pool_catalog(a)}
        return [("Cost pool", names.get(sid) or sid, sid, "$000s / engine period",
                 _money_k_series(vals), _MONEY_FMT) for sid, vals in cmap.items()]
    cp = results.get("cost_pools") or {}
    names = cp.get("names") or {}
    return [("Cost pool", names.get(sid) or sid, sid, cp.get("units") or "$000s / engine period", vals, _MONEY_FMT)
            for sid, vals in (cp.get("series") or {}).items()]


def _flatten_numeric_series(obj: Any, path: str = ""):
    if isinstance(obj, Mapping):
        for k, v in obj.items():
            np = f"{path}.{k}" if path else str(k)
            yield from _flatten_numeric_series(v, np)
        return
    if isinstance(obj, list):
        if obj and all(x is None or isinstance(x, (int, float)) for x in obj):
            yield path, list(obj)
            return
        for i, v in enumerate(obj):
            if isinstance(v, Mapping):
                ident = v.get("name") or v.get("role") or v.get("id") or i
                yield from _flatten_numeric_series(v, f"{path}[{ident}]")


def _all_series_rows(results, n):
    rows = []
    seen = set()
    for path, vals in _flatten_numeric_series(results):
        if path in seen or len(vals) not in (n, n + 1):
            continue
        seen.add(path)
        arr = vals[1:] if len(vals) == n + 1 else vals
        low = path.lower()
        if any(x in low for x in ("count", "customer")):
            units, fmt = "native / count", _COUNT_FMT
        elif any(x in low for x in ("rate", "ratio", "roa", "roe", "nim", "eff", "lev", "pct")):
            units, fmt = "native rate / ratio", _RATE_FMT
        else:
            units, fmt = "public run units", _MONEY_FMT
        section = path.split(".", 1)[0]
        rows.append((section, path, path, units, arr, fmt))
    return rows


def calculation_audit_workbook(cfg: Mapping[str, Any], results: Mapping[str, Any]) -> Workbook:
    """Build a multi-sheet reconciliation workbook from one validated Foundry run."""
    a = cfg.get("assumptions") or {}
    ppy = int(a.get("periods_per_year") or 4)
    n = int(a.get("n_periods") or 12)
    wb = Workbook()
    wb.remove(wb.active)

    idx = wb.create_sheet("Index")
    hr = _sheet_title(idx, "Foundry Calculation Audit Workbook",
                      "Native-cadence reconciliation output. Designed for model comparison and debugging; not a filing exhibit.", 8)
    meta = [
        ("Client", (results.get("engagement_echo") or {}).get("client") or cfg.get("proposed_bank")),
        ("Engagement ID", cfg.get("engagement_id")),
        ("Config hash", results.get("config_hash")),
        ("Run hash", results.get("run_hash")),
        ("Engine", results.get("engine_version")),
        ("Cadence", f"{ppy} periods/year"),
        ("Projection periods", n),
        ("Audit precision", "Detailed sheets use the unrounded deterministic base engine; monetary outputs are converted to $000s without display rounding unless a sheet states raw dollars."),
    ]
    for k, v in meta:
        idx.append([k, v])
    idx.append([])
    idx.append(["Worksheet", "Purpose"])
    sheet_notes = [
        ("Config Snapshot", "Exact authored scalar configuration values, including full numeric precision; sensitive credential-like fields are redacted."),
        ("Series Provenance", "Stable Series IDs, ownership metadata, and authored link targets/aggregation semantics."),
        ("Income Statement", "Every public native-cadence income-statement series."),
        ("Balance Sheet", "Every public balance-sheet series, including opening balances."),
        ("Ratios", "Native-cadence public ratios."),
        ("Operating Expense", "IS Opex decomposition plus recurring categories, additive components, settlement balances, and residual reconciliation."),
        ("Opex Component Detail", "Period-by-period linked/tiered/cost-pool intermediates, including tier observations, active bands, rates, and raw-dollar calculated expense."),
        ("Workforce", "Role-level count, compensation assumptions, additive compensation, and engine total."),
        ("Workforce Component Detail", "Period-by-period additive compensation diagnostics: observed FY metric/balance, weighted terms, active hurdle band, and posted Workforce expense."),
        ("CAC - AUC", "Customer-acquisition AUC EOP / average and customer measures by native period."),
        ("CAC Channels", "Annual channel summary aggregated from the canonical monthly CAC calculation."),
        ("CAC Monthly Acquisition", "Primary causal CAC audit: monthly resolved operands and equation outputs for every acquisition channel."),
        ("CAC Annual Rollforward", "Feed-level annual beginning/new/lost/ending AUC and customers, attrition, spend, and blended CAC."),
        ("CAC Monthly Canonical", "Canonical monthly beginning/EOP/average AUC and customer stocks, even when the engine itself is quarterly."),
        ("Product Calculations", "Every native numeric product series surfaced by the run."),
        ("Managed Securities", "Target-driven securities stock/flow/yield chain: target balance, allocation, runoff, signed net purchases/(sales), ending balance, annual yield, and interest income."),
        ("Fee Product Costs", "Product Details/public-run Fee revenue, Fee Product costs, and product Opex, followed by explicitly labeled exact-engine rows plus cost-factor diagnostics."),
        ("Fee Stream Quantities", "Stable observational driver-quantity Series. Row labels explicitly identify the quantity/throughput layer so a revenue-oriented stream name cannot be mistaken for revenue."),
        ("Fee Stream Economics", "Per-stream causal chain captured by the engine: driver quantity/throughput, pricing factor, gross revenue, contra-revenue, reported stream fee income, and direct operating cost layers."),
        ("Cost Pools", "Resolved non-posting cost-pool Series."),
        ("All Series", "Catch-all inventory of numeric period Series surfaced by the public run."),
    ]
    for r in sheet_notes:
        idx.append(list(r))
    index_header_row = idx.max_row - len(sheet_notes)
    _style_header(idx, index_header_row, 2)
    idx.column_dimensions["A"].width = 28
    idx.column_dimensions["B"].width = 95
    idx.freeze_panes = "A4"

    exact = _exact_base_run(cfg)
    bs = exact.get("bs") or {}
    is_ = exact.get("is") or {}
    rt = exact.get("ratios") or {}

    _write_long_rows(
        wb.create_sheet("Config Snapshot"),
        title="Configuration Snapshot · Audit",
        subtitle="Exact authored scalar leaves from the run configuration. Numeric cells preserve materially more precision than ordinary authoring displays; credential-like fields are redacted.",
        headers=["Config path", "Exact authored value", "Type"],
        rows=_config_snapshot_rows(cfg), widths=[78, 30, 16], formats={2: _RAW_NUM_FMT}, freeze_col=1)
    _write_long_rows(
        wb.create_sheet("Series Provenance"),
        title="Series Provenance / Link Inventory · Audit",
        subtitle="Configuration-declared stable IDs and link metadata. Use this with All Series to distinguish source ownership from downstream consumption.",
        headers=["Config path", "ID field", "Series ID / target", "Owner module", "Semantic type", "Source",
                 "Trajectory / mode", "Cadence / period", "Resolution", "Link kind", "Link target", "Aggregation"],
        rows=_series_provenance_rows(cfg), widths=[55, 22, 34, 28, 26, 14, 20, 18, 16, 28, 34, 16], freeze_col=3)

    _write_wide_rows(wb.create_sheet("Income Statement"), cfg,
                     _exact_financial_rows(is_, "Income Statement", "$000s / engine period", n),
                     title="Income Statement · Calculation Audit", subtitle="Unrounded deterministic base-engine IS series, converted to $000s without display rounding.", n=n, ppy=ppy)
    _write_wide_rows(wb.create_sheet("Balance Sheet"), cfg,
                     _exact_financial_rows(bs, "Balance Sheet", "$000s EOP", n, include_open=True),
                     title="Balance Sheet · Calculation Audit", subtitle="Unrounded base-engine opening and EOP balances, converted to $000s without display rounding.", n=n, ppy=ppy, include_open=True)
    _write_wide_rows(wb.create_sheet("Ratios"), cfg,
                     _exact_financial_rows(rt, "Ratios", "% / engine ratio units", n, ratio=True),
                     title="Ratios · Calculation Audit", subtitle="Unrounded native-cadence base-engine ratios.", n=n, ppy=ppy)

    _write_wide_rows(wb.create_sheet("Operating Expense"), cfg, _operating_expense_rows(cfg, results, n, ppy, exact=exact),
                     title="Operating Expense · Calculation Audit",
                     subtitle="Granular recurring and additive calculations. Residual line isolates Opex mechanics not represented by category/component rows.", n=n, ppy=ppy)
    _write_long_rows(
        wb.create_sheet("Opex Component Detail"),
        title="Operating Expense Component Detail · Audit",
        subtitle="Raw-dollar period diagnostics for additive components. Headcount links expose active FTE and amount/FTE; tiered/banded rows expose observation period, every weighted term, composite driver, active literal band, and calculated expense.",
        headers=["Category", "Component", "Component ID", "Driver", "Model period", "Period end", "Event due",
                 "Observation period", "Observation date", "Term #", "Term source", "Term Series ID", "Measure", "Weight",
                 "Raw term value ($)", "Weighted term value ($)", "Composite driver ($)", "Band lower ($)", "Band upper ($)",
                 "Band base fee ($)", "Marginal rate", "Linked rate / amount factor", "Recovery %", "Source base (native units)",
                 "Calculated expense ($)", "Notes"],
        rows=_opex_component_detail_rows(cfg, results, n, ppy, exact=exact),
        widths=[26, 30, 28, 24, 12, 13, 11, 16, 15, 9, 28, 32, 18, 11, 20, 22, 22, 18, 18, 20, 18, 18, 14, 20, 22, 54],
        formats={5: "0", 6: "yyyy-mm-dd", 8: "0", 9: "yyyy-mm-dd", 10: "0", 14: _RAW_NUM_FMT,
                 15: _RAW_MONEY_FMT, 16: _RAW_MONEY_FMT, 17: _RAW_MONEY_FMT, 18: _RAW_MONEY_FMT, 19: _RAW_MONEY_FMT,
                 20: _RAW_MONEY_FMT, 21: _RAW_NUM_FMT, 22: _RAW_NUM_FMT, 23: _RATE_FMT, 24: _RAW_MONEY_FMT, 25: _RAW_MONEY_FMT},
        freeze_col=4)
    _write_wide_rows(wb.create_sheet("Workforce"), cfg, _workforce_rows(cfg, results, n, ppy, exact=exact),
                     title="Workforce · Calculation Audit",
                     subtitle="Role-level headcount and payroll calculations, including resolved activation timing.", n=n, ppy=ppy)
    _write_long_rows(
        wb.create_sheet("Workforce Component Detail"),
        title="Workforce Additive Compensation Component Detail · Audit",
        subtitle="Raw-dollar diagnostics for tiered/banded compensation. Income-statement terms show the exact model-year flow observed at the event; every weighted term and active literal band is exposed.",
        headers=["Component", "Component ID", "Model period", "Period end", "Event due", "Observation period",
                 "Observation date", "Term #", "Term source", "Metric / Series ID", "Aggregation / measure", "Weight",
                 "Raw term value ($)", "Weighted term value ($)", "Composite driver ($)", "Band lower ($)", "Band upper ($)",
                 "Band base fee ($)", "Marginal rate", "Calculated compensation ($)", "Notes"],
        rows=_workforce_component_detail_rows(cfg, results, n, ppy, exact=exact),
        widths=[34, 30, 12, 13, 11, 16, 15, 9, 26, 34, 24, 11, 20, 22, 22, 18, 18, 20, 18, 24, 70],
        formats={3:"0", 4:"yyyy-mm-dd", 6:"0", 7:"yyyy-mm-dd", 8:"0", 12:_RAW_NUM_FMT,
                 13:_RAW_MONEY_FMT, 14:_RAW_MONEY_FMT, 15:_RAW_MONEY_FMT, 16:_RAW_MONEY_FMT, 17:_RAW_MONEY_FMT,
                 18:_RAW_MONEY_FMT, 19:_RAW_NUM_FMT, 20:_RAW_MONEY_FMT}, freeze_col=4)
    _write_wide_rows(wb.create_sheet("CAC - AUC"), cfg, _cac_rows(cfg, results, n, ppy),
                     title="Customer Acquisition / AUC · Calculation Audit",
                     subtitle="AUC period-end and canonical period-average paths plus customer-count measures.", n=n, ppy=ppy)
    _write_long_rows(
        wb.create_sheet("CAC Channels"),
        title="Customer Acquisition Channel Annual Summary · Audit",
        subtitle="Annual presentation summary aggregated from the canonical monthly acquisition calculation. Use CAC Monthly Acquisition for the causal operands and month-by-month equation results.",
        headers=["Feed", "Feed Series ID", "Model year", "Channel", "Method", "New customers",
                 "New AUC ($)", "Acquisition spend ($)", "Implied CAC ($/customer)", "Driver/source metadata"],
        rows=_cac_channel_rows(cfg, n, ppy),
        widths=[25, 30, 12, 30, 22, 20, 24, 22, 24, 65],
        formats={3: "0", 6: _RAW_NUM_FMT, 7: _RAW_MONEY_FMT, 8: _RAW_MONEY_FMT, 9: _RAW_MONEY_FMT}, freeze_col=5)
    _write_long_rows(
        wb.create_sheet("CAC Monthly Acquisition"),
        title="Customer Acquisition Canonical Monthly Equation Detail · Audit",
        subtitle="Primary causal CAC audit. Every source cadence is resolved to this monthly grid before Pool × Conversion, Spend ÷ CAC, FTE × Productivity, or Explicit Customers is evaluated.",
        headers=["Feed", "Feed Series ID", "Canonical month", "Model year", "Month in year", "Channel", "Method",
                 "Pool flow (customers)", "Conversion rate", "Spend input ($)", "CAC input ($/customer)", "FTEs",
                 "Productivity flow (customers/FTE)", "Explicit new customers", "Average AUC/new customer ($)",
                 "New customers", "New AUC ($)", "Acquisition spend ($)", "Implied CAC ($/customer)"],
        rows=_cac_monthly_acquisition_rows(cfg, n, ppy),
        widths=[24, 30, 16, 12, 14, 28, 22, 22, 18, 20, 22, 14, 30, 22, 28, 20, 24, 22, 24],
        formats={3:"0",4:"0",5:"0",8:_RAW_NUM_FMT,9:_RAW_NUM_FMT,10:_RAW_MONEY_FMT,11:_RAW_MONEY_FMT,
                 12:_RAW_NUM_FMT,13:_RAW_NUM_FMT,14:_RAW_NUM_FMT,15:_RAW_MONEY_FMT,16:_RAW_NUM_FMT,
                 17:_RAW_MONEY_FMT,18:_RAW_MONEY_FMT,19:_RAW_MONEY_FMT}, freeze_col=7)
    _write_long_rows(
        wb.create_sheet("CAC Annual Rollforward"),
        title="Customer Acquisition Annual Rollforward · Audit",
        subtitle="Feed-level annual presentation bridge aggregated from canonical monthly CAC calculations.",
        headers=["Feed", "Feed Series ID", "Model year", "Calculation cadence", "Attrition source period",
                 "Beginning AUC ($)", "New AUC ($)", "Effective annual attrition rate", "AUC lost ($)", "Ending AUC ($)",
                 "Beginning customers", "New customers", "Customers lost", "Ending customers",
                 "Total acquisition spend ($)", "Blended CAC ($/new customer)"],
        rows=_cac_rollforward_rows(cfg, n, ppy),
        widths=[25, 30, 12, 22, 25, 22, 22, 18, 20, 22, 20, 18, 18, 18, 25, 26],
        formats={3: "0", 6: _RAW_MONEY_FMT, 7: _RAW_MONEY_FMT, 8: _RAW_NUM_FMT, 9: _RAW_MONEY_FMT, 10: _RAW_MONEY_FMT,
                 11: _RAW_NUM_FMT, 12: _RAW_NUM_FMT, 13: _RAW_NUM_FMT, 14: _RAW_NUM_FMT, 15: _RAW_MONEY_FMT, 16: _RAW_MONEY_FMT}, freeze_col=5)
    _write_long_rows(
        wb.create_sheet("CAC Monthly Canonical"),
        title="Customer Acquisition Canonical Monthly Stocks · Audit",
        subtitle="Canonical monthly stock-and-flow observations underlying native-cadence AUC/customer measures. Rows beyond a partial terminal model year are retained but explicitly marked outside the projection.",
        headers=["Feed", "Feed Series ID", "Canonical month", "Month end", "Model year", "Month in model year",
                 "Native engine period", "Within projection", "Beginning AUC ($)", "AUC EOP ($)", "Average AUC ($)",
                 "Beginning customers", "Customers EOP", "Average customers", "New customers", "Customers lost",
                 "New AUC ($)", "AUC lost ($)", "Attrition rate on event", "Attrition event", "Acquisition spend ($)"],
        rows=_cac_monthly_rows(cfg, n, ppy),
        widths=[25, 30, 16, 14, 12, 18, 18, 16, 22, 22, 22, 20, 18, 18, 18, 18, 22, 20, 22, 16, 22],
        formats={3: "0", 4: "yyyy-mm-dd", 5: "0", 6: "0", 7: "0", 9: _RAW_MONEY_FMT, 10: _RAW_MONEY_FMT,
                 11: _RAW_MONEY_FMT, 12: _RAW_NUM_FMT, 13: _RAW_NUM_FMT, 14: _RAW_NUM_FMT, 15:_RAW_NUM_FMT,
                 16:_RAW_NUM_FMT,17:_RAW_MONEY_FMT,18:_RAW_MONEY_FMT,19:_RAW_NUM_FMT,21:_RAW_MONEY_FMT}, freeze_col=8)
    _write_wide_rows(wb.create_sheet("Product Calculations"), cfg, _product_rows(results, n, exact=exact),
                     title="Product Calculations · Audit",
                     subtitle="Every native numeric product series surfaced by Foundry.", n=n, ppy=ppy)
    _write_wide_rows(wb.create_sheet("Managed Securities"), cfg, _managed_securities_rows(cfg, exact, n),
                     title="Managed Securities · Calculation Audit",
                     subtitle="Target-driven stock/flow/yield chain. Monetary rows are exact engine dollars converted once to $000s; net purchases/(sales) remain signed; annual yields are periodized by the engine exactly once.", n=n, ppy=ppy)
    _write_wide_rows(wb.create_sheet("Fee Product Costs"), cfg, _fee_cost_rows(cfg, results, n, ppy, exact=exact),
                     title="Fee Product Costs · Audit",
                     subtitle="Headline rows reconcile exactly to Product Details/public-run values; separately labeled exact-engine rows retain unrounded precision for source-model reconciliation.", n=n, ppy=ppy)
    _write_wide_rows(wb.create_sheet("Fee Stream Quantities"), cfg, _quantity_rows(cfg, results, n, exact=exact),
                     title="Fee Stream Quantities · Audit",
                     subtitle="Stable driver-quantity Series available to downstream model components. These are quantities/throughput, NOT fee revenue; the economic layer is stated explicitly in every row label.", n=n, ppy=ppy)
    _write_wide_rows(wb.create_sheet("Fee Stream Economics"), cfg, _fee_stream_economics_rows(cfg, exact, n),
                     title="Fee Stream Economics · Audit",
                     subtitle="Engine-captured causal chain by stream. Monetary rows are exact engine dollars converted to $000s; percentage rows remain decimal rates displayed as percentages. Stream output is before any later bank-level post-processing such as Durbin overage adjustments.", n=n, ppy=ppy)
    _write_wide_rows(wb.create_sheet("Cost Pools"), cfg, _pool_rows(cfg, results, n, ppy),
                     title="Cost Pools · Audit",
                     subtitle="Resolved non-posting pricing/cost-recovery source Series.", n=n, ppy=ppy)
    _write_wide_rows(wb.create_sheet("All Series"), cfg, _all_series_rows(results, n),
                     title="All Public Run Series · Audit",
                     subtitle="Catch-all numeric Series inventory. Use subsystem sheets first; this sheet is the completeness backstop.", n=n, ppy=ppy)

    # Basic workbook-wide visual discipline: readable labels, consistent period widths.
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        for row in ws.iter_rows():
            for cell in row:
                if cell.row > 2 and cell.column <= 4:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return wb
