"""Accounting invariants, explicit legacy adoption and typed-source rejection."""
import copy, io, json
from openpyxl import load_workbook
from pathlib import Path
from .tests_interest_balances import _cfg
from .balance_components import BalanceComponents, legacy_to_components
from .engine_q_a import run_pf_a
from .run_q import run_v2
from .audit_workbook import calculation_audit_workbook


def flat(v,unit='scalar'):
    return {'source':'entered','trajectory':'flat','value':v,'unit':unit}


def component(sid,kind,amount,rate=.12,**kw):
    return {'id':sid,'name':'User-defined '+sid,'treatment':kind,'opening_balance':0,
            'risk_weight':.2 if kind in {'cash_allocation','residual_cash'} else 1,
            'yield_spec':flat(rate,'rate'),'terms':[{'source_spec':flat(amount,'money'),'factors':[]}],**kw}


def configured(comps):
    c=_cfg();a=c['assumptions'];a['n_periods']=36
    del a['interest_balance_model'];a['surplus_allocation_policy']='cash'
    a['balance_components_model']={'enabled':True,'inputs':[],'components':comps}
    return c


def main():
    # Explicit adoption preserves every common financial vector over three years,
    # including both cadence choices, stock bases, linked quantities and delayed income.
    count=0
    for ppy,n in [(12,36),(4,12)]:
        for measure in ['current_end','prior_end','average']:
            c=_cfg();a=c['assumptions'];a['periods_per_year']=ppy;a['n_periods']=n
            if ppy==4:
                ib=copy.deepcopy(a['interest_balance_model']);c=json.loads(Path('foundry/fixtures/core_bank_test_base.json').read_text());a=c['assumptions'];a['interest_balance_model']=ib;a['cash_yield']=0;a['periods_per_year']=4;a['n_periods']=12
            ib=a['interest_balance_model'];ib['operating_cash_interest_basis']=measure
            ib['affiliated_cash_interest_basis']=measure;ib['affiliated_cash_interest_start_period']=2
            old=run_pf_a(copy.deepcopy(c));a['balance_components_model']=legacy_to_components(ib);ib['enabled']=False
            new=run_pf_a(c)
            for section in ['bs','is','ratios']:
                for key,values in old[section].items():
                    if key in new[section] and isinstance(values,list):
                        assert len(values)==len(new[section][key]), (section,key)
                        assert all(x==y or x is not None and y is not None and abs(x-y)<1e-6 for x,y in zip(values,new[section][key])),(ppy,measure,section,key)
            count+=1
    # Count × share × raw dollars/customer gives a stock, not a monetary flow.
    stock=component('admin','off_book',0)
    stock['terms']=[{'source_spec':flat(10,'count'),'factors':[flat(.5),flat(20000,'money_per_unit')]}]
    stock['interest_balance_measure']='average'
    m=BalanceComponents(configured([stock])['assumptions'],{}, {},3,12)
    b={};vals=m.stocks(0,b);rows,_=m.period(0,b,vals,0,0,0)
    assert vals['admin']==100000 and rows[0]['income']==500
    m.commit(rows);rows,_=m.period(1,b,m.stocks(1,b),0,0,0);assert rows[0]['income']==1000
    # Opening and current stocks stay in the correct accounting destinations.
    cash=component('bank-cash','cash_allocation',200000,opening_balance=100000)
    asset=component('other-asset','earning_asset',300000,opening_balance=250000)
    admin=component('admin','off_book',1000000)
    c=configured([cash,asset,admin]);r=run_pf_a(copy.deepcopy(c))
    zero=configured([dict(cash,yield_spec=flat(0)),dict(asset,yield_spec=flat(0)),dict(admin,yield_spec=flat(0))]);z=run_pf_a(zero)
    assert abs(r['is']['cashInt'][0]-15000)<1e-6
    assert r['bs']['earningAssets'][0]==250000 and r['bs']['earningAssets'][1]==300000
    assert r['bs']['offBookBalances'][1]==1000000
    for q in range(37):
        assert abs(r['bs']['totalAssets'][q]-(r['bs']['cash'][q]+r['bs']['sec'][q]+r['bs']['afsBook'][q]+r['bs']['htmBook'][q]+r['bs']['netLoans'][q]+r['fixed_assets']['net'][q]+c['assumptions']['intangibles']+c['assumptions']['other_assets']+r['bs']['prepaidOpex'][q]+r['bs']['msr'][q]+r['bs']['dta'][q]+r['bs']['earningAssets'][q]))<1e-5
        assert r['bs']['cash'][q]>=0
    # Costs and fees post once, into their selected lines, without double counting cash interest.
    paid=component('fee-service','off_book',1200000,income_line='fee',cost_spec=flat(.06,'rate'),cost_line='operating_expense')
    c=configured([paid]);p=run_pf_a(copy.deepcopy(c));zero=configured([dict(paid,yield_spec=flat(0),cost_spec=flat(0))]);z=run_pf_a(zero)
    assert abs(p['is']['fees'][0]-z['is']['fees'][0]-12000)<1e-6
    assert abs(p['is']['prodOpex'][0]-z['is']['prodOpex'][0]-6000)<1e-6
    assert p['is']['cashInt'][0]==z['is']['cashInt'][0]==0
    # Full run and workbook expose exact unit conversion and source provenance.
    public=run_v2(copy.deepcopy(c));assert public['balance_components']['components'][0]['income'][0]==12
    buffer=io.BytesIO();calculation_audit_workbook(c,public).save(buffer);buffer.seek(0)
    workbook=load_workbook(buffer,data_only=True)
    assert 'Balance Components' in workbook.sheetnames and 'All Series' in workbook.sheetnames
    ws=workbook['Balance Components'];data=list(ws.values)
    assert any('fee-service.income' in row and 12 in row for row in data)
    assert any('fee-service.annual_rate' in row and .12 in row for row in data)
    assert any('fee-service.annual_cost_rate' in row and .06 in row for row in data)
    # Custom asset weights affect RWA; administered balances remain outside bank assets.
    weighted=configured([component('asset','earning_asset',300000,risk_weight=1)])
    exempt=copy.deepcopy(weighted);exempt['assumptions']['balance_components_model']['components'][0]['risk_weight']=0
    wr=run_v2(weighted);er=run_v2(exempt)
    assert all(abs(x-y-300)<1e-6 for x,y in zip(wr['capital']['standardized']['rwa'],er['capital']['standardized']['rwa']))
    assert wr['financials']==er['financials']
    # An opening recognized DTA and separate earning asset each enter the opening ledger once.
    opening=configured([component('asset','earning_asset',300000,opening_balance=250000)])
    opening['assumptions']['tax_policy']={'opening_nol':1000000,'opening_nol_dta_net':100000}
    opening['assumptions']['tax_detail']={'enabled':True,'va_mode':'none'}
    op=run_pf_a(opening)
    assert op['bs']['dta'][0]==100000 and op['bs']['earningAssets'][0]==250000
    assert abs(op['bs']['totalAssets'][0]-(op['bs']['equity'][0]+sum(p['bal'][0] for p in op['products'] if p['family']=='deposit')+op['bs']['borrow'][0]+op['bs']['otherLiab'][0]))<1e-5
    # Canonical linked count and missing/dimensional sources are checked by the runtime.
    linked=_cfg();linked['assumptions']['n_periods']=36;ib=linked['assumptions']['interest_balance_model']
    ib['mab_source']={'source':'fee_stream_quantity','series_id':'fee-qty-migrated-universal'}
    lr=run_pf_a(copy.deepcopy(linked));linked['assumptions']['balance_components_model']=legacy_to_components(ib);ib['enabled']=False
    ln=run_pf_a(linked);assert all(abs(x-y)<1e-6 for x,y in zip(lr['is']['ni'],ln['is']['ni']))
    # Fail closed on dimensional mistakes, missing references and current equity circularity.
    failures=[]
    for source in [{'source':'input_ref','input_id':'missing'}, {'source':'bank_series','series_id':'equity','timing':'current_period'},flat(5,'count')]:
        item=component('bad','earning_asset',0);item['terms']=[{'source_spec':source,'factors':[]}];failures.append(configured([item])['assumptions'])
    bad=configured([component('bad','earning_asset',0,yield_spec=flat(1,'money'))])['assumptions'];failures.append(bad)
    bad=configured([])['assumptions'];bad['balance_components_model']['inputs']=[{'id':'a','unit':'money','spec':{'source':'input_ref','input_id':'b'}},{'id':'b','unit':'money','spec':{'source':'input_ref','input_id':'a'}}];failures.append(bad)
    for bad in failures:
        try:BalanceComponents(bad,{}, {},36,12)
        except ValueError:pass
        else:raise AssertionError('Invalid source accepted')
    print(f'PASS: {count} legacy adoption cases, stock/unit math, accounting destinations, exact audit and {len(failures)} rejected bad models')


if __name__=='__main__':main()
