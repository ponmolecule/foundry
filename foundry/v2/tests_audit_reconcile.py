"""Audit/Product Details reconciliation and preview-lifecycle regression tests.

Product Details now owns a dedicated unrounded diagnostic series so presentation precision can
be increased without being capped by the historical 0.01-$000s parity seam. The Calculation
Audit workbook must reconcile to that same diagnostic series while retaining separately-labelled
raw-engine rows as a second audit trail.
"""
import copy
import io
import json
import sys

from openpyxl import load_workbook

sys.path.insert(0, ".")

from foundry.v2.audit_workbook import _fee_cost_rows, calculation_audit_workbook
from foundry.v2.engine_q_a import run_pf_a
from foundry.v2.run_q import run_v2

_P = _F = 0


def ck(name, cond, detail=""):
    global _P, _F
    if cond:
        _P += 1
        print(f"  PASS  {name}" + (f" — {detail}" if detail else ""))
    else:
        _F += 1
        print(f"  FAIL  {name}" + (f" — {detail}" if detail else ""))


def cfg_with_rounding_visible():
    c = json.load(open("foundry/fixtures/universal_template_bank.json", encoding="utf-8"))
    a = c["assumptions"]
    a["obs_exposures"] = [p for p in (a.get("obs_exposures") or []) if not p.get("_fee_product")]
    a["obs_exposures"].append({
        "name": "Audit Reconcile",
        "call_report_line": "obs",
        "_fee_product": True,
        "fee_streams": [{
            "name": "Throughput",
            "basis": "transaction",
            "driver": {"source": "constant", "trajectory": "flat", "params": {"base": 123456.789}},
            "rate": {"params": {"per_unit": 0.01789}},
            "cost": {"kind": "pct_of_revenue_opex", "params": {"pct": 0.333333}},
            "timing": {"start_period": 1},
        }],
    })
    return c




def cfg_with_subcent_costs():
    """Three stream costs whose exact product total is 0.038779351215693 $000s.

    The frozen public parity conversion rounds that product total to 0.04 $000s. Product Details
    must retain the exact diagnostic alongside the legacy public value so a precision toggle can
    expose the difference instead of formatting 0.04 as 0.040.
    """
    c = json.load(open("foundry/fixtures/universal_template_bank.json", encoding="utf-8"))
    a = c["assumptions"]
    a["obs_exposures"] = [p for p in (a.get("obs_exposures") or []) if not p.get("_fee_product")]
    costs_000s = [0.019282552032793, 0.019282552032793, 0.000214247150107]
    streams = []
    for i, cost_k in enumerate(costs_000s, 1):
        streams.append({
            "name": f"Cost stream {i}",
            "basis": "transaction",
            "quantity_series_id": f"fee-qty-subcent-{i}",
            "driver": {"source": "constant", "trajectory": "flat", "params": {"base": 1_000_000.0}},
            "rate": {"behavior": "flat", "params": {"per_unit": 0.0}},
            # cost_k is $000s; convert to dollars and solve factor against $1,000,000 throughput.
            "cost": {"kind": "pct_of_throughput_opex", "params": {"pct": (cost_k * 1000.0) / 1_000_000.0}},
            "timing": {"start_period": 1},
        })
    a["obs_exposures"].append({
        "name": "Subcent Cost Reconcile",
        "call_report_line": "obs",
        "_fee_product": True,
        "fee_streams": streams,
    })
    return c

def main():
    print("AUDIT / PRODUCT DETAILS RECONCILIATION\n")
    cfg = cfg_with_rounding_visible()
    a = cfg["assumptions"]
    n = int(a.get("n_periods") or 12)
    ppy = int(a.get("periods_per_year") or 4)
    public = run_v2(copy.deepcopy(cfg))
    exact = run_pf_a(copy.deepcopy(cfg))
    product = next(p for p in public["products"] if p.get("name") == "Audit Reconcile")

    rows = _fee_cost_rows(cfg, public, n, ppy, exact=exact)
    by_label = {(r[0], r[1]): r for r in rows}
    detail = product.get("detailExact") or {}
    for key, label in (("fees", "Fee revenue"), ("passCost", "Fee Product cost"), ("opex", "Product operating expense")):
        row = by_label.get(("Audit Reconcile", label))
        want = detail.get(key) if isinstance(detail.get(key), list) else product[key]
        ck(f"headline {label} row equals Product Details diagnostic series exactly",
           row is not None and list(row[4]) == list(want))

    exact_fee = by_label.get(("Audit Reconcile", "Fee revenue · exact engine"))
    ck("exact fee row remains separately available and ties to Product Details diagnostic",
       exact_fee is not None
       and abs(exact_fee[4][0] - exact["products"][-1]["fees"][0] / 1000.0) < 1e-12
       and abs(exact_fee[4][0] - detail["fees"][0]) < 1e-12
       and abs(product["fees"][0] - detail["fees"][0]) > 1e-6,
       f"public={product['fees'][0]} detail={detail.get('fees',[None])[0]} exact={exact_fee[4][0] if exact_fee else None}")

    buf = io.BytesIO()
    calculation_audit_workbook(cfg, public).save(buf)
    wb = load_workbook(io.BytesIO(buf.getvalue()), data_only=True)
    ws = wb["Fee Product Costs"]
    found = {}
    for row in ws.iter_rows(values_only=True):
        if len(row) >= 5 and row[0] == "Audit Reconcile" and row[1] in {
            "Fee revenue", "Fee Product cost", "Product operating expense",
            "Fee revenue · exact engine", "Fee Product cost · exact engine", "Product operating expense · exact engine",
        }:
            found[row[1]] = row
    ck("workbook Fee revenue cell reconciles exactly to Product Details",
       found.get("Fee revenue") is not None and found["Fee revenue"][4] == detail["fees"][0])
    ck("workbook Fee Product cost cell reconciles exactly to Product Details",
       found.get("Fee Product cost") is not None and found["Fee Product cost"][4] == detail["passCost"][0])
    ck("workbook Product operating expense cell reconciles exactly to Product Details",
       found.get("Product operating expense") is not None and found["Product operating expense"][4] == detail["opex"][0])
    ck("workbook keeps separately labelled exact-engine rows",
       all(k in found for k in ("Fee revenue · exact engine", "Fee Product cost · exact engine", "Product operating expense · exact engine")))

    # First-principles regression for the user-observed 0.040 discrepancy. The three exact stream
    # costs total 0.038779351215693 $000s; the legacy public parity series rounds that to 0.04.
    # Product Details must receive both, and its diagnostic must equal the exact stream sum.
    sub_cfg = cfg_with_subcent_costs()
    sub_public = run_v2(copy.deepcopy(sub_cfg))
    sub_exact = run_pf_a(copy.deepcopy(sub_cfg))
    sub_prod = next(p for p in sub_public["products"] if p.get("name") == "Subcent Cost Reconcile")
    sub_detail = (sub_prod.get("detailExact") or {}).get("passCost") or []
    sub_econ = sub_exact.get("fee_stream_economics") or {}
    stream_sum_k = sum((sub_econ[f"fee-qty-subcent-{i}"]["fee_product_cost"][0] / 1000.0) for i in (1,2,3))
    expected_sum = 0.038779351215693
    ck("legacy parity seam explains 0.040 while Product Details retains 0.038779351215693",
       sub_prod["passCost"][0] == 0.04
       and len(sub_detail) > 0
       and abs(sub_detail[0] - expected_sum) < 1e-15
       and abs(stream_sum_k - expected_sum) < 1e-15,
       f"public={sub_prod['passCost'][0]} detail={sub_detail[0] if sub_detail else None} streams={stream_sum_k}")

    html = open("web/console_v2.html", encoding="utf-8").read()
    ck("audit export freezes a config snapshot and pins expected hashes",
       "async function _calculationAuditSnapshot()" in html
       and "const cfgText = JSON.stringify(cfg)" in html
       and "expected_config_hash:snap.result.config_hash" in html
       and "expected_run_hash:snap.result.run_hash" in html)
    ck("audit export fails closed when inputs change during snapshot run",
       "JSON.stringify(cfg) !== cfgText" in html and "r.status===409" in html)

    # r95 lifecycle hardening: a failed audit preparation must never poison the shared preview
    # sequence or leave output tabs stuck on the null-result ``Running…`` fallback.
    audit_start = html.index("async function _calculationAuditSnapshot()")
    audit_end = html.index("async function exportCalculationAudit()", audit_start)
    audit_js = html[audit_start:audit_end]
    preview_start = html.index("async function preview()")
    preview_end = html.index("function paintChrome()", preview_start)
    preview_js = html[preview_start:preview_end]
    ck("audit snapshot synchronizes modules before freezing config",
       audit_js.index("syncModules()") < audit_js.index("const cfgText = JSON.stringify(cfg)"))
    ck("failed audit preparation does not advance shared preview sequence",
       "++seq" not in audit_js.split("if(!pr.ok)")[0])
    ck("successful audit snapshot invalidates older preview only after success",
       "++seq;" in audit_js and audit_js.index("++seq;") > audit_js.index("snapRes = await pr.json()"))
    ck("audit preparation has timeout and useful failure detail",
       "AbortController" in audit_js and "Audit snapshot preview failed" in audit_js)
    ck("audit failure schedules a normal preview recovery",
       "if(!snap.result){ refresh();" in html)
    ck("normal preview cannot spin forever on server/network failure",
       "AbortController" in preview_js and "Engine preview timed out after 60 seconds." in preview_js
       and "else if(!r.ok)" in preview_js)
    ck("output tabs surface model-run failures instead of indefinite Running",
       "if(lastRunFailure)" in html and "Model run failed — no stale financials are being shown" in html)

    # Direct endpoint regression: the server must reject a workbook request whose expected
    # Product Details hashes do not match its deterministic rerun of the frozen config.
    import app as app_module
    bad = app_module.v2_calculation_audit({
        "config": copy.deepcopy(cfg),
        "expected_config_hash": "definitely-wrong",
        "expected_run_hash": "also-wrong",
    }, _="test")
    ck("server rejects an audit snapshot hash mismatch", getattr(bad, "status_code", None) == 409)

    print(f"\nRESULT: {_P} passed, {_F} failed")
    if _F:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
