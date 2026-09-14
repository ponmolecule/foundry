"""r96 regression: Fee Stream audit rows must expose the economic layer they contain.

A revenue-oriented stream name must never make a driver-quantity row look like fee revenue.
The audit workbook also exposes the engine-captured quantity -> pricing -> revenue -> cost chain.
"""
import copy
import io
import json
import sys

from openpyxl import load_workbook

sys.path.insert(0, ".")

from foundry.v2.audit_workbook import (
    _fee_stream_economics_rows,
    _quantity_rows,
    calculation_audit_workbook,
)
from foundry.v2.engine_q_a import run_pf_a
from foundry.v2.run_q import run_v2

_P = _F = 0


def ck(name, cond, detail=""):
    global _P, _F
    if cond:
        _P += 1
        print(f"  PASS  {name}" + (f" — {detail}" if detail else ""))
    else:
        _F += 1
        print(f"  FAIL  {name}" + (f" — {detail}" if detail else ""))


def cfg_case():
    c = json.load(open("foundry/fixtures/universal_template_bank.json", encoding="utf-8"))
    a = c["assumptions"]
    a["obs_exposures"] = [p for p in (a.get("obs_exposures") or []) if not p.get("_fee_product")]
    a["obs_exposures"].append({
        "name": "BaaS APIs",
        "call_report_line": "obs",
        "_fee_product": True,
        "fee_streams": [
            {
                "name": "Enabled Partners",
                "basis": "transaction",
                "quantity_series_id": "fee-qty-audit-partners",
                "driver": {"source": "constant", "trajectory": "flat", "params": {"base": 2.0}},
                "rate": {"behavior": "flat", "params": {"per_unit": 0.0}},
                "cost": {"kind": "none", "params": {}},
                "timing": {"start_period": 1},
            },
            {
                "name": "Per-transaction API Revenue",
                "basis": "transaction",
                "quantity_series_id": "fee-qty-audit-api",
                "driver": {
                    "source": "stream_ref", "ref": "Enabled Partners", "trajectory": "derived",
                    "params": {"coefficient": {
                        "kind": "amount_per_source_unit", "value": 500_000.0,
                        "period": "quarter", "trajectory": "flat"
                    }}
                },
                "rate": {"behavior": "flat", "params": {"per_unit": 0.0002}},  # 0.02% of throughput
                "cost": {"kind": "pct_of_revenue_opex", "params": {"pct": 0.05}},  # 5% of gross fee revenue
                "timing": {"start_period": 1},
            },
        ],
    })
    return c


def main():
    print("FEE STREAM ECONOMICS AUDIT\n")
    cfg = cfg_case()
    n = int(cfg["assumptions"].get("n_periods") or 12)
    ppy = int(cfg["assumptions"].get("periods_per_year") or 4)
    exact = run_pf_a(copy.deepcopy(cfg))
    public = run_v2(copy.deepcopy(cfg))
    sid = "fee-qty-audit-api"

    econ = exact.get("fee_stream_economics") or {}
    rec = econ.get(sid) or {}
    ck("engine captures per-stream economics under stable Series ID", sid in econ)
    ck("captured quantity is throughput, not revenue", abs(rec["quantity"][0] - 1_000_000.0) < 1e-9)
    ck("captured pricing factor is 0.02% as decimal", abs(rec["pricing_factor"][0] - 0.0002) < 1e-12)
    ck("captured gross fee revenue equals throughput x fee rate", abs(rec["gross_fee_revenue"][0] - 200.0) < 1e-9)
    ck("captured reported fee income preserves gross revenue for operating-cost mode", abs(rec["reported_fee_income"][0] - 200.0) < 1e-9)
    ck("captured direct cost rate remains distinct from multiplier", abs(rec["direct_cost_factor"][0] - 0.05) < 1e-12 and abs(rec["cost_multiplier"][0] - 1.0) < 1e-12)
    ck("captured Fee Product cost equals gross revenue x cost rate", abs(rec["fee_product_cost"][0] - 10.0) < 1e-9)

    qrows = {r[2]: r for r in _quantity_rows(cfg, public, n, exact=exact)}
    qr = qrows[sid]
    ck("quantity row explicitly labels Transaction throughput / driver quantity",
       qr[1].endswith("Transaction throughput / driver quantity") and "Per-transaction API Revenue" in qr[1], qr[1])
    ck("quantity row converts monetary throughput to $000s without calling it revenue",
       qr[3].startswith("$000s") and abs(qr[4][0] - 1000.0) < 1e-12 and "revenue" not in qr[3].lower(), qr[3])

    erows = _fee_stream_economics_rows(cfg, exact, n)
    by_label = {r[1]: r for r in erows if r[0] == "BaaS APIs › Per-transaction API Revenue"}
    ck("economics sheet exposes throughput, rate, gross revenue, reported income, and cost as separate rows",
       all(k in by_label for k in (
           "Transaction throughput / driver quantity", "Fee rate / pricing factor",
           "Gross fee revenue · before contra-revenue", "Reported fee income · stream output",
           "Direct cost rate / factor", "Cost multiplier", "Effective cost rate / factor",
           "Fee Product cost · stream output")))
    ck("economics rows preserve first-principles values in display units",
       abs(by_label["Transaction throughput / driver quantity"][4][0] - 1000.0) < 1e-12
       and abs(by_label["Fee rate / pricing factor"][4][0] - 0.0002) < 1e-12
       and abs(by_label["Gross fee revenue · before contra-revenue"][4][0] - 0.2) < 1e-12
       and abs(by_label["Fee Product cost · stream output"][4][0] - 0.01) < 1e-12)

    buf = io.BytesIO(); calculation_audit_workbook(cfg, public).save(buf)
    wb = load_workbook(io.BytesIO(buf.getvalue()), data_only=True)
    ck("workbook ships dedicated Fee Stream Economics troubleshooting sheet", "Fee Stream Economics" in wb.sheetnames)
    qws = wb["Fee Stream Quantities"]
    qhit = next((r for r in qws.iter_rows(values_only=True) if len(r) > 4 and r[2] == sid), None)
    ck("workbook quantity sheet itself cannot visually masquerade as revenue",
       qhit is not None and "Transaction throughput / driver quantity" in str(qhit[1]) and abs(qhit[4] - 1000.0) < 1e-12,
       str(qhit[:5] if qhit else None))
    ews = wb["Fee Stream Economics"]
    rows = [r for r in ews.iter_rows(values_only=True) if len(r) > 4 and r[0] == "BaaS APIs › Per-transaction API Revenue"]
    vals = {r[1]: r[4] for r in rows}
    ck("workbook economics sheet gross fee revenue is not throughput",
       abs(vals.get("Gross fee revenue · before contra-revenue", -1) - 0.2) < 1e-12
       and abs(vals.get("Transaction throughput / driver quantity", -1) - 1000.0) < 1e-12)

    print(f"\n{_P} passed, {_F} failed")
    if _F:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
