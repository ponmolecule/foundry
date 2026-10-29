"""r416: reporting manifest — presentation-only grouping of whole products, checked by result and by period."""
import copy
import io
import json
import os
import unittest

from openpyxl import load_workbook

from .bpt_cover import build_bpt_cover
from .reporting import build_view, override_problem
from .run_q import run_v2

_FIX = os.path.join(os.path.dirname(__file__), "..", "fixtures", "universal_template_bank.json")
_BASE = None


def _base():
    global _BASE
    if _BASE is None:
        with open(_FIX, encoding="utf-8") as fh:
            cfg = json.load(fh)
        _BASE = (cfg, run_v2(copy.deepcopy(cfg)))
    return copy.deepcopy(_BASE[0]), copy.deepcopy(_BASE[1])


def _grouped(extra=None):
    cfg, _ = _base()
    cfg["reporting_manifest"] = {"reporting_products": [{
        "id": "rp-oa", "name": "Operating Accounts",
        "members": [{"family": "deposit", "name": "Retail Demand"}, {"family": "deposit", "name": "MMDA / Savings"},
                    {"family": "obs", "name": "Stream and pricing mechanics"}],
        "line_overrides": dict(extra or {})}]}
    return cfg, run_v2(copy.deepcopy(cfg))


def _rows(ws):
    return {r[0]: list(r[1:]) for r in ws.iter_rows(values_only=True) if r and r[0]}


def _strip(res):
    out = copy.deepcopy(res)
    for k in ("reporting", "config_hash", "run_hash"):
        out.pop(k, None)
    out.get("engagement_echo", {}).pop("config_hash", None)
    return out


class ReportingManifestTests(unittest.TestCase):
    def test_engine_never_reads_manifest(self):
        cfg0, res0 = _base()
        cfg1, res1 = _grouped()
        self.assertEqual(json.dumps(_strip(res0), sort_keys=True), json.dumps(_strip(res1), sort_keys=True))

    def test_standalone_by_default(self):
        cfg, res = _base()
        v = res["reporting"]
        self.assertTrue(v["ok"], v.get("problems"))
        self.assertEqual(v["problems"], [])
        self.assertTrue(all(e["kind"] == "standalone" for e in v["entries"]))
        self.assertEqual(len(v["entries"]), len([p for p in res["products"]]))
        ids = [r["id"] for r in v["results"]]
        self.assertEqual(len(ids), len(set(ids)), "result IDs are unique")
        self.assertEqual(v["window"]["periods"], cfg["assumptions"]["n_periods"])
        self.assertEqual(v["window"]["label"], "M1–M36")

    def test_whole_product_results_and_schedule(self):
        cfg, res = _grouped()
        v = res["reporting"]
        self.assertTrue(v["ok"], v["problems"])
        e = next(x for x in v["entries"] if x["id"] == "rp-oa")
        self.assertEqual(len(e["members"]), 3)
        # every result of every member goes to the reporting product, and nowhere else
        mine = [r for r in v["results"] if r["product"] in e["members"]]
        self.assertTrue(mine and all(r["destinations"] == ["rp-oa"] for r in mine))
        # streams: parts sum exactly to the engine's product series
        prod = next(p for p in res["products"] if p["name"] == "Stream and pricing mechanics")
        for measure in ("fees", "passCost"):
            parts = [r for r in mine if r["product"] == "obs/Stream and pricing mechanics" and r["measure"] == measure]
            self.assertTrue(any(r["stream"] for r in parts), measure)
            tot = [sum(r["series"][i] for r in parts) for i in range(len(prod[measure]))]
            self.assertTrue(all(abs(a - b) < 1e-9 for a, b in zip(tot, prod[measure])), measure)
        sch = e["schedule"]
        for ln in sch["lines"]:
            if ln["detail"]:   # detail is inside its subtotal, never added again
                tot = [sum(d["series"][i] for d in ln["detail"]) for i in range(len(ln["series"]))]
                self.assertTrue(all(abs(a - b) < 1e-9 for a, b in zip(tot, ln["series"])), ln["label"])
        flows = [ln for ln in sch["lines"] if not ln["stock"]]
        dc = [sum((1 if ln["section"] == "income" else -1) * ln["series"][i] for ln in flows)
              for i in range(cfg["assumptions"]["n_periods"])]
        self.assertTrue(all(abs(a - b) < 1e-9 for a, b in zip(dc, sch["direct_contribution"])))
        self.assertEqual([u["label"] for u in sch["unresolved"]], ["Earnings on deposits"])
        labels = [ln["label"] for ln in sch["lines"]]
        self.assertIn("Attributable costs", labels)
        self.assertIn("Cost of funds", labels)
        self.assertNotIn("Earnings on deposits", labels, "never inferred from balances")

    def test_overrides_stay_within_class(self):
        cfg, res = _grouped()
        rid = next(r["id"] for r in res["reporting"]["results"]
                   if r["product"] == "deposit/Retail Demand" and r["class"] == "noninterest_expense")
        self.assertIsNone(override_problem("noninterest_expense", "Deposit servicing"))
        self.assertIn("Not allowed", override_problem("noninterest_expense", "Fee income"))
        self.assertIn("Not allowed", override_problem("noninterest_expense", "cost of funds"))
        self.assertIn("computed line", override_problem("noninterest_income", "Direct contribution"))
        cfg["reporting_manifest"]["reporting_products"][0]["line_overrides"] = {rid: "Deposit servicing"}
        v = build_view(cfg, res)
        self.assertTrue(v["ok"], v["problems"])
        r = next(x for x in v["results"] if x["id"] == rid)
        self.assertEqual((r["line"], r["line_source"]), ("Deposit servicing", "override"))
        e = next(x for x in v["entries"] if x["id"] == "rp-oa")
        self.assertIn("Deposit servicing", [ln["label"] for ln in e["schedule"]["lines"]])
        cfg["reporting_manifest"]["reporting_products"][0]["line_overrides"] = {rid: "Fee income"}
        v = build_view(cfg, res)
        self.assertFalse(v["ok"])
        self.assertEqual(next(x for x in v["results"] if x["id"] == rid)["line"], "Attributable costs")
        self.assertIn("invalid_override", [p["code"] for p in v["problems"]])

    def test_overlap_and_stale_members_do_not_reconcile(self):
        cfg, res = _grouped()
        cfg["reporting_manifest"]["reporting_products"].append(
            {"id": "rp-b", "name": "Second", "members": [{"family": "deposit", "name": "Retail Demand"},
                                                         {"family": "lending", "name": "No such loan"}]})
        v = build_view(cfg, res)
        self.assertFalse(v["ok"])
        codes = {p["code"] for p in v["problems"]}
        self.assertTrue({"overlap", "stale_member"} <= codes, codes)
        self.assertTrue(v["reconciliation"]["coverage"]["reported_twice"])
        bad = {m["class"] for m in v["reconciliation"]["measures"] if not m["ok"]}
        self.assertIn("deposits", bad)

    def test_reconciles_by_measure_and_period(self):
        cfg, res = _grouped()
        rec = res["reporting"]["reconciliation"]
        self.assertTrue(rec["ok"])
        self.assertEqual([m["class"] for m in rec["measures"]],
                         ["loans", "deposits", "interest_income", "interest_expense", "noninterest_income",
                          "noninterest_expense", "credit_losses"])
        for m in rec["measures"]:
            self.assertIsNotNone(m["engine_bank"], m["class"])
            want = cfg["assumptions"]["n_periods"] + (1 if m["kind"] == "stock" else 0)
            self.assertEqual(len(m["diff_bank"]), want)
        # tampering with one result in one period is caught for that measure only
        v = copy.deepcopy(res["reporting"])
        r = next(x for x in v["results"] if x["class"] == "noninterest_expense")
        r["series"][5] += 1.0
        from .reporting import reconcile
        rec2 = reconcile(v, res)
        self.assertFalse(rec2["ok"])
        self.assertEqual([m["class"] for m in rec2["measures"] if not m["ok"]], ["noninterest_expense"])

    def test_business_plan_tables_consume_manifest(self):
        cfg0, res0 = _base()
        wb0 = load_workbook(io.BytesIO(build_bpt_cover(cfg0, res0)[0]), data_only=True)
        self.assertNotIn("Reporting Products", wb0.sheetnames)
        cfg, res = _grouped()
        before = copy.deepcopy((cfg, res))
        wb = load_workbook(io.BytesIO(build_bpt_cover(cfg, res)[0]), data_only=True)
        self.assertEqual((cfg, res), before, "export does not change configuration or results")
        self.assertEqual(wb.sheetnames.index("Reporting Products"), wb.sheetnames.index("Product Map & Notes") - 1)
        for sh in ("Non-Interest Income", "Non-Interest Expense"):
            a, b = _rows(wb0[sh]), _rows(wb[sh])
            self.assertEqual(a["Total"], b["Total"], sh)
        nii, nie = _rows(wb["Non-Interest Income"]), _rows(wb["Non-Interest Expense"])
        self.assertIn("Operating Accounts", nii)
        self.assertNotIn("Retail Demand Fees", nii)
        self.assertIn("Cost of Revenue: Operating Accounts", nie)
        self.assertNotIn("Cost of Revenue: Stream and pricing mechanics", nie)
        rp = _rows(wb["Reporting Products"])
        self.assertIn("Operating Accounts — reporting product ($000)", rp)
        self.assertEqual(rp["Earnings on deposits · unresolved"][:3], ["unresolved"] * 3)
        self.assertIn("Direct contribution · excludes unresolved earnings", rp)
        pm = _rows(wb["Product Map & Notes"])
        self.assertEqual(pm["Retail Demand"][6], "Operating Accounts")
        self.assertEqual(pm["Brokered"][6], "Itself")
        texts = [v for row in wb["Reporting Products"].iter_rows(values_only=True) for v in row if isinstance(v, str)]
        self.assertTrue(any("reporting reconciles" in t for t in texts))

    def test_horizon_follows_engagement(self):
        cfg, _ = _grouped()
        cfg["assumptions"]["n_periods"] = 48
        res = run_v2(copy.deepcopy(cfg))
        v = res["reporting"]
        self.assertEqual(v["window"]["label"], "M1–M48")
        e = next(x for x in v["entries"] if x["id"] == "rp-oa")
        self.assertEqual(len(e["schedule"]["direct_contribution"]), 48)
        self.assertTrue(v["ok"], v["problems"])


    def test_console_wiring_and_browser(self):
        import subprocess
        from pathlib import Path
        root = Path(__file__).resolve().parents[2]
        h = (root / "web/console_v2.html").read_text(encoding="utf-8")
        self.assertNotIn("Fee Product cost (\\u2192 overhead)", h)
        self.assertIn("rowBSProduct('Total direct operating costs'", h)
        self.assertIn("if(typeof rptDetailBar==='function')", h)
        r = subprocess.run(["node", str(root / "web/tests_r416_reporting_manifest.cjs")], capture_output=True,
                           text=True, cwd=str(root), timeout=480)
        self.assertEqual(r.returncode, 0, (r.stdout[-1500:] + r.stderr[-1500:]))


if __name__ == "__main__":
    unittest.main(verbosity=2)
