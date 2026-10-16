"""r338 tests: enhanced, engagement-driven Business Plan Tables (foundry/v2/bpt_enhanced.py).

Run: python3 -m foundry.v2.tests_bpt_enhanced

The payments engagement below is synthesized from the public core fixture (no new fixture file):
four asset products, a non-interest-bearing and an interest-bearing program deposit, and five
off-balance-sheet products, one of which has no revenue or cost (a driver-only source).
"""
import copy
import io
import json
import os
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
FIX = os.path.join(HERE, "..", "fixtures")

ASSETS = ["Prefund Facility", "AP Working Capital", "Instant Settlement", "Merchant Working Capital"]
NIB = "Program Deposits - Non-interest-bearing"
IB = "Program Deposit - Interest-bearing"
OBS_REV = ["Merchant acquiring net revenue", "Money Movement Revenue", "Issuing Revenue",
           "Intercompany services fee from HoldCo (markup on the unmigrated book)"]


def payments_engagement():
    base = json.load(open(os.path.join(FIX, "core_bank_test_base.json")))
    c = copy.deepcopy(base)
    a = c["assumptions"]
    L, D = a["lending_products"], a["deposit_products"]
    O = {o["name"]: o for o in a["obs_exposures"]}
    a["lending_products"] = [dict(copy.deepcopy(L[i]), name=n) for i, n in zip([1, 1, 2, 1], ASSETS)]
    a["deposit_products"] = [dict(copy.deepcopy(D[0]), name=NIB, rate_type="fixed", rate_paid_ann=0.0),
                             dict(copy.deepcopy(D[4]), name=IB)]
    a["obs_exposures"] = [
        dict(copy.deepcopy(O["Unused Commercial Commitments"]), name="Activity Sources",
             fee_yield_ann=0.0, opex_pct_ann=0.0, opex_fixed_m=0.0),
        dict(copy.deepcopy(O["Payments"]), name=OBS_REV[0]),
        dict(copy.deepcopy(O["BaaS"]), name=OBS_REV[1]),
        dict(copy.deepcopy(O["Interchange"]), name=OBS_REV[2]),
        dict(copy.deepcopy(O["Service charges"]), name=OBS_REV[3]),
    ]
    c["proposed_bank"] = c["scenario_name"] = "Payments Program Bank"
    return c


def rows(ws):
    """label -> value cells, limited to the table's own columns (header row width)."""
    width = max(max((i + 1 for i, c in enumerate(r) if c is not None), default=0)
                for r in ws.iter_rows(min_row=2, max_row=3, values_only=True))
    return {r[0]: list(r[1:width]) for r in ws.iter_rows(values_only=True) if r and isinstance(r[0], str)}


def main():
    from foundry.v2.bpt_cover import build_bpt_cover, build_bpt_cover_legacy
    from foundry.v2.run_q import run_v2

    passed = failed = 0

    def ck(name, cond, detail=""):
        nonlocal passed, failed
        if cond:
            passed += 1; print(f"PASS {name}")
        else:
            failed += 1; print(f"FAIL {name} — {detail}")

    cfg = payments_engagement()
    res = run_v2(copy.deepcopy(cfg))
    data, fname = build_bpt_cover(cfg, res)
    wb = load_workbook(io.BytesIO(data), data_only=True)
    names = wb.sheetnames

    # 1. sheet structure follows the enhanced template, with one assumption sheet per ASSET product
    head = ["Balance Sheet EOP (Annual)", "Income Statement (Annual)", "Balance Sheet % of Asset",
            "Income Statement % of Asset", "Financial Performance Highlight", "Balance Sheet Quarterly",
            "Income Statement Quarterly", "Interest Rate Assumptions", "Net Interest Income",
            "Interest Income Calculation"]
    tail = ["Forecasted Capital Levels", "Non-Interest Income", "Non-Interest Expense",
            "De Novo Bank Operating Expenses", "Product Map & Notes"]
    asset_sheets = ["Prefund Facility Assumptions", "AP Working Capital Assumptions",
                    "Instant Settlement Assumptions", "Merchant Working Capital Assump"]
    ck("S1 sheet order = template head + per-asset sheets + template tail",
       names == head + asset_sheets + tail, str(names))
    ck("S2 filename tracks saved engagement name", fname == "Payments_Program_Bank_Foundry_Business_Plan_Table.xlsx", fname)
    ck("S3 sheet titles respect Excel's 31-character limit", all(len(n) <= 31 for n in names))

    # 2. product-driven rows come from this engagement's registry, not the template
    fph = rows(wb["Financial Performance Highlight"])
    ck("P1 Loan Mix rows are this engagement's asset products", all(n in fph for n in ASSETS))
    ck("P2 Program Migration rows are the off-BS revenue programs (driver-only source excluded)",
       all(n in fph for n in OBS_REV) and "Activity Sources" not in fph)
    ck("P3 Program Migration cells are blank (not an engine output)",
       all(len(fph[n]) == 3 and all(v is None for v in fph[n]) for n in OBS_REV))
    mix = [fph[n] for n in ASSETS]
    ck("P4 Loan Mix sums to 100% each year", all(abs(sum(col) - 1) < 1e-6 for col in zip(*mix)))
    dm = [fph["Non-Interest-Bearing Deposits"], fph["Interest-Bearing Deposits"]]
    ck("P5 Deposit Mix sums to 100% and classifies the NIB product as non-interest-bearing",
       all(abs(sum(col) - 1) < 1e-6 for col in zip(*dm)) and dm[0][0] > 0.5)
    nii = rows(wb["Non-Interest Income"])
    ck("P6 Non-Interest Income has one row per off-BS revenue product", all(n in nii for n in OBS_REV))
    ck("P7 driver-only off-BS product is not given an income row", "Activity Sources" not in nii)
    pm = rows(wb["Product Map & Notes"])
    ck("P8 product map lists every registry product with its class",
       all(pm[n][0] == "Asset" for n in ASSETS) and pm[NIB][0] == "Liability" and pm[IB][0] == "Liability"
       and all(pm[n][0] == "Off-BS" for n in OBS_REV + ["Activity Sources"]))

    # 3. ties to the engine
    is_ = res["financials"]["is"]
    y = lambda k, i: sum(x or 0 for x in is_[k][4 * i:4 * i + 4])
    isa = rows(wb["Income Statement (Annual)"])
    ck("T1 annual net income = engine quarters summed",
       all(abs(isa["Net Income"][i] - y("ni", i)) < 0.01 for i in range(3)))
    tot_ii = res["presentation"]["derived"]["totalIntInc"]
    ck("T2 interest income uses the engine total (designated-book interest not double counted)",
       abs(isa["Interest Income"][0] - sum(tot_ii[:4])) < 0.01 and
       abs(isa["Interest Income"][0] - (y("loanInt", 0) + y("secInt", 0) + y("cashInt", 0))) < 0.01)
    ck("T3 Non-Interest Income total ties to the income statement",
       all(abs(nii["Total"][i] - isa["Non-Interest Income"][i]) < 0.01 for i in range(3)))
    nie = rows(wb["Non-Interest Expense"])
    ck("T4 Non-Interest Expense total ties to the income statement",
       all(abs(nie["Total"][i] - isa["Non-Interest Expense"][i]) < 0.01 for i in range(3)))
    dn = rows(wb["De Novo Bank Operating Expenses"])
    ck("T5 De Novo total ties to the Non-Interest Expense line",
       all(abs(dn["Total"][i] - nie["De Novo Bank Operating Expense"][i]) < 0.01 for i in range(3)))
    for nm in ("Balance Sheet EOP (Annual)", "Balance Sheet Quarterly"):
        b = rows(wb[nm])
        parts = ["Cash and Balances Due from Banks", "Securities", "Loans, Net of Allowance for Credit Losses",
                 "Other Assets", "Intangible Assets"]
        ok = all(abs(sum(b[k][j] or 0 for k in parts) - b["Total Assets"][j]) < 0.02 and
                 abs(b["Total Liabilities"][j] + b["Equity"][j] - b["Total Assets"][j]) < 0.05
                 for j in range(len(b["Total Assets"])) if b["Total Assets"][j] is not None)
        ck(f"T6 {nm}: assets foot and liabilities + equity = total assets", ok)
    bq = rows(wb["Balance Sheet Quarterly"])
    ck("T7 quarterly deposit split sums to engine deposits",
       all(abs(a + b - d) < 0.02 for a, b, d in zip(bq["Non-Interest-Bearing Deposits"],
                                                     bq["Interest-Bearing Deposits"],
                                                     res["financials"]["bs"]["deposits"])))
    isp = rows(wb["Income Statement % of Asset"])
    ck("T8 net income % of average assets equals ROAA",
       all(abs(isp["Net Income (Loss)"][i] - fph["ROAA"][i]) < 1e-9 for i in range(3)))
    ra = rows(wb["Interest Rate Assumptions"])
    ck("T9 NIB deposits cost 0%; IB label carries the configured index pricing",
       all(v == 0 for v in ra["Non-Interest-Bearing Deposits"]) and
       any(k.startswith("Interest-Bearing Deposits (SOFR") for k in ra))
    ck("T10 Fed funds row is populated from the engagement's curves",
       all(isinstance(v, float) and 0 < v < 0.2 for v in ra["Federal Funds Rate (EFFR)"]))

    # 4. a different engagement produces different rows from the same code
    uni = json.load(open(os.path.join(FIX, "universal_template_bank.json")))
    ures = run_v2(copy.deepcopy(uni))
    uwb = load_workbook(io.BytesIO(build_bpt_cover(uni, ures)[0]), data_only=True)
    ulend = [p["name"] for p in ures["products"] if p["family"] == "lending"]
    ck("E1 monthly universal engagement: one assumption sheet per own asset product, none of the payments ones",
       (uwb.sheetnames.index("Forecasted Capital Levels") - uwb.sheetnames.index("Interest Income Calculation") - 1)
       == len(ulend)
       and "Prefund Facility Assumptions" not in uwb.sheetnames, str(uwb.sheetnames))
    uq = rows(uwb["Income Statement Quarterly"])
    uni_ni = ures["financials"]["is"]["ni"]
    ck("E2 monthly engine aggregates to calendar quarters (Q1 = M1..M3)",
       abs(uq["Net Income"][0] - sum(uni_ni[0:3])) < 0.01 and len(uq["Net Income"]) == 12)

    # 6. r339 reporting fixes
    from foundry.v2 import tests_balance_components as BC
    from foundry.v2.bpt_enhanced import _Model
    bc = BC.configured([BC.component("other-asset", "earning_asset", 30000000, rate=.06, opening_balance=25000000)])
    bres = run_v2(copy.deepcopy(bc))
    bm = _Model(bc, bres)
    bbs = bres["financials"]["bs"]
    core = [(bbs["cash"][i] or 0) + sum((bbs.get(k) or [0] * 37)[i] or 0 for k in ("htm", "afs", "sec", "htmBook", "afsBook"))
            + (bbs["grossLoans"][i] or 0) for i in range(len(bbs["totalAssets"]))]
    ea = bm.earning_assets_series()
    ck("R1 earning-asset base includes other interest-bearing assets (earningAssets)",
       all(abs(ea[i] - core[i] - bbs["earningAssets"][i]) < 1e-6 for i in range(len(core))) and bbs["earningAssets"][12] > 0)
    dup = copy.deepcopy(bres); dup["financials"]["bs"]["frbStock"] = list(bbs["earningAssets"])
    ck("R2 earningAssets replaces (never adds to) Federal Reserve stock",
       _Model(bc, dup).earning_assets_series() == ea)
    bwb = load_workbook(io.BytesIO(build_bpt_cover(bc, bres)[0]), data_only=True)
    bf = rows(bwb["Financial Performance Highlight"]); bi = rows(bwb["Interest Income Calculation"])
    ann_nii = sum(bres["financials"]["is"]["nii"][:12])
    avg_ea = sum(((ea[i] or 0) + (ea[i + 1] or 0)) / 2 for i in range(12)) / 12
    ck("R3 NIM and yield use the full earning-asset base",
       abs(bf["Net Interest Margin"][0] - ann_nii / avg_ea) < 1e-6 and abs(bi["Total Average Earning Assets"][0] - avg_ea) < 1e-6)
    eb = res["financials"]["is"]["ebtda"]
    lbl = "Earnings before D&A and taxes (bank EBITDA)"
    eq = rows(wb["Income Statement Quarterly"])
    ck("R4 bank EBITDA restored on annual and quarterly income statements and ties to the engine",
       lbl in isa and all(abs(isa[lbl][i] - sum(eb[4 * i:4 * i + 4])) < 0.01 for i in range(3)) and
       lbl in eq and all(abs(eq[lbl][q] - eb[q]) < 0.01 for q in range(12)) and lbl in rows(wb["Income Statement % of Asset"]))
    off = copy.deepcopy(res)
    tgt = next(p for p in off["products"] if p["name"] == "Activity Sources")
    tgt["fees"] = [0.0] * len(tgt["fees"]); tgt["fees"][0] = 1000.0; tgt["fees"][4] = -1000.0
    owb = load_workbook(io.BytesIO(build_bpt_cover(cfg, off)[0]), data_only=True)
    on = rows(owb["Non-Interest Income"]); of = rows(owb["Financial Performance Highlight"])
    ck("R5 offsetting activity (+1,000 then -1,000) keeps the product row",
       "Activity Sources" in on and on["Activity Sources"][:2] == [1000.0, -1000.0] and "Activity Sources" in of)
    ck("R6 a product with no activity in any period is still excluded",
       "Activity Sources" not in nii and not _Model(cfg, res).has_window_activity([0.0] * 12))

    # 5. legacy layout remains reachable and unchanged in shape
    lwb = load_workbook(io.BytesIO(build_bpt_cover_legacy(cfg, res)[0]), read_only=True)
    ck("L1 legacy two-sheet layout still builds",
       lwb.sheetnames == ["Business Plan Tables Annual", "Business Plan Tables Quarterly"])

    print(f"{passed} passed, {failed} failed")
    return failed


if __name__ == "__main__":
    sys.exit(1 if main() else 0)
