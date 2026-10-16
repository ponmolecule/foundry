"""Enhanced Business Plan Tables — engagement-driven, multi-sheet workbook.

Layout follows the enhanced Business Plan Table template (one table per sheet):

    Balance Sheet EOP (Annual), Income Statement (Annual), Balance Sheet % of Asset,
    Income Statement % of Asset, Financial Performance Highlight, Balance Sheet Quarterly,
    Income Statement Quarterly, Interest Rate Assumptions, Net Interest Income,
    Interest Income Calculation, <one assumption sheet per ASSET product>,
    Forecasted Capital Levels, Non-Interest Income, Non-Interest Expense,
    De Novo Bank Operating Expenses, Product Map & Notes.

Engagement-specific rows come from ONE place: the engine's product registry
(``res["products"]``), the same registry the console shows on Detail -> Per-Product
Contributions with its ASSET / LIABILITY / OFF-BS chips.  The straight line is:

    ASSET (lending)       -> Loan Mix rows; one "<Product> Assumptions" sheet each
    LIABILITY (deposit)   -> Deposit Mix + quarterly deposit split, classified as
                             non-interest-bearing / interest-bearing; rate rows
    OFF-BS (obs)          -> Non-Interest Income rows (products with revenue),
                             "Cost of Revenue: <Product>" rows (products with cost),
                             Program Migration rows

Operating-expense categories, other-income lines and fee rows are likewise read from the
engagement's own configuration / engine output, so a new engagement produces its own rows
without code changes.  No template product name is hard-coded.

Sourcing rules (unchanged from the earlier exhibit): every number is engine output from a
fresh ``run_v2`` of the posted configuration, in $ thousands.  Lines the engine does not
produce are left blank (never zero-filled) and are listed in the notes.  Residual lines
("other") exist only so a table ties exactly to its engine total, and are labelled as such.

Window: the regulator-facing submission window (``timebase.submission_quarters``, default
12 calendar quarters), independent of the computational horizon and cadence.
"""

import copy
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .bpt_cover import (_annual_periods, _bank_name, _engagement_slug, _norm,
                        _quarterly_periods)

_ARIAL = "Arial"
_HDR_FILL = PatternFill("solid", fgColor="2C2C2C")
_SEC_FILL = PatternFill("solid", fgColor="E6E5E1")
_HDR_FONT = Font(name=_ARIAL, bold=True, color="FFFFFF", size=10)
_SEC_FONT = Font(name=_ARIAL, bold=True, color="2C2C2C", size=10)
_LBL_FONT = Font(name=_ARIAL, size=10)
_TOTAL_FONT = Font(name=_ARIAL, bold=True, size=10)
_TITLE_FONT = Font(name=_ARIAL, bold=True, color="2C2C2C", size=12)
_NOTE_FONT = Font(name=_ARIAL, italic=True, size=8, color="6A6A66")
_MONEY = '#,##0;(#,##0);"-"'
_PCT = '0.0%;(0.0%);"-"'
_RATE = '0.00%;(0.00%);"-"'
_YEARS = '#,##0.00;(#,##0.00);"-"'

SHEETS_FIXED_HEAD = [
    "Balance Sheet EOP (Annual)", "Income Statement (Annual)", "Balance Sheet % of Asset",
    "Income Statement % of Asset", "Financial Performance Highlight", "Balance Sheet Quarterly",
    "Income Statement Quarterly", "Interest Rate Assumptions", "Net Interest Income",
    "Interest Income Calculation",
]
SHEETS_FIXED_TAIL = ["Forecasted Capital Levels", "Non-Interest Income", "Non-Interest Expense",
                     "De Novo Bank Operating Expenses", "Product Map & Notes"]

_SEC_KEYS = ("htm", "afs", "sec", "htmBook", "afsBook")   # securities lines (bs_layout)
_EPS = 0.005                                              # $5 in $000s: "has activity" tolerance
_RESID = 0.5                                              # $500: rounding noise below this is not a line


# ---------------------------------------------------------------------------------------
# small numeric helpers
# ---------------------------------------------------------------------------------------
def _num(x):
    return float(x) if isinstance(x, (int, float)) else 0.0


def _addv(*vs):
    out = None
    for v in vs:
        if v is None:
            continue
        out = list(v) if out is None else [(a if a is not None else 0.0) + (b if b is not None else 0.0)
                                            for a, b in zip(out, v)]
    return out


def _div(a, b):
    if a is None or b is None or not isinstance(b, (int, float)) or abs(b) < 1e-12:
        return None
    return a / b


def _sheet_title(name, used):
    s = re.sub(r"[\[\]\:\*\?\/\\]", "-", str(name)).strip().strip("'") or "Product"
    s = s[:31].rstrip()
    base, k = s, 2
    while s.lower() in used:
        suffix = f" ({k})"
        s = base[:31 - len(suffix)].rstrip() + suffix
        k += 1
    used.add(s.lower())
    return s


# ---------------------------------------------------------------------------------------
# engagement model: everything the sheets need, derived once from cfg + res
# ---------------------------------------------------------------------------------------
class _Model:
    def __init__(self, cfg, res):
        from .timebase import submission_quarters
        self.cfg = cfg
        self.res = res
        a = cfg.get("assumptions") or {}
        self.a = a
        self.ppy = int(a.get("periods_per_year") or 4)
        self.n = int(a.get("n_periods") or 12)
        self.subq = submission_quarters(cfg)
        self.annual = _annual_periods(self.ppy, self.n, self.subq)
        self.quarterly = _quarterly_periods(self.ppy, self.n, self.subq)
        self.quarterly_ok = self.ppy >= 4
        self.years = [p for p in self.annual if not p.get("day1")]
        fin = res.get("financials") or {}
        self.bs = dict(fin.get("bs") or {})
        self.is_ = dict(fin.get("is") or {})
        derived = ((res.get("presentation") or {}).get("derived") or {})
        for k in ("totalLiab", "totalIntInc", "totalIntExp", "nie"):
            v = derived.get(k)
            if isinstance(v, list):
                (self.is_ if k in ("totalIntInc", "totalIntExp", "nie") else self.bs).setdefault(k, v)
        cap = ((res.get("capital") or {}).get("standardized") or {})
        self.cap = cap.get("ratios") or {}
        self.products = list(res.get("products") or [])
        self.lend = [p for p in self.products if p.get("family") == "lending"]
        self.dep = [p for p in self.products if p.get("family") == "deposit"]
        self.obs = [p for p in self.products if p.get("family") == "obs"]
        self.lend_cfg = self._match_cfg(self.lend, a.get("lending_products") or [])
        self.dep_cfg = self._match_cfg(self.dep, a.get("deposit_products") or [])
        self.dep_class = [self._classify_deposit(p, c) for p, c in zip(self.dep, self.dep_cfg)]
        self._curves = None
        self._curves_err = None

    # -- config matching (ordinal first, name fallback) ---------------------------------
    @staticmethod
    def _match_cfg(prods, cfgs):
        out = []
        for i, p in enumerate(prods):
            c = cfgs[i] if i < len(cfgs) and cfgs[i].get("name") == p.get("name") else None
            if c is None:
                c = next((x for x in cfgs if x.get("name") == p.get("name")), {})
            out.append(c or {})
        return out

    @staticmethod
    def _classify_deposit(p, c):
        """'nib' only when the product is contractually and actually non-interest-bearing."""
        if c.get("rate_type") == "float":
            return "ib"
        if _num(c.get("rate_paid_ann")) > 0 or _num(c.get("yield_ann")) > 0:
            return "ib"
        if any(abs(_num(x)) > 1e-9 for x in (p.get("intExp") or [])):
            return "ib"
        return "nib"

    # -- period aggregation -------------------------------------------------------------
    def stock(self, series, periods):
        s, d1 = _norm(series, periods, self.n)
        out = []
        for p in periods:
            if p.get("day1"):
                out.append(d1)
            else:
                i = p["stock_idx"]
                out.append(s[i] if 0 <= i < len(s) and s[i] is not None else None)
        return out

    def flow(self, series, periods):
        if series is None:
            return [None] * len(periods)
        s, _ = _norm(series, periods, self.n)
        out = []
        for p in periods:
            if p.get("day1"):
                out.append(None)
                continue
            xs = [s[i] for i in p["flow_idxs"] if 0 <= i < len(s) and s[i] is not None]
            out.append(sum(xs) if xs else 0.0)
        return out

    def period_avg(self, series):
        """Per-engine-period average balance: (open + close) / 2 when an opening slot exists."""
        s = list(series or [])
        if len(s) == self.n + 1:
            return [((_num(s[i]) + _num(s[i + 1])) / 2.0) for i in range(self.n)]
        return [_num(x) for x in s[:self.n]]

    def avg(self, per_period, periods):
        out = []
        for p in periods:
            if p.get("day1"):
                out.append(None)
                continue
            xs = [per_period[i] for i in p["flow_idxs"] if 0 <= i < len(per_period)]
            out.append(sum(xs) / len(xs) if xs else None)
        return out

    def annualizer(self, periods):
        return [None if p.get("day1") else (self.ppy / float(len(p["flow_idxs"])) if p["flow_idxs"] else None)
                for p in periods]

    # -- engine series ------------------------------------------------------------------
    def bsv(self, key):
        return self.bs.get(key)

    def securities_series(self):
        return _addv(*[self.bs.get(k) for k in _SEC_KEYS if isinstance(self.bs.get(k), list)])

    def intangibles_series(self):
        v = _num(self.a.get("intangibles")) / 1000.0
        ta = self.bs.get("totalAssets") or []
        return [v] * len(ta)

    def other_assets_series(self):
        ta = self.bs.get("totalAssets") or []
        parts = [self.bs.get("cash"), self.securities_series(), self.bs.get("netLoans"),
                 self.intangibles_series()]
        tot = _addv(*[p for p in parts if p is not None]) or [0.0] * len(ta)
        return [(_num(t) - _num(x)) for t, x in zip(ta, tot)]

    def total_liab_series(self):
        tl = self.bs.get("totalLiab")
        if isinstance(tl, list):
            return tl
        ta, eq = self.bs.get("totalAssets") or [], self.bs.get("equity") or []
        return [(_num(t) - _num(e)) for t, e in zip(ta, eq)]

    def other_liab_series(self):
        return [(_num(t) - _num(d)) for t, d in zip(self.total_liab_series(), self.bs.get("deposits") or [])]

    def earning_assets_series(self):
        return _addv(self.bs.get("cash"), self.securities_series(), self.bs.get("grossLoans"),
                     self.bs.get("frbStock"))

    def isv(self, key):
        return self.is_.get(key)

    def is_sum(self, keys):
        return _addv(*[self.is_.get(k) for k in keys if isinstance(self.is_.get(k), list)])

    def interest_income(self):
        return self.is_.get("totalIntInc") or self.is_sum(["loanInt", "secInt", "cashInt"])

    def interest_expense(self):
        return self.is_.get("totalIntExp") or self.is_sum(["depExp", "borrExp"])

    def nonint_income(self):
        return self.is_sum(["fees", "gos", "servNet", "fvPnl", "otherIncome"])

    def nonint_expense(self):
        return self.is_.get("nie") or self.is_sum(["prodOpex", "feeOpex", "overhead"])

    def nii(self):
        return self.is_.get("nii") or [(_num(a) - _num(b)) for a, b in zip(self.interest_income(), self.interest_expense())]

    # -- product helpers ----------------------------------------------------------------
    def window_total(self, series):
        return sum(v for v in self.flow(series, self.years) if v is not None)

    def has_window_activity(self, series):
        return abs(self.window_total(series)) > _EPS

    def obs_revenue_products(self):
        return [p for p in self.obs if self.has_window_activity(p.get("fees"))]

    def obs_cost_products(self):
        return [p for p in self.obs if self.has_window_activity(_addv(p.get("opex"), p.get("passCost")))]

    # -- rate curves --------------------------------------------------------------------
    def curves(self):
        if self._curves is None and self._curves_err is None:
            try:
                from .engine_q_a import _build_rate_curve_set
                from .validate_q import validate_config_v2
                c = copy.deepcopy(self.cfg)
                validate_config_v2(c)
                self._curves = _build_rate_curve_set(copy.deepcopy(c["assumptions"]), c, self.ppy)
            except Exception as e:      # presentation only; never block the exhibit
                self._curves_err = str(e)[:160]
        return self._curves

    def curve_path(self, name):
        cs = self.curves()
        fn = (cs or {}).get(name) if cs else None
        if fn is None:
            return None
        try:
            return [float(fn(t)) for t in range(1, self.n + 1)]
        except Exception:
            return None

    # -- opex detail --------------------------------------------------------------------
    def opex_categories(self):
        """[(name, $000s per engine period)] for each configured Operating Expense category.

        Base path is the engine's own recognized category series (dollars); additive linked
        components come from the engine result (already $000s)."""
        nd = (self.a.get("nie_detail") or {})
        cats = nd.get("categories") or []
        if not cats:
            return []
        try:
            from .growth import growth_context_from_cfg
            from .income_modules import nie_detail_series
            from .validate_q import validate_config_v2
            c = copy.deepcopy(self.cfg)
            validate_config_v2(c)
            nds = nie_detail_series(c["assumptions"], self.ppy, growth_context_from_cfg(c, self.ppy))
            base = (nds or {}).get("_category_series") or []
        except Exception:
            return []
        comps = ((self.res.get("nie_detail_series") or {}).get("components_by_category") or {})
        out = []
        for ci, cat in enumerate(cats):
            b = base[ci] if ci < len(base) else [0.0] * self.n
            lk = comps.get(str(ci)) or [0.0] * self.n
            out.append((cat.get("name") or f"Category {ci + 1}",
                        [(_num(b[i]) / 1000.0 if i < len(b) else 0.0) + (_num(lk[i]) if i < len(lk) else 0.0)
                         for i in range(self.n)]))
        return out


# ---------------------------------------------------------------------------------------
# sheet writer
# ---------------------------------------------------------------------------------------
class _Sheet:
    def __init__(self, wb, title, bank, first=False):
        self.ws = wb.active if first else wb.create_sheet()
        self.ws.title = title
        self.row = 1
        self.ncols = 1
        self.ws.column_dimensions["A"].width = 58
        try:
            self.ws.oddHeader.left.text = bank
            self.ws.oddFooter.right.text = "Page &P of &N"
        except Exception:
            pass

    def title(self, text):
        c = self.ws.cell(row=self.row, column=1, value=text)
        c.font = _TITLE_FONT
        self.row += 1

    def header(self, label, cols, width=13):
        r = self.row
        c = self.ws.cell(row=r, column=1, value=label)
        c.font, c.fill = _HDR_FONT, _HDR_FILL
        for j, h in enumerate(cols):
            cc = self.ws.cell(row=r, column=2 + j, value=h)
            cc.font, cc.fill = _HDR_FONT, _HDR_FILL
            cc.alignment = Alignment(horizontal="center")
            self.ws.column_dimensions[get_column_letter(2 + j)].width = width
        self.ncols = max(self.ncols, 1 + len(cols))
        self.row += 1
        self.ws.freeze_panes = self.ws.cell(row=self.row, column=2)

    def quarterly_header(self, label, periods, include_day1):
        """Two header rows: year groups (merged) over Q1..Q4 labels, as in the template."""
        r1, r2 = self.row, self.row + 1
        cols = [p for p in periods if include_day1 or not p.get("day1")]
        c = self.ws.cell(row=r1, column=1, value=label)
        c.font, c.fill = _HDR_FONT, _HDR_FILL
        self.ws.cell(row=r2, column=1).fill = _HDR_FILL
        self.ws.merge_cells(start_row=r1, start_column=1, end_row=r2, end_column=1)
        groups = []
        for j, p in enumerate(cols):
            col = 2 + j
            self.ws.column_dimensions[get_column_letter(col)].width = 11
            if p.get("day1"):
                cc = self.ws.cell(row=r1, column=col, value="Day 1")
                cc.font, cc.fill = _HDR_FONT, _HDR_FILL
                cc.alignment = Alignment(horizontal="center")
                self.ws.cell(row=r2, column=col).fill = _HDR_FILL
                self.ws.merge_cells(start_row=r1, start_column=col, end_row=r2, end_column=col)
                continue
            y, q = p["label"].split()
            if groups and groups[-1][0] == y:
                groups[-1][2] = col
            else:
                groups.append([y, col, col])
            cq = self.ws.cell(row=r2, column=col, value=q)
            cq.font, cq.fill = _HDR_FONT, _HDR_FILL
            cq.alignment = Alignment(horizontal="center")
        for y, a, b in groups:
            cy = self.ws.cell(row=r1, column=a, value="Year " + y.lstrip("Y"))
            cy.font, cy.fill = _HDR_FONT, _HDR_FILL
            cy.alignment = Alignment(horizontal="center")
            for col in range(a + 1, b + 1):
                self.ws.cell(row=r1, column=col).fill = _HDR_FILL
            if b > a:
                self.ws.merge_cells(start_row=r1, start_column=a, end_row=r1, end_column=b)
        self.ncols = max(self.ncols, 1 + len(cols))
        self.row += 2
        self.ws.freeze_panes = self.ws.cell(row=self.row, column=2)
        return cols

    def section(self, label):
        r = self.row
        c = self.ws.cell(row=r, column=1, value=label)
        c.font, c.fill = _SEC_FONT, _SEC_FILL
        for col in range(2, self.ncols + 1):
            self.ws.cell(row=r, column=col).fill = _SEC_FILL
        if self.ncols > 1:
            self.ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=self.ncols)
        self.row += 1

    def line(self, label, values, fmt=_MONEY, bold=False, indent=0):
        r = self.row
        c = self.ws.cell(row=r, column=1, value=label)
        c.font = _TOTAL_FONT if bold else _LBL_FONT
        if indent:
            c.alignment = Alignment(indent=indent)
        for j, v in enumerate(values):
            if isinstance(v, float) and (v != v):   # NaN guard
                v = None
            cc = self.ws.cell(row=r, column=2 + j, value=(round(v, 6) if isinstance(v, float) else v))
            cc.font = _TOTAL_FONT if bold else _LBL_FONT
            cc.number_format = fmt
            cc.alignment = Alignment(horizontal="right")
        self.row += 1

    def notes(self, lines):
        if not lines:
            return
        self.row += 1
        for t in lines:
            c = self.ws.cell(row=self.row, column=1, value=t)
            c.font = _NOTE_FONT
            c.alignment = Alignment(wrap_text=True, vertical="top")
            span = max(self.ncols, 6)
            self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=span)
            self.ws.row_dimensions[self.row].height = 11 * max(1, (len(t) // 200) + 1)
            self.row += 1


# ---------------------------------------------------------------------------------------
# sheets
# ---------------------------------------------------------------------------------------
def _bs_annual(wb, m, bank):
    sh = _Sheet(wb, "Balance Sheet EOP (Annual)", bank, first=True)
    P = m.annual
    sh.title("Balance Sheet Period-End Annual ($000)")
    sh.header("Balance Sheet", [p["label"] for p in P])
    sh.section("Assets")
    sh.line("Cash and Balances Due from Banks", m.stock(m.bsv("cash"), P))
    sh.line("Securities", m.stock(m.securities_series(), P))
    sh.line("Loans, Net of Allowance for Credit Losses", m.stock(m.bsv("netLoans"), P))
    sh.line("Other Assets", m.stock(m.other_assets_series(), P))
    sh.line("Intangible Assets", m.stock(m.intangibles_series(), P))
    sh.line("Total Assets", m.stock(m.bsv("totalAssets"), P), bold=True)
    sh.section("Liabilities and Equity")
    sh.line("Deposits", m.stock(m.bsv("deposits"), P))
    sh.line("Other Liabilities", m.stock(m.other_liab_series(), P))
    sh.line("Total Liabilities", m.stock(m.total_liab_series(), P), bold=True)
    sh.line("Equity", m.stock(m.bsv("equity"), P), bold=True)
    sh.notes(_BS_NOTES)


_BS_NOTES = [
    "Period-end balances: Day 1 is the opening balance sheet; each year is its final submission quarter-end.",
    "Securities = all AFS, HTM and designated securities books. Loans are net of the allowance and include "
    "loans held for sale at carrying value. Other Assets = Total Assets less the lines shown (premises, "
    "Federal Reserve stock, MSRs, deferred tax, prepaid and other assets). Intangible Assets = configured intangibles.",
    "Other Liabilities = Total Liabilities less Deposits (borrowings, accrued and other liabilities).",
]


def _is_rows(m, P):
    return [
        ("Interest Income", m.flow(m.interest_income(), P), False),
        ("Interest Expense", m.flow(m.interest_expense(), P), False),
        ("Provision for Credit Losses", m.flow(m.isv("prov"), P), False),
        ("Non-Interest Income", m.flow(m.nonint_income(), P), False),
        ("Non-Interest Expense", m.flow(m.nonint_expense(), P), False),
        ("Pre-Tax Net Income", m.flow(m.isv("pretax"), P), True),
        ("Taxes", m.flow(m.isv("tax"), P), False),
        ("Net Income", m.flow(m.isv("ni"), P), True),
    ]


def _is_annual(wb, m, bank):
    sh = _Sheet(wb, "Income Statement (Annual)", bank)
    sh.title("Income Statement (Annual) ($000)")
    sh.header("Income Statement", [p["label"] for p in m.years])
    for label, vals, bold in _is_rows(m, m.years):
        sh.line(label, vals, bold=bold)
    sh.notes(["Flows summed over each submission year. Pre-Tax Net Income = Interest Income − Interest Expense − "
              "Provision + Non-Interest Income − Non-Interest Expense."])


def _bs_pct(wb, m, bank):
    sh = _Sheet(wb, "Balance Sheet % of Asset", bank)
    Y = m.years
    ta = m.stock(m.bsv("totalAssets"), Y)
    pct = lambda s: [_div(v, t) for v, t in zip(m.stock(s, Y), ta)]
    sh.title("Balance Sheet Period-End (as % total assets)")
    sh.header("Balance Sheet (% of Total Assets)", [p["label"] for p in Y])
    sh.section("Assets")
    sh.line("Cash and Balances Due from Banks", pct(m.bsv("cash")), fmt=_PCT)
    sh.line("Securities", pct(m.securities_series()), fmt=_PCT)
    sh.line("Loans, Net of Allowance for Credit Losses", pct(m.bsv("netLoans")), fmt=_PCT)
    sh.line("Other Assets", pct(m.other_assets_series()), fmt=_PCT)
    sh.line("Intangible Assets", pct(m.intangibles_series()), fmt=_PCT)
    sh.line("Total Assets", pct(m.bsv("totalAssets")), fmt=_PCT, bold=True)
    sh.section("Liabilities and Equity")
    sh.line("Deposits", pct(m.bsv("deposits")), fmt=_PCT)
    sh.line("Other Liabilities", pct(m.other_liab_series()), fmt=_PCT)
    sh.line("Total Liabilities", pct(m.total_liab_series()), fmt=_PCT, bold=True)
    sh.line("Equity", pct(m.bsv("equity")), fmt=_PCT, bold=True)
    sh.notes(["Each line ÷ year-end Total Assets. Line definitions as on Balance Sheet EOP (Annual)."])


def _avg_total_assets(m, P):
    return m.avg(m.period_avg(m.bsv("totalAssets")), P)


def _is_pct(wb, m, bank):
    sh = _Sheet(wb, "Income Statement % of Asset", bank)
    Y = m.years
    ata = _avg_total_assets(m, Y)
    ann = m.annualizer(Y)
    sh.title("Income Statement (as % total assets)")
    sh.header("Income Statement (% of Total Assets)", [p["label"] for p in Y])
    for label, vals, bold in _is_rows(m, Y):
        if label == "Net Income":
            label = "Net Income (Loss)"
        sh.line(label, [_div(v * f, t) if v is not None and f else None for v, f, t in zip(vals, ann, ata)],
                fmt=_RATE, bold=bold)
    sh.notes(["Annual flow ÷ average total assets for the year (mean of engine-period average balances). "
              "Net Income (Loss) on this basis equals ROAA on Financial Performance Highlight."])


def _fph(wb, m, bank):
    sh = _Sheet(wb, "Financial Performance Highlight", bank)
    Y = m.years
    ann = m.annualizer(Y)
    sh.title("Financial Performance Highlights")
    sh.header("Financial Performance Highlights", [p["label"] for p in Y])
    lev = [(_div(v, 100.0)) for v in m.stock(m.cap.get("leverage"), Y)] if isinstance(m.cap.get("leverage"), list) \
        else [None] * len(Y)
    sh.line("Tier 1 Leverage Ratio", lev, fmt=_PCT)
    ni = m.flow(m.isv("ni"), Y)
    ii = m.flow(m.interest_income(), Y)
    nii = m.flow(m.nii(), Y)
    nonii = m.flow(m.nonint_income(), Y)
    nie = m.flow(m.nonint_expense(), Y)
    ata = _avg_total_assets(m, Y)
    aeq = m.avg(m.period_avg(m.bsv("equity")), Y)
    aea = m.avg(m.period_avg(m.earning_assets_series()), Y)
    annz = lambda vals, den: [_div(v * f, d) if v is not None and f else None for v, f, d in zip(vals, ann, den)]
    sh.section("Profitability")
    sh.line("ROAE", annz(ni, aeq), fmt=_PCT, indent=1)
    sh.line("ROAA", annz(ni, ata), fmt=_PCT, indent=1)
    sh.line("Yield on Earning Assets", annz(ii, aea), fmt=_PCT, indent=1)
    sh.line("Net Interest Margin", annz(nii, aea), fmt=_PCT, indent=1)
    sh.line("Efficiency Ratio", [_div(e, (_num(a) + _num(b))) if (_num(a) + _num(b)) > 0 else None
                                 for e, a, b in zip(nie, nii, nonii)], fmt=_PCT, indent=1)
    sh.section("Liquidity")
    ta = m.stock(m.bsv("totalAssets"), Y)
    dep = m.stock(m.bsv("deposits"), Y)
    gl = m.stock(m.bsv("grossLoans"), Y)
    sh.line("Cash/Total Assets", [_div(c, t) for c, t in zip(m.stock(m.bsv("cash"), Y), ta)], fmt=_PCT, indent=1)
    sh.line("Loans/Deposits", [_div(g, d) for g, d in zip(gl, dep)], fmt=_PCT, indent=1)
    sh.section("Deposit Mix")
    for cls, label in (("nib", "Non-Interest-Bearing Deposits"), ("ib", "Interest-Bearing Deposits")):
        members = [p for p, c in zip(m.dep, m.dep_class) if c == cls]
        bal = _addv(*[p.get("bal") for p in members]) if members else None
        sh.line(label, [_div(b, d) for b, d in zip(m.stock(bal, Y), dep)] if bal else [None] * len(Y),
                fmt=_PCT, indent=1)
        if len(members) > 1:
            for p in members:
                sh.line(p.get("name", "Deposit"), [_div(b, d) for b, d in zip(m.stock(p.get("bal"), Y), dep)],
                        fmt=_PCT, indent=2)
    hfs = m.stock(m.bsv("hfs"), Y) if isinstance(m.bsv("hfs"), list) else [None] * len(Y)
    sh.line("Loans held for sale", [_div(h, g) for h, g in zip(hfs, gl)], fmt=_PCT, indent=1)
    sh.section("Loan Mix")
    hfi = _addv(*[p.get("bal") for p in m.lend]) if m.lend else None
    hfi_ye = m.stock(hfi, Y) if hfi else [None] * len(Y)
    for p in m.lend:
        sh.line(p.get("name", "Loan Product"), [_div(b, t) for b, t in zip(m.stock(p.get("bal"), Y), hfi_ye)],
                fmt=_PCT, indent=1)
    if not m.lend:
        sh.line("(no asset products in this engagement)", [None] * len(Y), indent=1)
    sh.section("Program Migration")
    rev = m.obs_revenue_products()
    for p in rev:
        sh.line(p.get("name", "Program"), [None] * len(Y), fmt=_PCT, indent=1)
    if not rev:
        sh.line("(no off-balance-sheet revenue programs in this engagement)", [None] * len(Y), indent=1)
    sh.notes([
        "Tier 1 Leverage Ratio: standardized capital module, year-end. ROAE / ROAA: annual net income ÷ average "
        "equity / average total assets. Yield on Earning Assets and Net Interest Margin: annual total interest "
        "income / net interest income ÷ average earning assets (cash, securities, Federal Reserve stock and gross "
        "loans incl. held for sale). Efficiency Ratio: non-interest expense ÷ (net interest income + non-interest income).",
        "Liquidity, Deposit Mix and Loan Mix are year-end. Deposit Mix classifies each liability product: "
        "non-interest-bearing when its configured rate is zero (fixed) and it paid no interest in the horizon, "
        "otherwise interest-bearing. Loans held for sale ÷ gross loans. Loan Mix = each asset product ÷ total "
        "loans held for investment.",
        "Program Migration rows list this engagement's off-balance-sheet revenue programs. Migration share is "
        "not a separately produced engine output, so these cells are intentionally blank for entry.",
    ])


def _bs_quarterly(wb, m, bank):
    sh = _Sheet(wb, "Balance Sheet Quarterly", bank)
    P = m.quarterly
    sh.title("Balance Sheet Period-End (Quarterly) ($000)")
    cols = sh.quarterly_header("Balance Sheet", P, include_day1=True)
    q = (lambda s: m.stock(s, cols)) if m.quarterly_ok else (lambda s: [None] * len(cols))
    sh.section("Assets")
    sh.line("Cash and Balances Due from Banks", q(m.bsv("cash")))
    sh.line("Securities", q(m.securities_series()))
    sh.line("Loans, Net of Allowance for Credit Losses", q(m.bsv("netLoans")))
    sh.line("Other Assets", q(m.other_assets_series()))
    sh.line("Intangible Assets", q(m.intangibles_series()))
    sh.line("Total Assets", q(m.bsv("totalAssets")), bold=True)
    sh.section("Liabilities and Equity")
    for cls, label in (("nib", "Non-Interest-Bearing Deposits"), ("ib", "Interest-Bearing Deposits")):
        members = [p for p, c in zip(m.dep, m.dep_class) if c == cls]
        bal = _addv(*[p.get("bal") for p in members]) if members else None
        sh.line(label, q(bal) if bal else [None] * len(cols))
        if len(members) > 1:
            for p in members:
                sh.line(p.get("name", "Deposit"), q(p.get("bal")), indent=1)
    sh.line("Other Liabilities", q(m.other_liab_series()))
    sh.line("Total Liabilities", q(m.total_liab_series()), bold=True)
    sh.line("Equity", q(m.bsv("equity")), bold=True)
    notes = ["Quarter-end balances over the submission window; line definitions as on Balance Sheet EOP (Annual). "
             "Deposits are split by the Deposit Mix classification; product rows appear when a class holds more "
             "than one liability product."]
    if not m.quarterly_ok:
        notes.append("This engagement runs on an annual cadence, so quarter-end values are not produced.")
    sh.notes(notes)


def _is_quarterly(wb, m, bank):
    sh = _Sheet(wb, "Income Statement Quarterly", bank)
    P = [p for p in m.quarterly if not p.get("day1")]
    sh.title("Income Statement Period-End (Quarterly) ($000)")
    cols = sh.quarterly_header("Income Statement", P, include_day1=False)
    for label, vals, bold in _is_rows(m, cols):
        sh.line(label, vals if m.quarterly_ok else [None] * len(cols), bold=bold)
    notes = ["Flows within each submission quarter."]
    if not m.quarterly_ok:
        notes.append("This engagement runs on an annual cadence, so quarterly flows are not produced.")
    sh.notes(notes)


def _rate_descriptor(c):
    if c.get("rate_type") == "float":
        idx = str(c.get("index") or "sofr").upper()
        beta = c.get("deposit_beta")
        sp = _num(c.get("index_spread")) * 100.0
        b = "" if beta is None else f"{float(beta) * 100:g}% × "
        sign = "+" if sp >= 0 else "−"
        return f"{b}{idx} {sign} {abs(sp):.2f}%"
    r = c.get("rate_paid_ann")
    return f"fixed {(_num(r) * 100):.2f}%" if r is not None else None


def _rates(wb, m, bank):
    sh = _Sheet(wb, "Interest Rate Assumptions", bank)
    Y = m.years
    sh.title("Interest Rate Assumptions")
    sh.header("Rate Assumption", [p["label"] for p in Y])
    used_idx = {str(c.get("index") or "sofr").lower() for c in (m.lend_cfg + m.dep_cfg)
                if c.get("rate_type") == "float"}
    for key, label in (("effr", "Federal Funds Rate (EFFR)"), ("sofr", "SOFR"), ("prime", "Prime Rate")):
        if key == "prime" and "prime" not in used_idx:
            continue
        path = m.curve_path(key)
        sh.line(label, m.avg(path, Y) if path else [None] * len(Y), fmt=_RATE)
    ann = m.annualizer(Y)

    def realized(flow, bal_series):
        f = m.flow(flow, Y)
        b = m.avg(m.period_avg(bal_series) if bal_series is not None else [], Y)
        return [_div(v * k, d) if v is not None and k else None for v, k, d in zip(f, ann, b)]

    def realized_avg(flow, per_period_avg):
        f = m.flow(flow, Y)
        b = m.avg(per_period_avg, Y)
        return [_div(v * k, d) if v is not None and k else None for v, k, d in zip(f, ann, b)]

    sh.line("Cash and Balances Due from Banks (earned yield)", realized(m.isv("cashInt"), m.bsv("cash")), fmt=_RATE)
    sh.line("Securities (earned yield)", realized(m.isv("secInt"), m.securities_series()), fmt=_RATE)
    for cls, label in (("ib", "Interest-Bearing Deposits"), ("nib", "Non-Interest-Bearing Deposits")):
        members = [(p, c) for p, c, k in zip(m.dep, m.dep_cfg, m.dep_class) if k == cls]
        if not members:
            continue
        ie = _addv(*[p.get("intExp") for p, _ in members])
        avgb = _addv(*[p.get("avg") for p, _ in members])
        desc = _rate_descriptor(members[0][1]) if len(members) == 1 else None
        sh.line(label + (f" ({desc})" if desc and cls == "ib" else ""), realized_avg(ie, avgb or []), fmt=_RATE)
        if len(members) > 1:
            for p, c in members:
                d = _rate_descriptor(c)
                sh.line(p.get("name", "Deposit") + (f" ({d})" if d else ""),
                        realized_avg(p.get("intExp"), p.get("avg") or []), fmt=_RATE, indent=1)
    sh.line("Blended Cost of On-Balance-Sheet Deposits", realized(m.isv("depExp"), m.bsv("deposits")), fmt=_RATE)
    if m.has_window_activity(m.isv("borrExp")):
        bor = _addv(m.bsv("borrow"), m.bsv("borrowSched"))
        sh.line("Borrowed Funds (cost)", realized(m.isv("borrExp"), bor), fmt=_RATE)
    notes = ["Index rows are annual averages of the engagement's base-case curves as the engine evaluates them "
             "(Rates & Curves). Yields and costs are realized: annual interest ÷ average balance. Rate descriptors "
             "in labels are this engagement's configured pricing."]
    if m._curves_err:
        notes.append("Index curves could not be evaluated for this configuration; index rows are blank.")
    sh.notes(notes)


def _nii(wb, m, bank):
    sh = _Sheet(wb, "Net Interest Income", bank)
    Y = m.years
    sh.title("Net Interest Income ($000)")
    sh.header("", [p["label"] for p in Y])
    sh.line("Interest Income", m.flow(m.interest_income(), Y))
    sh.line("Interest Expense", m.flow(m.interest_expense(), Y))
    sh.line("Net Interest Income", m.flow(m.nii(), Y), bold=True)


def _ii_calc(wb, m, bank):
    sh = _Sheet(wb, "Interest Income Calculation", bank)
    Y = m.years
    ann = m.annualizer(Y)
    sh.title("Interest Income Calculation")
    sh.header("", [p["label"] for p in Y])
    aea = m.avg(m.period_avg(m.earning_assets_series()), Y)
    ii = m.flow(m.interest_income(), Y)
    sh.line("Total Average Earning Assets", aea)
    sh.line("Yield", [_div(v * k, d) if v is not None and k else None for v, k, d in zip(ii, ann, aea)], fmt=_RATE)
    sh.notes(["Average earning assets = cash, securities, Federal Reserve stock and gross loans (incl. held for "
              "sale), averaged over the year. Yield = total interest income ÷ average earning assets."])


def _loan_term_years(m, c):
    """(years, basis) from the product's configured repayment contract; None when not derivable."""
    prof = c.get("repayment_profile")
    if isinstance(prof, dict) and isinstance(prof.get("repayment_fractions"), list):
        fr = [_num(x) for x in prof["repayment_fractions"]]
        if sum(fr) > 0:
            return sum(i * f for i, f in enumerate(fr)) / sum(fr) / 12.0, "repayment-weighted average age of the monthly repayment curve"
    if c.get("structure") == "term" and _num(c.get("term_q")) > 0:
        return _num(c["term_q"]) / 4.0, "contractual term"
    ro = _num(c.get("runoff_per_period"))
    if ro > 0:
        return 1.0 / (ro * m.ppy), "implied by the configured runoff rate (1 ÷ annual runoff)"
    return None, None


def _product_sheets(wb, m, bank, used):
    titles = []
    for p, c in zip(m.lend, m.lend_cfg):
        name = p.get("name", "Loan Product")
        sh = _Sheet(wb, _sheet_title(f"{name} Assumptions", used), bank)
        titles.append(sh.ws.title)
        Y = m.years
        ann = m.annualizer(Y)
        sh.title(f"{name} Interest Income Assumptions")
        sh.header("", [x["label"] for x in Y])
        vol = p.get("fundedVolume") if isinstance(p.get("fundedVolume"), list) else p.get("origq")
        sh.line("Total Bank Originations ($000)", m.flow(vol, Y))
        ii = m.flow(p.get("intInc"), Y)
        av = m.avg([_num(x) for x in (p.get("avg") or [])], Y)
        sh.line("Average Asset Yield (%)", [_div(v * k, d) if v is not None and k else None
                                            for v, k, d in zip(ii, ann, av)], fmt=_RATE)
        yrs, basis = _loan_term_years(m, c)
        sh.line("Average Loan Term (years)", [yrs] * len(Y), fmt=_YEARS)
        n = ["Originations: funded volume (or new originations) summed over the year. Average Asset Yield: "
             "realized product interest income ÷ average product balance."]
        n.append(f"Average Loan Term: {basis}." if basis else
                 "Average Loan Term: not derivable from this product's configuration; left blank.")
        sh.notes(n)
    return titles


def _capital(wb, m, bank):
    sh = _Sheet(wb, "Forecasted Capital Levels", bank)
    P = m.annual
    sh.title("Forecasted Capital Levels")
    sh.header("", [p["label"] for p in P])
    for key, label in (("leverage", "Tier 1 Leverage Ratio"), ("cet1_rwa", "CET 1 Capital Ratio"),
                       ("tier1_rwa", "Tier 1 Capital Ratio"), ("total_rwa", "Total Risk-Based Capital Ratio")):
        s = m.cap.get(key)
        sh.line(label, [_div(v, 100.0) for v in m.stock(s, P)] if isinstance(s, list) else [None] * len(P), fmt=_PCT)
    sh.notes(["Standardized capital module, year-end. Day 1 is blank where the engine reports ratios only for "
              "projected periods."])


def _nonint_income(wb, m, bank):
    sh = _Sheet(wb, "Non-Interest Income", bank)
    Y = m.years
    sh.title("Non-Interest Income ($000)")
    sh.header("", [p["label"] for p in Y])
    shown = []
    for p in m.obs_revenue_products():
        v = m.flow(p.get("fees"), Y); shown.append(v); sh.line(p.get("name", "Program"), v)
    credit = _addv(*[x for p in m.lend for x in (p.get("fees"), p.get("gos"), p.get("servNet"))])
    if credit is not None and m.has_window_activity(credit):
        v = m.flow(credit, Y); shown.append(v); sh.line("Credit Product Fees and Servicing", v)
    for p in m.dep:
        if m.has_window_activity(p.get("fees")):
            v = m.flow(p.get("fees"), Y); shown.append(v); sh.line(f"{p.get('name', 'Deposit')} Fees", v)
    if m.has_window_activity(m.isv("fvPnl")):
        v = m.flow(m.isv("fvPnl"), Y); shown.append(v); sh.line("Net Gains (Losses) on Fair-Value Instruments", v)
    for ln in ((m.res.get("other_income_detail") or {}).get("lines") or []):
        if m.has_window_activity(ln.get("amount")):
            v = m.flow(ln.get("amount"), Y); shown.append(v)
            sh.line(ln.get("name") or ln.get("label") or "Other income", v)
    total = m.flow(m.nonint_income(), Y)
    resid = [(_num(t) - sum(_num(s[j]) for s in shown)) for j, t in enumerate(total)]
    if any(abs(x) > _RESID for x in resid):
        sh.line("Other Non-Interest Income, Net", resid)
    sh.line("Total", total, bold=True)
    sh.notes(["One row per off-balance-sheet product with revenue in the submission window, named as in this "
              "engagement. Credit Product Fees and Servicing = asset-product fees, gains on sale and net servicing. "
              "Total ties to Non-Interest Income on the Income Statement; any unattributed remainder is shown as Other."])


def _nonint_expense(wb, m, bank):
    sh = _Sheet(wb, "Non-Interest Expense", bank)
    Y = m.years
    sh.title("Non-Interest Expense ($000)")
    sh.header("", [p["label"] for p in Y])
    shown = []
    for p in m.obs_cost_products():
        v = m.flow(_addv(p.get("opex"), p.get("passCost")), Y); shown.append(v)
        sh.line(f"Cost of Revenue: {p.get('name', 'Program')}", v)
    for fam, label in ((m.lend, "Credit Product Operating Expense"), (m.dep, "Deposit Product Operating Expense")):
        s = _addv(*[p.get("opex") for p in fam]) if fam else None
        if s is not None and m.has_window_activity(s):
            v = m.flow(s, Y); shown.append(v); sh.line(label, v)
    oh = m.isv("overhead") or m.is_sum(["workforceComp", "otherOpex", "depreciationExpense"])
    v = m.flow(oh, Y); shown.append(v)
    sh.line("De Novo Bank Operating Expense", v)
    total = m.flow(m.nonint_expense(), Y)
    resid = [(_num(t) - sum(_num(s[j]) for s in shown)) for j, t in enumerate(total)]
    if any(abs(x) > _RESID for x in resid):
        sh.line("Other Non-Interest Expense", resid)
    sh.line("Total", total, bold=True)
    sh.notes(["Cost of Revenue rows: each off-balance-sheet product's operating and fee-stream costs. De Novo Bank "
              "Operating Expense is detailed on the next sheet. Total ties to Non-Interest Expense on the Income Statement."])


def _denovo(wb, m, bank):
    sh = _Sheet(wb, "De Novo Bank Operating Expenses", bank)
    Y = m.years
    sh.title("De Novo Bank Operating Expenses ($000)")
    sh.header("", [p["label"] for p in Y])
    shown = []
    comp = m.isv("workforceComp")
    if comp is not None:
        v = m.flow(comp, Y); shown.append(v); sh.line("Bank Compensation (executives and staff)", v)
    cats = m.opex_categories()
    for name, series in cats:
        v = m.flow(series, Y); shown.append(v); sh.line(name, v)
    dep = m.isv("depreciationExpense")
    if dep is not None and m.has_window_activity(dep):
        v = m.flow(dep, Y); shown.append(v); sh.line("Depreciation", v)
    oh = m.isv("overhead") or m.is_sum(["workforceComp", "otherOpex", "depreciationExpense"])
    total = m.flow(oh, Y)
    resid = [(_num(t) - sum(_num(s[j]) for s in shown)) for j, t in enumerate(total)]
    if any(abs(x) > _RESID for x in resid):
        sh.line("Regulatory Assessments, Gross-Up and Other Operating Expense" if cats else "Other Operating Expense",
                resid)
    sh.line("Total", total, bold=True)
    sh.notes(["Rows are this engagement's Operating Expense categories as configured (base path plus linked "
              "components). The remainder line carries engine-computed items not owned by a category, such as "
              "FDIC/OCC assessments and expense gross-up. Total ties to De Novo Bank Operating Expense."])


def _product_map(wb, m, bank, asset_sheets):
    sh = _Sheet(wb, "Product Map & Notes", bank)
    from .timebase import submission_quarters
    sh.title(f"{bank} — Business Plan Tables: product map")
    sh.header(f"Product (Q1–Q{submission_quarters(m.cfg)} submission totals, $000s)",
              ["Class", "End Balance", "Interest Income", "Interest Expense", "Fees", "Appears in"], width=15)
    sh.ws.column_dimensions[get_column_letter(7)].width = 70
    end = m.years[-1] if m.years else None
    rev = {id(p) for p in m.obs_revenue_products()}
    cost = {id(p) for p in m.obs_cost_products()}
    asset_iter = iter(asset_sheets)
    for p in m.products:
        fam = p.get("family")
        if fam == "lending":
            cls, where = "Asset", f"Loan Mix; '{next(asset_iter, '')}' sheet"
        elif fam == "deposit":
            k = m.dep_class[m.dep.index(p)]
            cls = "Liability"
            where = ("Deposit Mix and Balance Sheet Quarterly: " +
                     ("Non-Interest-Bearing" if k == "nib" else "Interest-Bearing") + "; Interest Rate Assumptions")
            if m.has_window_activity(p.get("fees")):
                where += "; Non-Interest Income"
        else:
            cls = "Off-BS"
            parts = []
            if id(p) in rev:
                parts += ["Non-Interest Income", "Program Migration"]
            if id(p) in cost:
                parts.append("Non-Interest Expense (Cost of Revenue)")
            where = "; ".join(parts) or "No revenue or cost in the submission window (driver/exposure only)"
        bal = m.stock(p.get("bal"), [end])[0] if end and isinstance(p.get("bal"), list) else None
        vals = [cls, bal, m.window_total(p.get("intInc")), m.window_total(p.get("intExp")),
                m.window_total(p.get("fees")), where]
        r = sh.row
        sh.line(p.get("name", "Product"), vals)
        sh.ws.cell(row=r, column=2).alignment = Alignment(horizontal="center")
        sh.ws.cell(row=r, column=7).alignment = Alignment(horizontal="left", wrap_text=True)
    gh = (m.res.get("run_hash") or m.res.get("config_hash") or "")[:12]
    sh.notes([
        "Rows on the product-driven sheets come from this engagement's product registry (Detail → Per-Product "
        "Contributions): Asset products drive Loan Mix and one assumption sheet each; Liability products drive the "
        "deposit split; Off-BS products drive Non-Interest Income, Cost of Revenue and Program Migration.",
        "All amounts in $ thousands from a fresh base-case engine run of the posted configuration. Annual: stocks "
        f"at submission year-end, flows summed over the year. Submission window: Q1–Q{m.subq} calendar quarters.",
        "Blank cells mark lines the engine does not produce (not zero).",
        f"Run hash: {gh}",
    ])


def build_enhanced_bpt(cfg, res):
    """Return (xlsx_bytes, filename) for the enhanced, engagement-driven Business Plan Tables."""
    m = _Model(cfg, res)
    bank = _bank_name(cfg)
    wb = Workbook()
    used = {s.lower() for s in SHEETS_FIXED_HEAD + SHEETS_FIXED_TAIL}
    _bs_annual(wb, m, bank)
    _is_annual(wb, m, bank)
    _bs_pct(wb, m, bank)
    _is_pct(wb, m, bank)
    _fph(wb, m, bank)
    _bs_quarterly(wb, m, bank)
    _is_quarterly(wb, m, bank)
    _rates(wb, m, bank)
    _nii(wb, m, bank)
    _ii_calc(wb, m, bank)
    asset_sheets = _product_sheets(wb, m, bank, used)
    _capital(wb, m, bank)
    _nonint_income(wb, m, bank)
    _nonint_expense(wb, m, bank)
    _denovo(wb, m, bank)
    _product_map(wb, m, bank, asset_sheets)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"{_engagement_slug(cfg)}_Foundry_Business_Plan_Table.xlsx"
