"""r313: every surface that reads a deposit rate uses the r312 beta / floor / cap, and both workbooks carry them."""
import copy
import io
import json
import sys
from pathlib import Path

FIX = "foundry/fixtures/universal_template_bank.json"
KEYS = ("deposit_beta", "deposit_rate_floor", "deposit_rate_cap")


def main():
    from foundry.v2 import run_q
    from foundry.v2.challenge_q import challenge_config
    from foundry.v2.peer_calibration import _flag_client_value
    from foundry.v2.excel_q import workbook_from_config_v2, parse_workbook_v2
    from foundry.v2.fiw import build_fiw, diff_import
    from openpyxl import load_workbook
    p = f = 0

    def ck(name, cond, detail=""):
        nonlocal p, f
        if cond: p += 1; print("  PASS ", name)
        else: f += 1; print("  FAIL ", name, detail)

    base = json.loads(Path(FIX).read_text(encoding="utf-8"))

    # --- funding checks read the rate the engine pays ---
    c = copy.deepcopy(base)
    d = next(x for x in c["assumptions"]["deposit_products"] if x.get("call_report_line") == "depDDA")
    d.update(rate_type="float", index="sofr", index_spread=0.0); d.pop("rate_paid_ann", None)

    def dda_flags(cfg):
        out = challenge_config(copy.deepcopy(cfg))
        out = out.get("flags", out) if isinstance(out, dict) else out
        return [x for x in out if x.get("id") == "FUND-DDA" and d["name"] in x.get("text", "")]

    ck("control: a DDA at 100% of SOFR is flagged as paying over 2%", len(dda_flags(c)) == 1)
    d["deposit_beta"] = 0.3
    paid = next(x for x in run_q.run_v2(copy.deepcopy(c))["products"] if x["name"] == d["name"])["rateQ"][0]
    ck("the engine pays about 30% of SOFR", 1.0 < paid < 1.3, paid)
    ck("the same DDA at a 30% beta is no longer flagged", dda_flags(c) == [])
    d["deposit_beta"] = 1.0; d["deposit_rate_cap"] = 0.01
    ck("a cap below 2% also clears the flag", dda_flags(c) == [])
    d["deposit_rate_cap"] = None; d["deposit_beta"] = 1.0; d["deposit_rate_floor"] = 0.06
    hot = [x for x in (lambda o: o.get("flags", o) if isinstance(o, dict) else o)(challenge_config(copy.deepcopy(c))) if x.get("id") == "FUND-HOT" and d["name"] in x.get("text", "")]
    ck("a floor above 5.5% raises the high-cost flag", len(hot) == 1)
    d["deposit_rate_floor"] = None; d["deposit_beta"] = 0.3
    full = copy.deepcopy(c); next(x for x in full["assumptions"]["deposit_products"] if x["name"] == d["name"])["deposit_beta"] = 1.0
    v30, v100 = _flag_client_value({"id": "FUND-DDA"}, c), _flag_client_value({"id": "FUND-DDA"}, full)
    ck("peer placement of the deposit-cost flag uses the beta rate", v30 is not None and v100 is not None and v30 < v100, (v30, v100))

    # --- the checks see the engine's Prime curve, not SOFR ---
    c3 = copy.deepcopy(base)
    cc = next(x for x in c3["assumptions"]["lending_products"] if x.get("index") == "prime")
    cc["index_spread"] = 0.20
    y = next(x for x in run_q.run_v2(copy.deepcopy(c3))["products"] if x["name"] == cc["name"])["rateQ"][0]
    fl = (lambda o: o.get("flags", o) if isinstance(o, dict) else o)(challenge_config(copy.deepcopy(c3)))
    us = [x for x in fl if x.get("id") == "PRICE-USURY" and cc["name"] in x.get("text", "")]
    ck("Prime + 20% (engine yield over 25%) is flagged at the engine's yield", y > 25 and len(us) == 1 and f"{y/100:.1%}" in us[0]["text"], (y, [x.get("text") for x in us]))

    # --- workbooks ---
    c2 = copy.deepcopy(base)
    m = next(x for x in c2["assumptions"]["deposit_products"] if x.get("rate_type") == "float")
    m.update(deposit_beta=0.4, deposit_rate_floor=0.005, deposit_rate_cap=0.03)
    b = io.BytesIO(); workbook_from_config_v2(copy.deepcopy(c2)).save(b)
    back = parse_workbook_v2(b.getvalue())
    bm = next(x for x in back["assumptions"]["deposit_products"] if x["name"] == m["name"])
    ck("config workbook carries beta, floor and cap", all(bm.get(k) == m[k] for k in KEYS), {k: bm.get(k) for k in KEYS})
    data, _gh = build_fiw(copy.deepcopy(c2))
    wb = load_workbook(io.BytesIO(data))
    ws = wb["ASSM_DEPOSITS"]
    idx = c2["assumptions"]["deposit_products"].index(m)
    keys = {str(r[0].value): r[0].row for r in ws.iter_rows(min_row=2)}
    ck("input workbook shows the three controls on each deposit", all(f"deposit.{idx}.{k}" in keys for k in KEYS)
       and ws.cell(row=keys[f"deposit.{idx}.deposit_beta"], column=4).value == 0.4)
    merged, _rep = diff_import(data, copy.deepcopy(c2))
    mm = next(x for x in merged["assumptions"]["deposit_products"] if x["name"] == m["name"])
    ck("input workbook zero-edit import keeps them", all(mm.get(k) == m[k] for k in KEYS))
    ws.cell(row=keys[f"deposit.{idx}.deposit_beta"], column=4).value = 0.25
    ws.cell(row=keys[f"deposit.{idx}.deposit_rate_cap"], column=4).value = None
    b2 = io.BytesIO(); wb.save(b2)
    merged2, rep2 = diff_import(b2.getvalue(), copy.deepcopy(c2))
    m2 = next(x for x in merged2["assumptions"]["deposit_products"] if x["name"] == m["name"])
    ck("a beta edited in the input workbook is imported", m2.get("deposit_beta") == 0.25, m2.get("deposit_beta"))
    ck("clearing the cap in the input workbook removes it", m2.get("deposit_rate_cap") is None and m2.get("deposit_rate_floor") == 0.005)

    # --- console ---
    html = Path("web/console_v2.html").read_text(encoding="utf-8")
    ck("beta, floor and cap sit under Pricing", "interest_balance_measure|deposit_beta|deposit_rate_floor|deposit_rate_cap/]" in html
       and '_depositPricingOrder' in html and 'rate = index × beta + spread' in html)
    ck("unsaved-changes list formats them as rates", '"index_spread","deposit_beta","deposit_rate_floor","deposit_rate_cap",' in html)
    ck("Bank Design Lab offers beta as a lever", "`${nm}: index beta`" in html)
    ck("the unsaved-changes Download keeps the whole engagement", 'if(V31 && typeof downloadFIW==="function") return downloadFIW();' in html)
    from foundry.v2.config_diff import _NAMES
    ck("change history names the three controls", _NAMES.get("deposit_beta") == "index beta")
    # A coverage regression must survive a successor release's BUILD_STAMP.
    # Exact release stamp validation lives in the new release's authoring test.
    import re
    stamp = re.fullmatch(r"foundry-r(\d+)-[a-z0-9-]+", Path("BUILD_STAMP").read_text().strip())
    ck("build stamp is r313 or a successor", bool(stamp and int(stamp.group(1)) >= 313))
    print(f"\n{p} passed, {f} failed")
    return 0 if f == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
