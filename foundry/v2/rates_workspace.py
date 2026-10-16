"""Opt-in, calendar-dated rate authoring. Published evidence is never a mutable curve.

No workspace means the r335 builder is called without changing a single operand.
Quarterly engines retain period-end index sampling; monthly previews are diagnostic,
not a claim that quarterly interest equals a monthly balance/interest simulation.
"""
from copy import deepcopy
from datetime import date
import math

INDICES = ('effr', 'sofr', 'prime')

def number(x, label):
    if isinstance(x, bool) or not isinstance(x, (int, float)) or not math.isfinite(x):
        raise ValueError(f'Rates & Curves: {label} must be a finite number')
    return float(x)

def iso(x):
    try:
        if not isinstance(x, str) or len(x) != 10: raise ValueError()
        return date.fromisoformat(x)
    except (ValueError, TypeError):
        raise ValueError('Rates & Curves: dates must be valid YYYY-MM-DD values') from None

def points(raw):
    if not isinstance(raw, list) or not raw or len(raw) > 4000:
        raise ValueError('Rates & Curves: supply 1–4000 dated rates')
    out = []
    for p in raw:
        if not isinstance(p, dict): raise ValueError('Rates & Curves: invalid dated rate')
        d, v = iso(p.get('date')), number(p.get('rate'), 'annualized rate')
        if abs(v) > 2: raise ValueError('Rates & Curves: annualized rates must be between -200% and 200%')
        out.append((d, v))
    if any(a[0] >= b[0] for a,b in zip(out,out[1:])):
        raise ValueError('Rates & Curves: dates must be unique and strictly increasing')
    return out

def lookup(pts, d, method='linear', tail='hold', target=None):
    if d < pts[0][0]:
        raise ValueError('Rates & Curves: schedule does not cover the first required date')
    for (d0,v0),(d1,v1) in zip(pts,pts[1:]):
        if d < d1:
            return v0 if method == 'step' else v0+(v1-v0)*(d-d0).days/(d1-d0).days
    d0,v0=pts[-1]
    if d == d0 or tail == 'hold': return v0
    if tail == 'error': raise ValueError('Rates & Curves: schedule ends before a required valuation date; choose an extension')
    step=.0005*max(0,(d-d0).days/365.2425*4)
    return max(target,v0-step) if v0>target else min(target,v0+step)

def validate(a, cfg):
    w=a.get('rates_workspace')
    if w is None: return
    if not isinstance(w,dict) or w.get('version') != 1: raise ValueError('Rates & Curves: unsupported workspace version')
    for field in ('published','anchor_overrides','basis','indices'):
        if field in w and not isinstance(w[field],dict): raise ValueError('Rates & Curves: '+field+' must be an object')
    pub=w.get('published') or {}
    for field in ('fomc','current_policy'):
        if field in pub and not isinstance(pub[field],dict): raise ValueError('Rates & Curves: invalid published snapshot')
    for k,c in (w.get('indices') or {}).items():
        if not isinstance(c,dict): raise ValueError('Rates & Curves: invalid index configuration')
        if c.get('schedule') is not None and not isinstance(c['schedule'],dict): raise ValueError('Rates & Curves: invalid schedule')
        ranges=c.get('overrides') or []
        if not isinstance(ranges,list) or len(ranges)>100: raise ValueError('Rates & Curves: at most 100 range overrides per index')
        for r in ranges:
            if not isinstance(r,dict): raise ValueError('Rates & Curves: invalid range override')
            if iso(r.get('start'))>iso(r.get('end')): raise ValueError('Rates & Curves: invalid range dates')
            number(r.get('rate'),'range override')
    if '_scenario_shift' in w or '_scenario_base' in w: raise ValueError('Rates & Curves: reserved scenario fields')
    if w.get('policy_interpolation','linear') not in ('linear','step'): raise ValueError('Rates & Curves: choose Step or Linear')
    for k,v in (w.get('anchor_overrides') or {}).items():
        if k != 'longer_run': iso(k)
        number(v,'policy anchor')
    for k,v in (w.get('basis') or {}).items():
        if k not in INDICES: raise ValueError('Rates & Curves: unknown basis index')
        number(v,'basis')
    if w.get('policy_tail','glide') not in ('hold','glide'): raise ValueError('Rates & Curves: invalid policy extension')
    for k,c in (w.get('indices') or {}).items():
        if k not in INDICES or not isinstance(c,dict): raise ValueError('Rates & Curves: unknown index')
        if c.get('source') not in ('existing','derived','schedule'): raise ValueError('Rates & Curves: invalid active source')
        # Inactive schedules remain saved but do not block another active source.
        if c.get('source')=='schedule':
            s=c.get('schedule') or {}; points(s.get('points'))
            if s.get('interpolation') not in ('step','linear'): raise ValueError('Rates & Curves: schedule requires Step or Linear')
            if s.get('extension') not in ('hold','derived','error'): raise ValueError('Rates & Curves: choose a schedule extension')
            if not str(s.get('label','')).strip(): raise ValueError('Rates & Curves: schedule source label is required')
            if s.get('units') not in ('percent','decimal'): raise ValueError('Rates & Curves: confirm Percent or Decimal units')
            if s.get('through'): iso(s['through'])
    ppy=a.get('periods_per_year',4);n=a.get('n_periods',12)
    if ppy not in (4,12) or isinstance(n,bool) or not isinstance(n,int) or not 1<=n<=7*ppy:
        raise ValueError('Rates & Curves: invalid engagement cadence or horizon')
    # Validate the complete computation horizon, not just the first displayed year.
    build(a,cfg,int(a.get('periods_per_year') or 4), validate_first=False)

def build(a, cfg, ppy, legacy_builder=None, validate_first=False):
    from .engine_q_a import _build_rate_curve_set
    from .timebase import model_period_end_date
    base=deepcopy(a); w=base.pop('rates_workspace',None)
    legacy=(legacy_builder or _build_rate_curve_set)(base,cfg,ppy)
    if not w: return legacy
    rc=base.get('rate_curves') or {}; f=rc.get('fomc') or {}; cp=rc.get('current_policy') or {}
    published=w.get('published') or w.get('_scenario_base') or {}; f=published.get('fomc') or f; cp=published.get('current_policy') or cp
    anchors=dict(f.get('anchors') or {})
    pd=cp.get('observation_date') or cp.get('statement_date')
    if pd and cp.get('mid') is not None: anchors[pd]=cp['mid']
    ov=w.get('anchor_overrides') or {}
    anchors.update({d:v for d,v in ov.items() if d!='longer_run'})
    lr=ov.get('longer_run',f.get('lr'))
    offsets={'effr':0.,'sofr':.0005,'prime':.03}; offsets.update(rc.get('offsets') or {}); offsets.update(w.get('basis') or {})
    half=rc.get('half_band',max(0,float(cp.get('top') or 0)-float(cp.get('mid') or 0)))
    shock=number(w.get('_scenario_shift',0),'scenario shift')
    out={}
    for k in INDICES:
        c=(w.get('indices') or {}).get(k,{'source':'existing'}); source=c.get('source','existing')
        if source=='existing': out[k]=legacy[k]; continue
        if (source=='derived' or (c.get('schedule') or {}).get('extension')=='derived') and (not anchors or lr is None):
            raise ValueError('Rates & Curves: Fed-based projection needs dated policy anchors and a longer-run value')
        off=number(offsets[k],'basis')+(float(half) if k=='prime' else 0)
        pp=points([{'date':d,'rate':v+off} for d,v in sorted(anchors.items())]) if anchors else [(date(1900,1,1),0.)]
        # Current Prime anchor uses the actual range top; year-end width stays an explicit assumption.
        if k=='prime' and pd and pd not in ov and cp.get('top') is not None:
            pp=[(d,float(cp['top'])+offsets[k] if d==iso(pd) else v) for d,v in pp]
        target=number(lr if lr is not None else 0,'longer-run')+off
        def derived(d,pp=pp,target=target):
            # Historical dates before the first policy observation retain the first snapshot value.
            return lookup(pp,max(d,pp[0][0]),w.get('policy_interpolation','linear'),w.get('policy_tail','glide'),target)
        s=c.get('schedule') or {}; sp=points(s.get('points')) if source=='schedule' else None
        def at(d,derived=derived,sp=sp,s=s,source=source,ranges=deepcopy(c.get('overrides') or [])):
            value=derived(d) if source=='derived' else 0.0
            if source=='schedule':
                end=iso(s['through']) if s.get('through') else sp[-1][0]
                if end < sp[-1][0]: raise ValueError('Rates & Curves: coverage end precedes last entered rate')
                if d>end and s['extension']=='derived': value=derived(d)
                else:
                    if d>end and s['extension']=='error': raise ValueError('Rates & Curves: incomplete engagement schedule')
                    value=lookup(sp,d,s['interpolation'],'hold' if d<=end else s['extension'])
            for r in ranges:
                if iso(r['start']) <= d <= iso(r['end']): value=number(r['rate'],'range override')
            return max(0,value+shock) if shock else value
        for r in c.get('overrides') or []:
            if iso(r['start'])>iso(r['end']): raise ValueError('Rates & Curves: override end precedes start')
            number(r['rate'],'range override')
        ranges=c.get('overrides') or []
        if any(max(iso(x['start']),iso(y['start']))<=min(iso(x['end']),iso(y['end'])) for i,x in enumerate(ranges) for y in ranges[i+1:]):
            raise ValueError('Rates & Curves: overlapping index overrides')
        def fn(t,at=at): return at(model_period_end_date(cfg,max(1,int(t)),ppy))
        fn.at_date=at; out[k]=fn
        n=int(a.get('n_periods') or a.get('n_quarters') or 12)
        for t in range(1,n+1): fn(t)
    out['sofr']._curves=out
    return out

def preview(cfg):
    from .engine_q_a import _build_rate_curve_set
    from .timebase import model_period_end_date
    a=cfg.get('assumptions') or {}; validate(a,cfg)
    ppy=int(a.get('periods_per_year') or 4); n=int(a.get('n_periods') or a.get('n_quarters') or 12)
    if ppy not in (4,12) or not 1<=n<=7*ppy: raise ValueError('Rates & Curves: invalid engagement cadence or horizon')
    engine=_build_rate_curve_set(a,cfg,ppy)
    monthly_cfg=deepcopy(cfg)
    if ppy==4:
        from .timebase import model_period_year_month
        y,m=model_period_year_month(cfg,1,ppy)
        monthly_cfg.setdefault('charter_profile',{})['target_opening']=f'{y:04d}-{m:02d}-01'
    monthly=_build_rate_curve_set(a,monthly_cfg,12)
    w=a.get('rates_workspace') or {}; rows=[]
    for m in range(1,n*(12//ppy)+1):
        d=model_period_end_date(monthly_cfg,m,12).isoformat(); row={'month':m,'date':d}
        for k in INDICES:
            c=(w.get('indices') or {}).get(k,{}); src=c.get('source','existing')
            tag={'existing':'Saved curve (r335 behavior)','derived':'Fed-based projection','schedule':'Engagement schedule'}[src]
            sch=c.get('schedule') or {}
            if src=='schedule' and sch.get('extension')=='derived' and d>(sch.get('through') or sch['points'][-1]['date']): tag='Fed-based projection · after schedule'
            for r in c.get('overrides') or []:
                if src!='existing' and r['start']<=d<=r['end']: tag='Month-range override'
            row[k]={'rate':monthly[k](m),'source':tag,'label':(c.get('schedule') or {}).get('label','') if src=='schedule' else ''}
        rows.append(row)
    comparison=[]
    if w and ((w.get('published') or a.get('rate_curves') or {}).get('fomc') or {}).get('anchors'):
        dc=deepcopy(monthly_cfg);dw=dc['assumptions']['rates_workspace']
        dw['indices']={k:{'source':'derived'} for k in INDICES}
        df=_build_rate_curve_set(dc['assumptions'],dc,12)
        comparison=[{'date':r['date'],**{k:df[k](r['month']) for k in INDICES}} for r in rows]
    native=[{'period':t,'date':model_period_end_date(cfg,t,ppy).isoformat(),**{k:engine[k](t) for k in INDICES}} for t in range(1,n+1)]
    return {'monthly':rows,'comparison':comparison,'native':native,'periods_per_year':ppy,'sampling':'Period-end index sampling; product balance measure and accrual rules apply separately.','workspace':deepcopy(w)}

def audit_sheets(wb, cfg):
    """Export the same canonical preview used by the authoring UI, without rate rounding."""
    from openpyxl.styles import Font, PatternFill, Alignment
    import json
    data=preview(cfg); a=cfg.get('assumptions') or {}; w=a.get('rates_workspace') or {}
    rc=a.get('rate_curves') or {}; pub=w.get('published') or rc
    ws=wb.create_sheet('Rate Assumptions')
    ws.append(['Layer','Index / target','Value / assumption','Source / treatment'])
    f=pub.get('fomc') or {}; cp=pub.get('current_policy') or {}
    ws.append(['Published evidence','Source',pub.get('evidence','Saved snapshot; original publication status not independently verified'),'FRED · FEDTARMD / FEDTARMDLR / DFEDTARU / DFEDTARL'])
    ws.append(['Published evidence','Retrieved',pub.get('retrieved_at'),'Stored snapshot; no network fetch during run'])
    ws.append(['Published evidence','SEP release',f.get('source_vintage'),'Not asserted where absent'])
    for k,v in sorted((f.get('anchors') or {}).items()): ws.append(['Published policy anchor',k,v,'Year-end target midpoint; annualized decimal rate'])
    ws.append(['Published policy anchor','Longer run',f.get('lr'),'Annualized decimal rate'])
    for k in ('mid','top','bottom','observation_date'): ws.append(['Current policy',k,cp.get(k),'Observed target range, not EFFR'])
    for k,v in (w.get('anchor_overrides') or {}).items(): ws.append(['Engagement policy override',k,v,'Replaces policy input, preserves published evidence'])
    ws.append(['Derivation','Interpolation',w.get('policy_interpolation','Legacy linear'),'Calendar days'])
    ws.append(['Derivation','Policy extension',w.get('policy_tail','Legacy 5bp/quarter glide'),'Foundry assumption'])
    ws.append(['Derivation','Prime range half-width',rc.get('half_band'),'Added to year-end midpoint before Prime basis'])
    for k in INDICES:
        c=(w.get('indices') or {}).get(k,{}); s=c.get('schedule') or {}
        ws.append(['Active source',k,c.get('source','existing'),'Scenario shifts apply subsequently; product pricing is separate'])
        ws.append(['Index basis',k,(w.get('basis') or rc.get('offsets') or {}).get(k),'Annualized decimal offset; bypassed by engagement schedule'])
        if c.get('source','existing')=='existing': ws.append(['Saved curve',k,json.dumps(rc.get(k) or {},sort_keys=True),'Exact legacy dated/path behavior retained'])
        if s:
            for field in ('label','units','input_cadence','interpolation','extension','through','entered_at'):
                ws.append(['Engagement schedule',k+' · '+field,s.get(field),'Active' if c.get('source')=='schedule' else 'Saved, inactive'])
            for pt in s.get('points') or []: ws.append(['Entered schedule rate',k+' · '+pt['date'],pt['rate'],'Annualized decimal rate'])
        for r in c.get('overrides') or []: ws.append(['Month-range override',k+' · '+r['start']+' / '+r['end'],r['rate'],'Applies after active source; inactive in saved-curve mode'])
    sheets=[ws]
    ms=wb.create_sheet('Effective Rates');sheets.append(ms)
    ms.append(['Month','Date','Index','Annualized rate','Source','Engagement source label','Fed-based comparison rate','Difference (bp)'])
    for r in data['monthly']:
        for k in INDICES:
            comp=data['comparison'][r['month']-1][k] if data['comparison'] else None
            ms.append([r['month'],r['date'],k,r[k]['rate'],r[k]['source'],r[k]['label'],comp,(r[k]['rate']-comp)*10000 if comp is not None else None])
    ns=wb.create_sheet('Engine Rates');sheets.append(ns)
    ns.append(['Period','Period end','EFFR','SOFR','Prime','Sampling'])
    for r in data['native']: ns.append([r['period'],r['date'],r['effr'],r['sofr'],r['prime'],data['sampling']])
    for sh in sheets:
        sh.freeze_panes='C2';sh.auto_filter.ref=sh.dimensions
        for cell in sh[1]: cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='30312F')
        for col,width in zip('ABCDEF',[26,34,44,65,32,55]): sh.column_dimensions[col].width=width
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(cell.value,str): cell.data_type='s'  # labels are evidence, never workbook formulas
    for row in ms.iter_rows(min_row=2):
        row[3].number_format='0.000000%'; row[6].number_format='0.000000%';row[7].number_format='0.000000'
    ms.column_dimensions['G'].width=27;ms.column_dimensions['H'].width=24
    for row in ns.iter_rows(min_row=2):
        for cell in row[2:5]:cell.number_format='0.000000%'
