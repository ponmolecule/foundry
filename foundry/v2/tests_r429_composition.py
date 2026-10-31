"""User-selected allocation rounding and causal active population sources."""
import copy,json,unittest
from pathlib import Path
from .tests_r428_pool_sources import setup as acquisition_setup
from .cac_feeder import cac_auc_rollforward,cac_customer_count_measure_series_map
from .run_q import run_v2
from .validate_q import validate_errors_v2
from .tests_r425_cross_product_links import account


def setup():
 a=acquisition_setup();f=a['cac_feeds']['Practices'];f.update(series_id='practices-auc',customer_count_series_id='practices-count',driver_specs={'attrition_rate':{'source':'entered','trajectory':'flat','value':.1/12,'period':'month'}})
 names=['Small','Medium','Large','Largest'];sizes=[7.5,17.5,37,75]
 bands=[dict(id='b'+str(i),series_id='count-b'+str(i),name=name,opening=0,units_per_customer=size) for i,(name,size) in enumerate(zip(names,sizes))]
 f['population_composition']={'rounding':'none','population_name':'Associated people','population_series_id':'people','bands':bands,'channel_shares':{
 'direct':dict(zip([b['id'] for b in bands],[0,0,.667694,.332306])),
 'digital':dict(zip([b['id'] for b in bands],[.792386,.207614,0,0])),
 'referrals':dict(zip([b['id'] for b in bands],[.735415,.192687,.048006,.023892]))}}
 return a

class Tests(unittest.TestCase):
 def test_none_identity_and_live_mix(self):
  a=setup();f=a['cac_feeds']['Practices'];old=copy.deepcopy(f);old.pop('population_composition');before=copy.deepcopy(a)
  r=cac_auc_rollforward(f,60,12,assumptions=a);legacy=cac_auc_rollforward(old,60,12,assumptions=a)
  self.assertEqual(a,before);self.assertEqual(r['customer_end_by_month'],legacy['customer_end_by_month']);self.assertEqual(r['auc_end_by_month'],legacy['auc_end_by_month'])
  self.assertTrue(all(m['rounding_increment']==0 for m in r['monthly']))
  bands=r['population_composition']['bands']
  for i,m in enumerate(r['monthly']):
   self.assertAlmostEqual(sum(b['customer_end_by_month'][i] for b in bands.values()),m['end_cust'])
   self.assertAlmostEqual(r['population_composition']['population']['customer_end_by_month'][i],sum(b['customer_end_by_month'][i]*b['units_per_customer'] for b in bands.values()))
  first=bands['b0']['monthly'][0];second=bands['b0']['monthly'][1];self.assertAlmostEqual(second['lost'],first['ending']*.1/12)
 def test_round_up_only_by_choice_and_raw_referrals(self):
  a=setup();f=a['cac_feeds']['Practices'];raw=cac_auc_rollforward(f,60,12,assumptions=a)
  f['population_composition']['rounding']='up';r=cac_auc_rollforward(f,60,12,assumptions=a)
  self.assertGreater(r['monthly'][0]['new_cust'],raw['monthly'][0]['new_cust']);self.assertEqual(r['monthly'][0]['total_spend'],21624)
  self.assertAlmostEqual(r['monthly'][0]['channels'][0]['equation_new_customers'],.937)
  self.assertAlmostEqual(r['derived_series']['direct']['canonical_values'][0],.28)
  for month in r['monthly']:
   self.assertTrue(all(v==int(v) for ch in month['channels'] for v in ch['band_additions'].values()))
   self.assertAlmostEqual(month['new_cust']-month['equation_new_cust'],month['rounding_increment'])
  f['population_composition'].pop('rounding');default=cac_auc_rollforward(f,60,12,assumptions=a);self.assertEqual(default['customer_end_by_month'],raw['customer_end_by_month'])
 def test_cadence_and_opening(self):
  a=setup();f=a['cac_feeds']['Practices'];f['beginning_customers']=10;f['population_composition']['bands'][0]['opening']=10
  m=cac_auc_rollforward(f,60,12,assumptions=a);q=cac_auc_rollforward(f,20,4,assumptions=a)
  self.assertEqual(m['population_composition']['bands']['b0']['monthly'],q['population_composition']['bands']['b0']['monthly'])
  self.assertAlmostEqual(m['population_composition']['bands']['b0']['monthly'][0]['lost'],10*.1/12)
  maps=cac_customer_count_measure_series_map(a,20,4);self.assertEqual(maps['people']['period_end'],q['population_composition']['population']['customer_level_by_period'])
 def test_public_catalog_and_product_cost(self):
  from .source_catalog import catalog
  a=setup();c=json.loads(Path('foundry/fixtures/core_bank_test_base.json').read_text());c['assumptions'].update(a)
  st=account('population-servicing','unused',100,'year');st['driver']={'source':'customer_acquisition_count','ref':'people','measure':'period_end','trajectory':'flat','params':{}}
  c['assumptions']['deposit_products'][0]['fee_streams']=[st]
  self.assertEqual(validate_errors_v2(copy.deepcopy(c)),[])
  r=run_v2(c);pop=r['customer_acquisition']['Practices']['populationComposition']['population'];self.assertEqual(r['fee_stream_quantities']['series']['population-servicing'],pop['customer_level_by_period'])
  rows={x['series_id']:x for x in catalog(c['assumptions']) if x['driver']=='customer_acquisition_count'}
  self.assertIn('people',rows);self.assertIn('count-b0',rows);self.assertIn('Associated people',rows['people']['label'])
 def test_deposit_counts_and_workbook(self):
  from .tests_r380_customer_deposits import config
  from .deposit_balance import customer_balance
  from .audit_workbook import calculation_audit_workbook
  from io import BytesIO
  from openpyxl import load_workbook
  c=config(4);c['assumptions'].update(setup());a=c['assumptions'];p=a['deposit_products'][0]
  p['balance_driver'].update(series_id='people',balance_per_customer_spec={'source':'entered','trajectory':'flat','value':100})
  counts=cac_auc_rollforward(a['cac_feeds']['Practices'],20,4,assumptions=a)['population_composition']['population']['customer_end_by_month']
  detail=customer_balance(p,a,20,4)
  for i,value in enumerate(detail['endingBalanceDollars']):self.assertAlmostEqual(value,counts[i*3+2]*100)
  r=run_v2(c);wb=calculation_audit_workbook(c,r);buf=BytesIO();wb.save(buf);buf.seek(0);wb=load_workbook(buf,data_only=True)
  self.assertIn('Population Composition',wb.sheetnames)
  self.assertTrue(any('Associated population' in str(cell.value) for row in wb['Population Composition'] for cell in row))
 def test_rejections(self):
  for case in ('share','opening','negative','duplicate','rounding','interpolated','missing_channel'):
   with self.subTest(case=case):
    a=setup();f=a['cac_feeds']['Practices'];comp=f['population_composition']
    if case=='share':comp['channel_shares']['direct']['b0']=.1
    if case=='opening':comp['bands'][0]['opening']=1
    if case=='negative':comp['bands'][0]['units_per_customer']=-1
    if case=='duplicate':comp['bands'][0]['series_id']=comp['bands'][1]['series_id']
    if case=='rounding':comp['rounding']='automatic'
    if case=='interpolated':f['intra_period_path']='straight_line'
    if case=='missing_channel':comp['channel_shares'].pop('digital')
    with self.assertRaises(ValueError):cac_auc_rollforward(f,60,12,assumptions=a)
if __name__=='__main__':unittest.main()
